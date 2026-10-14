"""Work hours report: one person's shifts over a date range, on screen and as Excel."""

from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.utils import timezone
from django.utils.dateformat import format as date_format
from django.utils.translation import get_language, gettext as _, pgettext

from shifts.access import can_check_in, display_name
from shifts.clock import format_clock
from shifts.models import ShiftAssignment
from shifts.services import (
    CHECKED_IN,
    COMPLETED,
    MISSED,
    clock_length,
    effective_interval,
    effective_times,
    settle_check_in,
    shift_outcome,
)

REPORT_MAX_DAYS = 366

PRESENT = "present"
ABSENT = "absent"
UPCOMING = "upcoming"
SCHEDULED = "scheduled"
# Shifts that count toward the totals. Scheduled is a finished shift for someone
# whose role does not use check-in, so there is nothing to check it against.
COUNTED = (PRESENT, SCHEDULED)


def status_label(status):
    return {
        PRESENT: _("Present"),
        ABSENT: _("Absent"),
        UPCOMING: _("Upcoming"),
        SCHEDULED: _("Scheduled"),
    }[status]


@dataclass
class ReportRow:
    number: int
    day: date
    weekday: str
    starts: str
    ends: str
    hours: Decimal
    status: str

    @property
    def status_text(self):
        return status_label(self.status)

    @property
    def counted(self):
        return self.status in COUNTED


@dataclass
class WorkHoursReport:
    person: object
    start: date
    end: date
    night_only: bool
    rows: list = field(default_factory=list)

    @property
    def name(self):
        return display_name(self.person)

    @property
    def title(self):
        if self.night_only:
            return _("Statement of work hours in the night shifts")
        return _("Statement of work hours")

    @property
    def days_present(self):
        return len({row.day for row in self.rows if row.counted})

    @property
    def total_hours(self):
        return sum((row.hours for row in self.rows if row.counted), Decimal(0))


def _hours(minutes):
    """Whole hours stay whole (5, 12); otherwise up to two decimals (7.5)."""
    value = (Decimal(minutes) / Decimal(60)).quantize(Decimal("0.01"))
    return value.quantize(Decimal(1)) if value == value.to_integral() else value.normalize()


def _status(assignment, now, uses_check_in):
    # Same outcome the rota and My shifts show (shifts.services.shift_outcome).
    outcome = shift_outcome(assignment, now, uses_check_in)
    if outcome == CHECKED_IN:
        return PRESENT
    if outcome == COMPLETED:
        _start_dt, end_dt = effective_interval(assignment)
        attended = bool(assignment.checked_in_at) or bool(
            assignment.checked_out_at and assignment.checked_out_at >= end_dt
        )
        return PRESENT if attended else SCHEDULED
    if outcome == MISSED:
        return ABSENT
    return UPCOMING


def clamp_range(start, end):
    """Order the dates and keep the range to at most a year."""
    if end < start:
        start, end = end, start
    if (end - start).days >= REPORT_MAX_DAYS:
        end = start + timedelta(days=REPORT_MAX_DAYS - 1)
    return start, end


def build_report(person, start, end, night_only=False, now=None):
    now = now or timezone.now()
    start, end = clamp_range(start, end)
    query = ShiftAssignment.objects.filter(
        user=person, date__gte=start, date__lte=end,
    ).select_related("shift_type", "shift_type__department")
    if night_only:
        query = query.filter(shift_type__is_night=True)
    assignments = sorted(query, key=lambda row: (effective_interval(row)[0], row.pk))
    uses_check_in = can_check_in(person)
    report = WorkHoursReport(person=person, start=start, end=end, night_only=night_only)
    for number, assignment in enumerate(assignments, start=1):
        # A report is a read: settle an ended check-in in memory only.
        settle_check_in(assignment, now, save=False)
        start_t, end_t = effective_times(assignment)
        status = _status(assignment, now, uses_check_in)
        starts = format_clock(start_t)
        if status == ABSENT:
            # Nobody worked an absent shift, so it is worth no hours.
            hours = Decimal(0)
        elif status == PRESENT:
            # "From" shows when the agent pressed check in, but time is only
            # counted inside the shift: an early check-in starts at the shift
            # start, and a shift still running counts up to now, not its end.
            start_dt, end_dt = effective_interval(assignment)
            counted_from = start_dt
            if assignment.arrived_at:
                starts = format_clock(assignment.arrived_at)
                counted_from = max(assignment.arrived_at, start_dt)
            counted_to = min(end_dt, now)
            minutes = int((counted_to - counted_from).total_seconds() // 60)
            hours = _hours(max(minutes, 0))
        else:
            hours = _hours(clock_length(start_t, end_t)[0])
        report.rows.append(ReportRow(
            number=number,
            day=assignment.date,
            weekday=date_format(assignment.date, "l"),
            starts=starts,
            ends=format_clock(end_t),
            hours=hours,
            status=status,
        ))
    return report


def report_filename(report):
    return f"work-hours_{report.person.username}_{report.start:%Y-%m-%d}_{report.end:%Y-%m-%d}.xlsx"


def report_workbook(report):
    """The report as an .xlsx file, laid out like the paper statement it replaces."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    rtl = (get_language() or "").startswith("ar")
    book = Workbook()
    sheet = book.active
    sheet.title = _("Work hours")[:31]
    sheet.sheet_view.rightToLeft = rtl

    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    bold = Font(bold=True)
    fills = {
        PRESENT: PatternFill("solid", fgColor="C6EFCE"),
        ABSENT: PatternFill("solid", fgColor="FFC7CE"),
        UPCOMING: PatternFill("solid", fgColor="DDEBF7"),
        SCHEDULED: PatternFill("solid", fgColor="EDEDED"),
    }
    status_fonts = {
        PRESENT: Font(italic=True, color="006100"),
        ABSENT: Font(italic=True, color="9C0006"),
        UPCOMING: Font(italic=True, color="1F4E78"),
        SCHEDULED: Font(italic=True, color="404040"),
    }
    columns = 7

    def banner(row, text):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
        cell = sheet.cell(row=row, column=1, value=text)
        cell.alignment = center
        cell.font = bold
        for col in range(1, columns + 1):
            sheet.cell(row=row, column=col).border = box

    banner(1, f"{_('Name')}: {report.name}")
    banner(2, report.title)
    period = _("From %(start)s to %(end)s") % {
        "start": f"{report.start:%d/%m/%Y}", "end": f"{report.end:%d/%m/%Y}",
    }
    banner(3, period)

    headers = [_("No."), pgettext("weekday column", "Day"), _("Date"), _("Status"), _("From"), _("To"), _("Work hours")]
    header_row = 5
    for col, text in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=col, value=text)
        cell.font = bold
        cell.alignment = center
        cell.border = box

    row_index = header_row
    for row in report.rows:
        row_index += 1
        values = [row.number, row.weekday, row.day, row.status_text, row.starts, row.ends, float(row.hours)]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col, value=value)
            cell.alignment = center
            cell.border = box
        sheet.cell(row=row_index, column=3).number_format = "DD/MM/YYYY"
        sheet.cell(row=row_index, column=1).font = bold
        status_cell = sheet.cell(row=row_index, column=4)
        status_cell.fill = fills[row.status]
        status_cell.font = status_fonts[row.status]

    # Totals sit under the last two columns, as on the paper statement.
    totals_row = row_index + 2
    for offset, (label, value) in enumerate((
        (_("Days present"), report.days_present),
        (_("Sum of work hours"), float(report.total_hours)),
    )):
        value_cell = sheet.cell(row=totals_row + offset, column=columns - 1, value=value)
        label_cell = sheet.cell(row=totals_row + offset, column=columns, value=label)
        for cell in (value_cell, label_cell):
            cell.border = box
            cell.alignment = center
        label_cell.font = bold

    for col, width in zip("ABCDEFG", (6, 14, 13, 12, 12, 12, 16)):
        sheet.column_dimensions[col].width = width

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
