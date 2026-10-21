import re
from contextlib import contextmanager
from datetime import date, datetime, time, timedelta, timezone as dt_timezone
from decimal import Decimal
from unittest.mock import patch

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from core.models import Department, Role
from shifts.models import ShiftAssignment, ShiftType
from shifts.services import effective_interval

User = get_user_model()


@contextmanager
def _both_clocks(now, day):
    with patch("django.utils.timezone.now", return_value=now), \
         patch("django.utils.timezone.localdate", return_value=day):
        yield


def aware(y, m, d, hh, mm=0):
    return timezone.make_aware(datetime(y, m, d, hh, mm))


def make_department(name="Desk"):
    return Department.objects.create(name=name)


def make_role(name="Agent", **flags):
    return Role.objects.create(name=name, **flags)


def make_user(username, user_type, department=None, role=None, superuser=False, first_name="", last_name=""):
    user = User.objects.create_user(
        username=username,
        email=f"{username}@example.com",
        password="test-pass-123",
        user_type=user_type,
        department=department,
        role=role,
        first_name=first_name,
        last_name=last_name,
    )
    if superuser:
        user.is_superuser = True
        user.is_staff = True
        user.save()
    return user


def make_shift_type(department, name="Night", start=time(22, 0), end=time(6, 0), archived=False, is_night=None):
    if is_night is None:
        # Night is ticked by hand now; fixtures default to what the old automatic
        # rule said, so tests written before the change keep their meaning.
        from shifts.migrations._night_legacy import is_night_shift

        is_night = is_night_shift(start, end, time(18, 0), time(6, 0))
    return ShiftType.objects.create(
        department=department,
        name=name,
        start_time=start,
        end_time=end,
        colour="#336699",
        archived=archived,
        is_night=is_night,
    )


def make_assignment(user, shift_type, on, start=None, end=None):
    return ShiftAssignment.objects.create(
        user=user,
        shift_type=shift_type,
        date=on,
        start_time_override=start,
        end_time_override=end,
        times_set_by=user if start and end else None,
        times_set_at=timezone.now() if start and end else None,
    )


class ShiftTypeNightFlagTests(TestCase):
    def test_night_is_only_what_was_ticked(self):
        department = make_department()
        unticked = ShiftType.objects.create(
            department=department, name="Overnight", start_time=time(22, 0), end_time=time(6, 0), colour="#112233",
        )
        self.assertFalse(unticked.is_night)
        ticked = ShiftType.objects.create(
            department=department, name="Late", start_time=time(14, 0), end_time=time(22, 0), colour="#112233",
            is_night=True,
        )
        ticked.save()  # saving never recomputes it
        ticked.refresh_from_db()
        self.assertTrue(ticked.is_night)

    def test_colour_validator_rejects_a_non_hex_value(self):
        department = make_department("Colour")
        row = make_shift_type(department, name="Colour night")
        row.colour = "blue"
        with self.assertRaises(ValidationError):
            row.full_clean()

    def test_names_are_unique_per_department(self):
        department = make_department()
        make_shift_type(department, name="Night")
        with self.assertRaises(Exception):
            ShiftType.objects.create(
                department=department,
                name="Night",
                start_time=time(1, 0),
                end_time=time(2, 0),
                colour="#112233",
            )


class WeekAndHoursTests(TestCase):
    def test_wednesday_snaps_to_previous_saturday(self):
        from shifts.services import normalize_start, saturday_of, week_dates

        wednesday = date(2026, 10, 7)  # Wednesday
        saturday = date(2026, 10, 3)
        self.assertEqual(saturday_of(wednesday), saturday)
        self.assertEqual(saturday_of(saturday), saturday)
        self.assertEqual(normalize_start(wednesday, "week"), saturday)
        self.assertEqual(normalize_start(date(2026, 10, 7), "month"), date(2026, 10, 1))
        self.assertEqual(week_dates(saturday)[0], saturday)
        self.assertEqual(week_dates(saturday)[-1], date(2026, 10, 9))  # Friday
        self.assertEqual(len(week_dates(saturday)), 7)

    def test_clock_length_splits_after_midnight(self):
        from shifts.services import clock_length

        self.assertEqual(clock_length(time(22, 0), time(6, 0)), (480, 360))
        self.assertEqual(clock_length(time(9, 0), time(12, 30)), (210, 0))

    def test_half_open_touch_is_not_overlap(self):
        from shifts.services import interval_for, overlaps

        monday = date(2026, 10, 5)
        left = interval_for(monday, time(22, 0), time(6, 0))
        right = interval_for(date(2026, 10, 6), time(6, 0), time(14, 0))
        early = interval_for(date(2026, 10, 6), time(1, 0), time(9, 0))
        self.assertFalse(overlaps(left, right))
        self.assertTrue(overlaps(left, early))


class LandingTests(TestCase):
    def test_each_role_lands_on_its_page(self):
        department = make_department("Land")
        manager_role = make_role(name="Shift lead", can_manage_shifts=True)
        manager = make_user("lead", "support", department, manager_role)
        support = make_user("agent", "support", department)
        branch = make_user("branch", "branch")
        self.client.force_login(manager)
        self.assertRedirects(self.client.get(reverse("shifts_home")), reverse("shifts_rota"))
        self.client.force_login(support)
        self.assertRedirects(self.client.get(reverse("shifts_home")), reverse("shifts_mine"))
        self.client.force_login(branch)
        self.assertRedirects(self.client.get(reverse("shifts_home")), reverse("shifts_available"))

    def test_sidebar_link_is_present_for_support(self):
        department = make_department("Nav")
        support = make_user("navagent", "support", department)
        self.client.force_login(support)
        response = self.client.get(reverse("shifts_mine"))
        self.assertContains(response, 'data-nav-key="shifts"')


class ShiftTypeCrudTests(TestCase):
    def test_types_page_opens_on_a_department_that_has_types(self):
        make_department("Aaa empty")
        filled = make_department("Zzz filled")
        make_shift_type(filled, name="Desk night")
        role = make_role("Types default", can_manage_shifts=True)
        manager = make_user("types-default", "support", filled, role)
        self.client.force_login(manager)
        response = self.client.get(reverse("shifts_types"))
        self.assertContains(response, "Desk night")
        self.assertContains(response, f'value="{filled.pk}" selected')

    def test_support_cannot_open_types(self):
        department = make_department("Types")
        support = make_user("types-agent", "support", department)
        self.client.force_login(support)
        self.assertEqual(self.client.get(reverse("shifts_types")).status_code, 403)

    def test_manager_creates_and_archives(self):
        department = make_department("Types2")
        role = make_role("Types lead", can_manage_shifts=True)
        manager = make_user("types-lead", "support", department, role)
        self.client.force_login(manager)
        response = self.client.post(reverse("shifts_type_add"), {
            "department": department.pk,
            "name": "Night",
            "start_time": "22:00",
            "end_time": "06:00",
            "colour": "#336699",
            "is_night": "on",
        })
        self.assertEqual(response.status_code, 302)
        row = ShiftType.objects.get(name="Night")
        self.assertTrue(row.is_night)
        archive = self.client.post(reverse("shifts_type_archive", args=[row.pk]))
        self.assertEqual(archive.status_code, 302)
        row.refresh_from_db()
        self.assertTrue(row.archived)

    def test_equal_times_rejected(self):
        department = make_department("Types3")
        role = make_role("Types lead 2", can_manage_shifts=True)
        manager = make_user("types-lead-2", "support", department, role)
        self.client.force_login(manager)
        response = self.client.post(reverse("shifts_type_add"), {
            "department": department.pk,
            "name": "Bad",
            "start_time": "09:00",
            "end_time": "09:00",
            "colour": "#336699",
        })
        self.assertEqual(response.status_code, 200)
        self.assertFalse(ShiftType.objects.filter(name="Bad").exists())
        self.assertContains(response, "Start and end must differ.")

    def test_short_name_rejected(self):
        department = make_department("Types4")
        role = make_role("Types lead 3", can_manage_shifts=True)
        manager = make_user("types-lead-3", "support", department, role)
        self.client.force_login(manager)
        response = self.client.post(reverse("shifts_type_add"), {
            "department": department.pk,
            "name": " N ",
            "start_time": "09:00",
            "end_time": "17:00",
            "colour": "#336699",
        })
        self.assertEqual(response.status_code, 200)
        self.assertFalse(ShiftType.objects.filter(department=department).exists())
        self.assertContains(response, "at least 2 characters")


class RotaGridTests(TestCase):
    def setUp(self):
        self.department = make_department("Rota")
        self.role = make_role("Rota lead", can_manage_shifts=True)
        self.manager = make_user("rota-lead", "support", self.department, self.role)
        self.agent = make_user("rota-agent", "support", self.department)
        self.night = make_shift_type(self.department)
        self.client.force_login(self.manager)

    def test_rota_opens_on_a_department_that_has_types(self):
        empty = make_department("Aaa empty rota")
        response = self.client.get(reverse("shifts_rota"))
        self.assertContains(response, f'value="{self.department.pk}" selected')
        self.assertNotContains(response, f'value="{empty.pk}" selected')
        self.assertNotIn(">Show<", response.content.decode())

    def test_week_query_normalizes_to_saturday(self):
        response = self.client.get(reverse("shifts_rota"), {"view": "week", "start": "2026-10-07", "department": self.department.pk})
        self.assertContains(response, "2026-10-03")
        self.assertContains(response, "2026-10-09")

    def test_cell_post_blocks_overlap(self):
        make_assignment(self.agent, self.night, date(2026, 10, 5))
        response = self.client.post(reverse("shifts_cell"), {
            "user": self.agent.pk,
            "date": "2026-10-06",
            "department": self.department.pk,
            "shift_type": self.night.pk,
            "start_time_override": "01:00",
            "end_time_override": "09:00",
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "overlaps")
        self.assertEqual(ShiftAssignment.objects.filter(user=self.agent).count(), 1)

    def test_support_gets_403(self):
        self.client.force_login(self.agent)
        self.assertEqual(self.client.get(reverse("shifts_rota")).status_code, 403)

    def test_grid_query_budget(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        for index in range(6):
            person = make_user(f"budget-{index}", "support", self.department)
            make_assignment(person, self.night, date(2026, 10, 5))
        with CaptureQueriesContext(connection) as captured:
            response = self.client.get(reverse("shifts_rota_grid"), {
                "view": "week",
                "start": "2026-10-03",
                "department": self.department.pk,
            })
        self.assertEqual(response.status_code, 200)
        self.assertLessEqual(len(captured), 12)


class CopyWeekTests(TestCase):
    def test_worker_times_are_not_copied(self):
        from shifts.services import plan_copy

        department = make_department("Copy")
        agent = make_user("copy-agent", "support", department)
        manager = make_user("copy-lead", "support", department, make_role("Copy lead", can_manage_shifts=True))
        night = make_shift_type(department)
        source = make_assignment(agent, night, date(2026, 10, 3), time(18, 0), time(23, 0))
        planned = plan_copy([source], timedelta(days=7), manager)
        self.assertEqual(planned[0]["start"], None)
        self.assertEqual(planned[0]["end"], None)
        self.assertIsNone(planned[0]["times_set_by_id"])
        self.assertEqual(planned[0]["date"], date(2026, 10, 10))

    def test_repeat_rejects_13(self):
        department = make_department("Repeat")
        manager = make_user("repeat-lead", "support", department, make_role("Repeat lead", can_manage_shifts=True))
        self.client.force_login(manager)
        response = self.client.post(reverse("shifts_repeat_week"), {
            "department": department.pk,
            "start": "2026-10-03",
            "weeks": "13",
        })
        self.assertEqual(response.status_code, 302)
        followed = self.client.get(response["Location"])
        self.assertContains(followed, "1 to 12")
        self.assertEqual(ShiftAssignment.objects.count(), 0)

    def test_friday_tail_does_not_become_saturday(self):
        department = make_department("Tail")
        manager = make_user("tail-lead", "support", department, make_role("Tail lead", can_manage_shifts=True))
        agent = make_user("tail-agent", "support", department)
        night = make_shift_type(department, name="Tail night")
        make_assignment(agent, night, date(2026, 10, 9))  # Friday
        self.client.force_login(manager)
        self.client.post(reverse("shifts_copy_week"), {
            "department": department.pk,
            "start": "2026-10-10",
        })
        self.assertFalse(ShiftAssignment.objects.filter(user=agent, date=date(2026, 10, 10)).exists())
        copied = ShiftAssignment.objects.get(user=agent, date=date(2026, 10, 16))
        self.assertIsNone(copied.checked_in_at)
        self.assertIsNone(copied.start_time_override)

    def test_manager_set_override_is_copied_and_repeated(self):
        department = make_department("Mgr copy")
        manager = make_user("mgr-copy-lead", "support", department, make_role("Mgr copy lead", can_manage_shifts=True))
        agent = make_user("mgr-copy-agent", "support", department)
        night = make_shift_type(department, name="Mgr night")
        source = make_assignment(agent, night, date(2026, 10, 3), time(18, 0), time(23, 0))
        source.times_set_by = manager
        source.save()
        self.client.force_login(manager)
        self.client.post(reverse("shifts_copy_week"), {
            "department": department.pk,
            "start": "2026-10-10",
        })
        copied = ShiftAssignment.objects.get(user=agent, date=date(2026, 10, 10))
        self.assertEqual(copied.start_time_override, time(18, 0))
        self.assertEqual(copied.end_time_override, time(23, 0))
        self.assertEqual(copied.times_set_by_id, manager.pk)
        self.client.post(reverse("shifts_repeat_week"), {
            "department": department.pk,
            "start": "2026-10-10",
            "weeks": "1",
        })
        repeated = ShiftAssignment.objects.get(user=agent, date=date(2026, 10, 17))
        self.assertEqual(repeated.start_time_override, time(18, 0))
        self.assertEqual(repeated.end_time_override, time(23, 0))
        self.assertEqual(repeated.times_set_by_id, manager.pk)


class RotationTests(TestCase):
    def test_apply_does_not_overwrite(self):
        department = make_department("Keep")
        role = make_role("Keep lead", can_manage_shifts=True)
        manager = make_user("keep-lead", "support", department, role)
        agent = make_user("keep-agent", "support", department)
        night = make_shift_type(department, name="Keep night")
        day = timezone.localdate() + timedelta(days=7)
        existing = make_assignment(agent, night, day)
        self.client.force_login(manager)
        response = self.client.post(reverse("shifts_calc_rotation"), {
            "department": department.pk,
            "start": day.isoformat(),
            "end": day.isoformat(),
            "shift_type": night.pk,
            "users": [agent.pk],
            "action": "apply",
        })
        self.assertContains(response, "Added 0 shifts.")
        self.assertEqual(ShiftAssignment.objects.count(), 1)
        existing.refresh_from_db()
        self.assertEqual(existing.shift_type_id, night.pk)
    def test_archived_or_other_department_type_returns_404_and_writes_nothing(self):
        department = make_department("Gate")
        role = make_role("Gate lead", can_manage_shifts=True)
        manager = make_user("gate-lead", "support", department, role)
        agent = make_user("gate-agent", "support", department)
        archived = make_shift_type(department, name="Gate old night", archived=True)
        elsewhere = make_shift_type(make_department("Gate other"), name="Other night")
        day = make_shift_type(department, name="Gate day", start=time(9, 0), end=time(17, 0))
        on = (timezone.localdate() + timedelta(days=7)).isoformat()
        self.client.force_login(manager)
        payload = {"department": department.pk, "start": on, "end": on, "users": [agent.pk], "action": "apply"}
        for bad in (archived, elsewhere):
            response = self.client.post(reverse("shifts_calc_rotation"), {**payload, "shift_type": bad.pk})
            self.assertEqual(response.status_code, 404)
        self.assertEqual(ShiftAssignment.objects.count(), 0)
        # Bulk assign takes any active type, day shifts included.
        response = self.client.post(reverse("shifts_calc_rotation"), {**payload, "shift_type": day.pk})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(ShiftAssignment.objects.get().shift_type_id, day.pk)
    def test_calculator_lists_active_types_and_preview_uses_names(self):
        department = make_department("Picker")
        manager = make_user("picker-lead", "support", department, make_role("Picker lead", can_manage_shifts=True))
        agent = make_user("picker-agent", "support", department, first_name="Nora", last_name="Saleh")
        make_user("picker-day-only", "branch")
        night = make_shift_type(department, name="Picker night")
        make_shift_type(department, name="Picker day", start=time(9, 0), end=time(17, 0))
        make_shift_type(department, name="Picker old", archived=True)
        self.client.force_login(manager)
        page = self.client.get(reverse("shifts_calculator"))
        rotation = page.content.decode().split('id="shifts-rotation"', 1)[1].split("</section>", 1)[0]
        self.assertIn("Bulk assign", rotation)
        self.assertIn("Picker night (night)", rotation)
        self.assertIn("Picker day", rotation)
        self.assertNotIn("Picker old", rotation)
        self.assertIn("Nora Saleh", rotation)
        self.assertIn('name="users"', rotation)
        self.assertIn("Select all", rotation)
        self.assertIn('name="skip_weekday"', rotation)
        self.assertIn('name="skip_date"', rotation)
        on = (timezone.localdate() + timedelta(days=7)).isoformat()
        preview = self.client.post(reverse("shifts_calc_rotation"), {
            "department": department.pk,
            "start": on,
            "end": on,
            "shift_type": night.pk,
            "users": [agent.pk],
            "action": "preview",
        })
        self.assertEqual(preview.status_code, 200)
        self.assertContains(preview, "Nora Saleh")
        self.assertContains(preview, "shifts-grid--mini")
        self.assertContains(preview, f'name="plan" value="{agent.pk}:{on}"')
        hx = self.client.get(reverse("shifts_calculator"), HTTP_HX_REQUEST="true")
        self.assertEqual(hx.status_code, 200)
        self.assertNotContains(hx, "<main")
        empty = self.client.post(reverse("shifts_calc_rotation"), {
            "department": department.pk,
            "start": "",
            "end": "",
            "shift_type": "",
            "action": "preview",
        })
        # No type picked (a department without types): a hint, not a silent 404.
        self.assertEqual(empty.status_code, 200)
        self.assertContains(empty, "Add a shift type to this department first")
    def test_copy_conflict_names_the_person_and_redirects(self):
        department = make_department("Named copy")
        manager = make_user("name-lead", "support", department, make_role("Name lead", can_manage_shifts=True))
        agent = make_user("name-agent", "support", department, first_name="Mona", last_name="Saleh")
        night = make_shift_type(department, name="Named night")
        make_assignment(agent, night, date(2026, 10, 3))
        make_assignment(agent, night, date(2026, 10, 10))
        self.client.force_login(manager)
        page = self.client.get(reverse("shifts_rota"), {"department": department.pk, "start": "2026-10-10"})
        self.assertContains(page, "Copy last week into this week?")
        self.assertContains(page, "Repeat this week forward?")
        self.client.post(reverse("set_language"), {"language": "ar"})
        arabic = self.client.get(reverse("shifts_rota"), {"department": department.pk, "start": "2026-10-10"})
        body = arabic.content.decode()
        self.assertEqual(body.count('data-confirm-title="تأكيد الإجراء"'), 2)
        self.assertEqual(body.count('data-confirm-ok="نعم، متابعة"'), 2)
        self.assertEqual(body.count('data-confirm-cancel="إلغاء"'), 2)
        self.client.post(reverse("set_language"), {"language": "en"})
        response = self.client.post(reverse("shifts_copy_week"), {
            "department": department.pk,
            "start": "2026-10-10",
        })
        self.assertEqual(response.status_code, 302)
        self.assertIn("/shifts/rota/", response["Location"])
        followed = self.client.get(response["Location"])
        self.assertContains(followed, "Mona Saleh")
        hx = self.client.post(reverse("shifts_copy_week"), {
            "department": department.pk,
            "start": "2026-10-10",
        }, HTTP_HX_REQUEST="true")
        self.assertEqual(hx.status_code, 200)
        self.assertNotContains(hx, 'id="shifts-shell-pane"')
        self.assertContains(hx, "Mona Saleh")
        self.assertNotContains(hx, f">{agent.pk}<")

    def test_back_to_back_preview_warns(self):
        department = make_department("Reason")
        manager = make_user("reason-lead", "support", department, make_role("Reason lead", can_manage_shifts=True))
        agent = make_user("reason-agent", "support", department, first_name="Mona", last_name="Saleh")
        night = make_shift_type(department, name="Reason night")
        first = timezone.localdate() + timedelta(days=7)
        self.client.force_login(manager)
        preview = self.client.post(reverse("shifts_calc_rotation"), {
            "department": department.pk,
            "start": first.isoformat(),
            "end": (first + timedelta(days=1)).isoformat(),
            "shift_type": night.pk,
            "users": [agent.pk],
            "action": "preview",
        })
        self.assertEqual(preview.status_code, 200)
        self.assertContains(preview, "1 night is next to another night for the same person.")
        self.assertContains(preview, "shifts-mini--warn", count=1)
    def test_missing_rotation_dates_ask_for_dates(self):
        department = make_department("Dates")
        manager = make_user("dates-lead", "support", department, make_role("Dates lead", can_manage_shifts=True))
        night = make_shift_type(department, name="Dates night")
        self.client.force_login(manager)
        response = self.client.post(reverse("shifts_calc_rotation"), {
            "department": department.pk,
            "start": "",
            "end": "",
            "shift_type": night.pk,
            "action": "preview",
        })
        self.assertContains(response, "Enter a start and an end date.")

class CalculatorTests(TestCase):
    def test_length_reports_after_midnight_separately(self):
        department = make_department("Calc")
        manager = make_user("calc-lead", "support", department, make_role("Calc lead", can_manage_shifts=True))
        self.client.force_login(manager)
        response = self.client.get(reverse("shifts_calc_length"), {"start": "22:00", "end": "06:00"})
        self.assertContains(response, "8")
        self.assertContains(response, "6")
        self.assertContains(response, "after midnight")

    def test_support_forbidden(self):
        department = make_department("Calc2")
        agent = make_user("calc-agent", "support", department)
        self.client.force_login(agent)
        self.assertEqual(self.client.get(reverse("shifts_calculator")).status_code, 403)

    def test_quarter_hour_length(self):
        department = make_department("Calc3")
        manager = make_user("calc-lead-3", "support", department, make_role("Calc lead 3", can_manage_shifts=True))
        self.client.force_login(manager)
        response = self.client.get(reverse("shifts_calc_length"), {"start": "09:00", "end": "12:30"})
        self.assertContains(response, "3 hours 30 minutes")
        self.assertContains(response, "0 hours 0 minutes after midnight")

class WorkerHoursTests(TestCase):
    def setUp(self):
        self.department = make_department("Mine")
        self.agent = make_user("mine-agent", "support", self.department)
        self.other = make_user("mine-other", "support", self.department)
        self.night = make_shift_type(self.department, name="Mine night")
        self.day = make_shift_type(self.department, name="Mine day", start=time(9, 0), end=time(17, 0))
        self.today = timezone.localdate()
        self.own = make_assignment(self.agent, self.night, self.today)
        self.client.force_login(self.agent)

    def test_mine_page_shows_a_colour_chip_not_the_hex(self):
        response = self.client.get(reverse("shifts_mine"))
        self.assertContains(response, 'class="shifts-chip')
        self.assertContains(response, "background-color: #336699")
        self.assertNotContains(response, ">#336699<")

    def test_time_ranges_keep_start_before_end_in_an_ltr_wrapper(self):
        response = self.client.get(reverse("shifts_mine"))
        self.assertContains(response, '<bdi class="shifts-ltr" dir="ltr">10:00 PM–6:00 AM</bdi>')
        self.assertNotContains(response, 'dir="auto"')
        self.assertNotContains(response, "13:00")

    def test_board_rota_and_mine_show_a_12_hour_clock(self):
        department = make_department("Clock")
        manager = make_user("clock-lead", "support", department, make_role("Clock lead", can_manage_shifts=True))
        agent = make_user("clock-agent", "support", department, first_name="Hana", last_name="Nabil")
        day = make_shift_type(department, name="Clock day", start=time(9, 0), end=time(13, 0))
        today = date(2026, 10, 6)
        row = make_assignment(agent, day, today)
        row.checked_in_at = aware(2026, 10, 6, 9, 30)
        row.save()
        during = aware(2026, 10, 6, 10, 0)
        self.client.force_login(manager)
        with _both_clocks(during, today):
            rota = self.client.get(reverse("shifts_rota"), {
                "department": department.pk,
                "view": "week",
                "start": "2026-10-06",
            })
        self.assertContains(rota, "1:00 PM")
        self.assertContains(rota, "9:00 AM")
        self.assertNotContains(rota, "13:00")
        self.client.force_login(agent)
        with _both_clocks(during, today):
            mine = self.client.get(reverse("shifts_mine"))
        self.assertContains(mine, "1:00 PM")
        self.assertNotContains(mine, "13:00")
        self.client.force_login(make_user("clock-branch", "branch"))
        with _both_clocks(during, today):
            board = self.client.get(reverse("shifts_available"))
        self.assertContains(board, "1:00 PM")
        self.assertNotContains(board, "13:00")

    @override_settings(TIME_ZONE="Africa/Cairo")
    def test_arabic_times_are_isolated_and_english_ranges_stay_left_to_right(self):
        department = make_department("Bidi")
        manager = make_user("bidi-lead", "support", department, make_role("Bidi lead", can_manage_shifts=True))
        agent = make_user("bidi-agent", "support", department, first_name="Sara", last_name="Nabil")
        day = make_shift_type(department, name="Bidi day", start=time(9, 0), end=time(17, 0))
        today = date(2026, 10, 6)
        row = make_assignment(agent, day, today)
        # 06:05 UTC is 9:05 AM in Africa/Cairo, independent of the server TIME_ZONE.
        row.checked_in_at = datetime(2026, 10, 6, 6, 5, tzinfo=dt_timezone.utc)
        row.save()
        params = {"department": department.pk, "view": "week", "start": "2026-10-06"}
        noon_utc = datetime(2026, 10, 6, 9, 0, tzinfo=dt_timezone.utc)
        self.client.force_login(manager)
        with _both_clocks(noon_utc, today):
            english = self.client.get(reverse("shifts_rota"), params)
        self.assertContains(english, '<bdi class="shifts-ltr" dir="ltr">9:00 AM–5:00 PM</bdi>')
        self.assertNotContains(english, 'dir="auto"')
        self.client.force_login(agent)
        with _both_clocks(noon_utc, today):
            mine_en = self.client.get(reverse("shifts_mine"))
        self.assertContains(mine_en, '<bdi class="shifts-ltr" dir="ltr">9:05 AM</bdi>')
        self.assertContains(mine_en, "Checked in at")
        self.client.force_login(manager)
        self.client.post(reverse("set_language"), {"language": "ar"})
        with _both_clocks(noon_utc, today):
            arabic = self.client.get(reverse("shifts_rota"), params)
        self.assertContains(arabic, '<bdi dir="auto">9:00 ص</bdi> – <bdi dir="auto">5:00 م</bdi>')
        self.assertNotContains(arabic, 'dir="ltr">9:00 ص')
        self.assertNotContains(arabic, 'class="shifts-ltr" dir="ltr">9:00 ص–5:00 م')
        self.client.force_login(agent)
        self.client.post(reverse("set_language"), {"language": "ar"})
        with _both_clocks(noon_utc, today):
            mine = self.client.get(reverse("shifts_mine"))
            board = self.client.get(reverse("shifts_available"))
        self.assertContains(mine, '<bdi dir="auto">9:00 ص</bdi> – <bdi dir="auto">5:00 م</bdi>')
        self.assertContains(mine, '<bdi dir="auto">9:05 ص</bdi>')
        self.assertContains(mine, "تم تسجيل الحضور في")
        self.assertNotContains(mine, 'dir="ltr">9:05')
        self.assertContains(board, '<bdi dir="auto">9:00 ص</bdi> – <bdi dir="auto">5:00 م</bdi>')
        self.assertNotContains(board, 'dir="ltr">5:00 م')

    def test_owner_can_set_same_evening_hours(self):
        response = self.client.post(reverse("shifts_mine_hours", args=[self.own.pk]), {
            "start_time": "18:00",
            "end_time": "23:00",
            "action": "save",
        })
        self.assertEqual(response.status_code, 200)
        self.own.refresh_from_db()
        self.assertEqual(self.own.start_time_override, time(18, 0))
        self.assertEqual(self.own.times_set_by_id, self.agent.pk)

    def test_other_user_and_day_shift_are_403(self):
        foreign = make_assignment(self.other, self.night, self.today)
        self.assertEqual(self.client.post(reverse("shifts_mine_hours", args=[foreign.pk]), {
            "start_time": "18:00", "end_time": "23:00", "action": "save",
        }).status_code, 403)
        foreign.refresh_from_db()
        self.assertIsNone(foreign.start_time_override)
        day_row = make_assignment(self.agent, self.day, self.today + timedelta(days=1))
        self.assertEqual(self.client.post(reverse("shifts_mine_hours", args=[day_row.pk]), {
            "start_time": "10:00", "end_time": "15:00", "action": "save",
        }).status_code, 403)

    def test_eight_days_ago_rejected_and_seven_accepted(self):
        old = make_assignment(self.agent, self.night, self.today - timedelta(days=8))
        response = self.client.post(reverse("shifts_mine_hours", args=[old.pk]), {
            "start_time": "18:00", "end_time": "23:00", "action": "save",
        })
        self.assertContains(response, "last 7 days")
        old.refresh_from_db()
        self.assertIsNone(old.start_time_override)
        recent = make_assignment(self.agent, self.night, self.today - timedelta(days=7))
        ok = self.client.post(reverse("shifts_mine_hours", args=[recent.pk]), {
            "start_time": "18:00", "end_time": "23:00", "action": "save",
        })
        self.assertEqual(ok.status_code, 200)
        recent.refresh_from_db()
        self.assertEqual(recent.start_time_override, time(18, 0))

    def test_branch_post_forbidden(self):
        branch = make_user("mine-branch", "branch")
        self.client.force_login(branch)
        self.assertEqual(self.client.post(reverse("shifts_mine_hours", args=[self.own.pk]), {
            "start_time": "18:00", "end_time": "23:00", "action": "save",
        }).status_code, 403)

    def test_length_cap_equal_times_overlap_and_reset(self):
        too_long = self.client.post(reverse("shifts_mine_hours", args=[self.own.pk]), {
            "start_time": "00:00", "end_time": "16:30", "action": "save",
        })
        self.assertContains(too_long, "16 hours")
        self.own.refresh_from_db()
        self.assertIsNone(self.own.start_time_override)
        equal = self.client.post(reverse("shifts_mine_hours", args=[self.own.pk]), {
            "start_time": "18:00", "end_time": "18:00", "action": "save",
        })
        self.assertContains(equal, "16 hours")
        crossing = self.client.post(reverse("shifts_mine_hours", args=[self.own.pk]), {
            "start_time": "18:00", "end_time": "02:00", "action": "save",
        })
        self.assertEqual(crossing.status_code, 200)
        self.own.refresh_from_db()
        self.assertEqual(self.own.end_time_override, time(2, 0))
        self.own.start_time_override = None
        self.own.end_time_override = None
        self.own.times_set_by = None
        self.own.times_set_at = None
        self.own.save()
        neighbor = make_assignment(self.agent, self.night, self.today + timedelta(days=1), time(1, 0), time(9, 0))
        blocked = self.client.post(reverse("shifts_mine_hours", args=[self.own.pk]), {
            "start_time": "18:00", "end_time": "02:00", "action": "save",
        })
        self.assertContains(blocked, "overlaps")
        self.own.refresh_from_db()
        self.assertIsNone(self.own.start_time_override)
        neighbor.delete()
        self.client.post(reverse("shifts_mine_hours", args=[self.own.pk]), {
            "start_time": "18:00", "end_time": "23:00", "action": "save",
        })
        reset = self.client.post(reverse("shifts_mine_hours", args=[self.own.pk]), {"action": "reset"})
        self.assertEqual(reset.status_code, 200)
        self.own.refresh_from_db()
        self.assertIsNone(self.own.start_time_override)
        self.assertIsNone(self.own.times_set_by_id)
        self.assertIsNone(self.own.times_set_at)


class CheckInTests(TestCase):
    def setUp(self):
        self.department = make_department("Check")
        self.agent = make_user(
            "check-agent",
            "support",
            self.department,
            make_role("Check agent", can_check_in=True),
        )
        self.other = make_user("check-other", "support", self.department)
        self.day = make_shift_type(self.department, name="Check day", start=time(9, 0), end=time(17, 0))
        self.night = make_shift_type(self.department, name="Check night")
        self.today = date(2026, 6, 2)
        self.row = make_assignment(self.agent, self.day, self.today)
        self.client.force_login(self.agent)

    def _at(self, hh, mm=0):
        return _both_clocks(aware(2026, 6, 2, hh, mm), date(2026, 6, 2))

    def test_window_and_idempotent_press(self):
        with self._at(8, 0):
            too_soon = self.client.post(reverse("shifts_check_in", args=[self.row.pk]))
        self.assertContains(too_soon, "30 minutes")
        self.row.refresh_from_db()
        self.assertIsNone(self.row.checked_in_at)
        with self._at(8, 40):
            first = self.client.post(reverse("shifts_check_in", args=[self.row.pk]))
        self.row.refresh_from_db()
        stamped = self.row.checked_in_at
        self.assertIsNotNone(stamped)
        with self._at(9, 0):
            second = self.client.post(reverse("shifts_check_in", args=[self.row.pk]))
        self.row.refresh_from_db()
        self.assertEqual(self.row.checked_in_at, stamped)
        self.assertContains(second, "Checked in at")

    def test_hidden_for_night_other_and_branch(self):
        night = make_assignment(self.agent, self.night, self.today + timedelta(days=1))
        with self._at(8, 40):
            self.assertEqual(self.client.post(reverse("shifts_check_in", args=[night.pk])).status_code, 403)
        # Freezing timezone.now rewrites the session expiry into the past
        # because SESSION_SAVE_EVERY_REQUEST is on. Sign in again before the
        # unpatched request.
        self.client.force_login(self.agent)
        foreign = make_assignment(self.other, self.day, self.today + timedelta(days=2))
        self.assertEqual(self.client.post(reverse("shifts_check_in", args=[foreign.pk])).status_code, 403)
        self.client.force_login(make_user("check-branch", "branch"))
        self.assertEqual(self.client.post(reverse("shifts_check_in", args=[self.row.pk])).status_code, 403)

    def test_manager_clear_worker_cannot(self):
        self.row.checked_in_at = aware(2026, 6, 2, 8, 40)
        self.row.save()
        self.assertEqual(self.client.post(reverse("shifts_cell"), {
            "clear_check_in": "1",
            "user": self.agent.pk,
            "date": "2026-06-02",
            "department": self.department.pk,
            "shift_type": self.day.pk,
        }).status_code, 403)
        manager = make_user("check-lead", "support", self.department, make_role("Check lead", can_manage_shifts=True))
        self.client.force_login(manager)
        self.client.post(reverse("shifts_cell"), {
            "clear_check_in": "1",
            "user": self.agent.pk,
            "date": "2026-06-02",
            "department": self.department.pk,
            "shift_type": self.day.pk,
        })
        self.row.refresh_from_db()
        self.assertIsNone(self.row.checked_in_at)

    def test_panel_only_for_the_owning_support_user(self):
        with patch("django.utils.timezone.now", return_value=aware(2026, 6, 2, 8, 40)), \
             patch("django.utils.timezone.localdate", return_value=date(2026, 6, 2)):
            mine = self.client.get(reverse("shifts_mine"))
        self.assertContains(mine, 'id="shifts-check-in"')
        self.client.force_login(self.other)
        with patch("django.utils.timezone.now", return_value=aware(2026, 6, 2, 8, 40)), \
             patch("django.utils.timezone.localdate", return_value=date(2026, 6, 2)):
            other = self.client.get(reverse("shifts_mine"))
        self.assertNotContains(other, 'id="shifts-check-in"')
        manager = make_user(
            "check-panel-lead",
            "support",
            self.department,
            make_role("Panel lead", can_manage_shifts=True, can_check_in=True),
        )
        make_assignment(manager, self.day, self.today)
        self.client.force_login(manager)
        with patch("django.utils.timezone.now", return_value=aware(2026, 6, 2, 8, 40)), \
             patch("django.utils.timezone.localdate", return_value=date(2026, 6, 2)):
            rota = self.client.get(reverse("shifts_rota"))
        self.assertContains(rota, 'id="shifts-check-in"')

    def test_day_to_night_clears_check_in(self):
        self.row.checked_in_at = aware(2026, 6, 2, 9, 0)
        self.row.save()
        manager = make_user("type-lead", "support", self.department, make_role("Type lead", can_manage_shifts=True))
        self.client.force_login(manager)
        self.client.post(reverse("shifts_cell"), {
            "user": self.agent.pk,
            "date": "2026-06-02",
            "department": self.department.pk,
            "assignment": self.row.pk,
            "shift_type": self.night.pk,
        })
        self.row.refresh_from_db()
        self.assertEqual(self.row.shift_type_id, self.night.pk)
        self.assertIsNone(self.row.checked_in_at)

    def test_day_shift_check_in_window_30_minutes(self):
        # 08:29 is 31 minutes before 09:00 - check-in panel should NOT be visible
        with self._at(8, 29):
            res_early = self.client.get(reverse("shifts_mine"))
            self.assertNotContains(res_early, f'id="shifts-check-in-{self.row.pk}"')
            post_early = self.client.post(reverse("shifts_check_in", args=[self.row.pk]))
            self.assertContains(post_early, "30 minutes")
        self.row.refresh_from_db()
        self.assertIsNone(self.row.checked_in_at)

        # 08:30 is exactly 30 minutes before 09:00 - check-in panel appears and POST succeeds
        with self._at(8, 30):
            res_open = self.client.get(reverse("shifts_mine"))
            self.assertContains(res_open, f'id="shifts-check-in-{self.row.pk}"')
            post_open = self.client.post(reverse("shifts_check_in", args=[self.row.pk]))
            self.assertContains(post_open, "Checked in at")
        self.row.refresh_from_db()
        self.assertIsNotNone(self.row.checked_in_at)

    def test_night_shift_cross_midnight_check_in_window(self):
        night_row = make_assignment(self.agent, self.night, self.today)
        # 21:29 is 31 minutes before 22:00 - check-in is early/hidden
        with self._at(21, 29):
            res_early = self.client.get(reverse("shifts_mine"))
            self.assertNotContains(res_early, f'id="shifts-check-in-{night_row.pk}"')
            post_early = self.client.post(reverse("shifts_check_in", args=[night_row.pk]))
            self.assertContains(post_early, "30 minutes")
        night_row.refresh_from_db()
        self.assertIsNone(night_row.checked_in_at)

        # 21:30 is 30 minutes before 22:00 - check-in opens
        with self._at(21, 30):
            res_open = self.client.get(reverse("shifts_mine"))
            self.assertContains(res_open, f'id="shifts-check-in-{night_row.pk}"')

        # Next day at 00:30 (cross-midnight, during the 22:00-06:00 shift)
        next_day = date(2026, 6, 3)
        with patch("django.utils.timezone.now", return_value=aware(2026, 6, 3, 0, 30)), \
             patch("django.utils.timezone.localdate", return_value=next_day):
            res_cross = self.client.get(reverse("shifts_mine"))
            self.assertContains(res_cross, f'id="shifts-check-in-{night_row.pk}"')
            post_cross = self.client.post(reverse("shifts_check_in", args=[night_row.pk]))
            self.assertContains(post_cross, "Checked in at")
        night_row.refresh_from_db()
        self.assertIsNotNone(night_row.checked_in_at)

    def test_check_in_panels_htmx_partial_endpoint(self):
        with self._at(8, 20):
            res = self.client.get(reverse("shifts_check_in_panels"))
            self.assertEqual(res.status_code, 200)
            self.assertNotContains(res, f'id="shifts-check-in-{self.row.pk}"')

        with self._at(8, 35):
            res = self.client.get(reverse("shifts_check_in_panels"))
            self.assertEqual(res.status_code, 200)
            self.assertContains(res, f'id="shifts-check-in-{self.row.pk}"')


class AutoCheckOutTests(TestCase):
    def setUp(self):
        self.department = make_department("Auto out")
        self.agent = make_user(
            "auto-agent",
            "support",
            self.department,
            make_role("Auto agent", can_check_in=True),
            first_name="Nada",
            last_name="Farid",
        )
        self.day = make_shift_type(self.department, name="Auto day", start=time(9, 0), end=time(17, 0))
        self.night = make_shift_type(self.department, name="Auto night")
        self.today = date(2026, 6, 2)
        self.row = make_assignment(self.agent, self.day, self.today)
        self.client.force_login(self.agent)

    def _stamp_checked_in(self, at):
        self.row.checked_in_at = at
        self.row.checked_out_at = None
        self.row.save(update_fields=["checked_in_at", "checked_out_at"])

    @override_settings(TIME_ZONE="Africa/Cairo")
    def test_lazy_clear_after_effective_end_in_cairo(self):
        checked_in = timezone.make_aware(datetime(2026, 6, 2, 9, 5))
        end = timezone.make_aware(datetime(2026, 6, 2, 17, 0))
        self._stamp_checked_in(checked_in)
        manager = make_user(
            "auto-lead",
            "support",
            self.department,
            make_role("Auto lead", can_manage_shifts=True),
        )
        self.client.force_login(manager)
        with _both_clocks(end, self.today):
            rota = self.client.get(reverse("shifts_rota"), {
                "department": self.department.pk,
                "view": "week",
                "start": "2026-05-30",
            })
        self.assertNotContains(rota, "Checked in at")
        self.client.force_login(self.agent)
        with _both_clocks(end, self.today):
            mine = self.client.get(reverse("shifts_mine"))
            board = self.client.get(reverse("shifts_available"))
        self.assertNotContains(mine, "Checked in at")
        # The ended shift stays on today's timeline as finished, not on shift.
        self.assertContains(board, "0 people on shift")
        self.assertContains(board, "shifts-tl-bar--done")
        self.assertNotContains(board, "shifts-tl-bar--on ")
        # Viewing settles the check-in in memory only; page loads never write.
        self.row.refresh_from_db()
        self.assertEqual(self.row.checked_in_at, checked_in)
        self.assertIsNone(self.row.checked_out_at)

    @override_settings(TIME_ZONE="Africa/Cairo")
    def test_override_end_is_the_auto_check_out_time(self):
        self.row.start_time_override = time(9, 0)
        self.row.end_time_override = time(15, 0)
        self.row.times_set_by = self.agent
        self.row.times_set_at = timezone.now()
        self.row.checked_in_at = timezone.make_aware(datetime(2026, 6, 2, 9, 5))
        self.row.save()
        still_open = timezone.make_aware(datetime(2026, 6, 2, 14, 59))
        with _both_clocks(still_open, self.today):
            self.client.get(reverse("shifts_mine"))
        self.row.refresh_from_db()
        self.assertIsNotNone(self.row.checked_in_at)
        ended = timezone.make_aware(datetime(2026, 6, 2, 15, 0))
        with _both_clocks(ended, self.today):
            mine = self.client.get(reverse("shifts_mine"))
        self.assertNotContains(mine, "Checked in at")
        self.row.refresh_from_db()
        self.assertIsNotNone(self.row.checked_in_at)
        self.assertIsNone(self.row.checked_out_at)

    @override_settings(TIME_ZONE="Africa/Cairo")
    def test_agent_cannot_check_in_again_after_auto_or_manager_clear(self):
        self._stamp_checked_in(timezone.make_aware(datetime(2026, 6, 2, 9, 5)))
        end = timezone.make_aware(datetime(2026, 6, 2, 17, 0))
        with _both_clocks(end, self.today):
            self.client.get(reverse("shifts_mine"))
            blocked = self.client.post(reverse("shifts_check_in", args=[self.row.pk]))
        self.row.refresh_from_db()
        self.assertIsNone(self.row.checked_in_at)
        self.assertIsNotNone(self.row.checked_out_at)
        self.assertNotContains(blocked, "Checked in at")
        self.row.start_time_override = time(9, 0)
        self.row.end_time_override = time(21, 0)
        self.row.times_set_by = self.agent
        self.row.times_set_at = timezone.now()
        self.row.save()
        later = timezone.make_aware(datetime(2026, 6, 2, 18, 0))
        with _both_clocks(later, self.today):
            self.client.post(reverse("shifts_check_in", args=[self.row.pk]))
        self.row.refresh_from_db()
        self.assertIsNone(self.row.checked_in_at)
        self.assertIsNotNone(self.row.checked_out_at)

        fresh = make_assignment(self.agent, self.day, self.today + timedelta(days=1))
        # Manager clear is on today's row; use a new open day shift that is today.
        fresh.delete()
        self.row.end_time_override = None
        self.row.start_time_override = None
        self.row.times_set_by = None
        self.row.times_set_at = None
        self.row.checked_out_at = None
        self.row.checked_in_at = timezone.make_aware(datetime(2026, 6, 2, 9, 5))
        self.row.save()
        manager = make_user(
            "auto-clear-lead",
            "support",
            self.department,
            make_role("Auto clear lead", can_manage_shifts=True),
        )
        during = timezone.make_aware(datetime(2026, 6, 2, 10, 0))
        self.client.force_login(manager)
        with _both_clocks(during, self.today):
            self.client.post(reverse("shifts_cell"), {
                "clear_check_in": "1",
                "user": self.agent.pk,
                "date": "2026-06-02",
                "department": self.department.pk,
                "shift_type": self.day.pk,
            })
        self.client.force_login(self.agent)
        with _both_clocks(during, self.today):
            again = self.client.post(reverse("shifts_check_in", args=[self.row.pk]))
        self.row.refresh_from_db()
        self.assertIsNone(self.row.checked_in_at)
        self.assertIsNotNone(self.row.checked_out_at)
        self.assertNotContains(again, "Checked in at")
        self.assertNotContains(again, "Check out")

    @override_settings(TIME_ZONE="Africa/Cairo")
    def test_manager_can_restore_check_in_without_the_agent_flag(self):
        self._stamp_checked_in(timezone.make_aware(datetime(2026, 6, 2, 9, 5)))
        manager = make_user(
            "auto-restore-lead",
            "support",
            self.department,
            make_role("Auto restore lead", can_manage_shifts=True),
        )
        during = timezone.make_aware(datetime(2026, 6, 2, 10, 0))
        self.client.force_login(manager)
        with _both_clocks(during, self.today):
            self.client.post(reverse("shifts_cell"), {
                "clear_check_in": "1",
                "user": self.agent.pk,
                "date": "2026-06-02",
                "department": self.department.pk,
                "shift_type": self.day.pk,
            })
            form = self.client.get(reverse("shifts_cell"), {
                "user": self.agent.pk,
                "date": "2026-06-02",
                "department": self.department.pk,
                "assignment": self.row.pk,
            }, HTTP_HX_REQUEST="true")
        self.assertContains(form, 'name="restore_check_in"')
        self.assertNotContains(form, "Clear check-in")
        with _both_clocks(during, self.today):
            self.client.post(reverse("shifts_cell"), {
                "restore_check_in": "1",
                "user": self.agent.pk,
                "date": "2026-06-02",
                "department": self.department.pk,
                "shift_type": self.day.pk,
            })
        self.row.refresh_from_db()
        self.assertIsNotNone(self.row.checked_in_at)
        self.assertIsNone(self.row.checked_out_at)
        self.client.force_login(self.agent)
        with _both_clocks(during, self.today):
            mine = self.client.get(reverse("shifts_mine"))
            board = self.client.get(reverse("shifts_available"))
        self.assertContains(mine, "Checked in at")
        self.assertContains(board, "Nada Farid")

    @override_settings(TIME_ZONE="Africa/Cairo")
    def test_night_shift_checks_out_when_it_ends(self):
        night = make_assignment(self.agent, self.night, date(2026, 6, 1))
        night.checked_in_at = timezone.make_aware(datetime(2026, 6, 1, 22, 5))
        night.save(update_fields=["checked_in_at"])
        later = timezone.make_aware(datetime(2026, 6, 2, 8, 0))
        with _both_clocks(later, self.today):
            self.client.get(reverse("shifts_available"))
            posted = self.client.post(reverse("shifts_check_in", args=[night.pk]))
        self.assertEqual(posted.status_code, 403)
        night.refresh_from_db()
        self.assertIsNone(night.checked_in_at)
        self.assertEqual(night.checked_out_at, timezone.make_aware(datetime(2026, 6, 2, 6, 0)))


    @override_settings(TIME_ZONE="Africa/Cairo")
    def test_viewing_pages_after_shift_end_does_not_write(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        self._stamp_checked_in(timezone.make_aware(datetime(2026, 6, 2, 9, 5)))
        after = timezone.make_aware(datetime(2026, 6, 2, 18, 0))
        with _both_clocks(after, self.today), CaptureQueriesContext(connection) as queries:
            mine = self.client.get(reverse("shifts_mine"))
            panels = self.client.get(reverse("shifts_check_in_panels"), HTTP_HX_REQUEST="true")
            team = self.client.get(reverse("shifts_team"))
        self.assertEqual([mine.status_code, panels.status_code, team.status_code], [200, 200, 200])
        self.assertNotContains(mine, "Checked in at")
        self.assertNotContains(team, "Checked in at")
        writes = [
            q["sql"] for q in queries.captured_queries
            if q["sql"].lstrip().upper().startswith("UPDATE") and "shifts_shiftassignment" in q["sql"]
        ]
        self.assertEqual(writes, [])


class CheckInPermissionTests(TestCase):
    def setUp(self):
        self.department = make_department("Check flag")
        self.day = make_shift_type(self.department, name="Flag day", start=time(9, 0), end=time(17, 0))
        self.today = date(2026, 6, 2)
        self.during = aware(2026, 6, 2, 10, 0)

    def test_flag_allows_check_in_and_hides_it_when_missing(self):
        allowed_user = make_user(
            "flag-yes",
            "support",
            self.department,
            make_role("Flag yes", can_check_in=True),
        )
        denied_user = make_user("flag-no", "support", self.department)
        allowed = make_assignment(allowed_user, self.day, self.today)
        denied = make_assignment(denied_user, self.day, self.today)
        self.client.force_login(denied_user)
        with _both_clocks(self.during, self.today):
            hidden = self.client.get(reverse("shifts_mine"))
            blocked = self.client.post(reverse("shifts_check_in", args=[denied.pk]))
        self.assertNotContains(hidden, 'id="shifts-check-in"')
        self.assertEqual(blocked.status_code, 403)
        denied.refresh_from_db()
        self.assertIsNone(denied.checked_in_at)
        self.client.force_login(allowed_user)
        with _both_clocks(self.during, self.today):
            shown = self.client.get(reverse("shifts_mine"))
            opened = self.client.post(reverse("shifts_check_in", args=[allowed.pk]))
        self.assertContains(shown, 'id="shifts-check-in"')
        self.assertContains(opened, "Checked in at")
        allowed.refresh_from_db()
        self.assertIsNotNone(allowed.checked_in_at)

    def test_manager_without_check_in_flag_can_still_clear_and_restore(self):
        agent = make_user(
            "flag-agent",
            "support",
            self.department,
            make_role("Flag agent", can_check_in=True),
        )
        row = make_assignment(agent, self.day, self.today)
        row.checked_in_at = self.during
        row.save(update_fields=["checked_in_at"])
        manager = make_user(
            "flag-lead",
            "support",
            self.department,
            make_role("Flag lead", can_manage_shifts=True),
        )
        make_assignment(manager, self.day, self.today)
        self.client.force_login(manager)
        with _both_clocks(self.during, self.today):
            own = self.client.get(reverse("shifts_mine"))
        self.assertNotContains(own, 'id="shifts-check-in"')
        with _both_clocks(self.during, self.today):
            self.client.post(reverse("shifts_cell"), {
                "clear_check_in": "1",
                "user": agent.pk,
                "date": "2026-06-02",
                "department": self.department.pk,
                "shift_type": self.day.pk,
            })
            restored = self.client.post(reverse("shifts_cell"), {
                "restore_check_in": "1",
                "user": agent.pk,
                "date": "2026-06-02",
                "department": self.department.pk,
                "shift_type": self.day.pk,
            })
        self.assertEqual(restored.status_code, 200)
        row.refresh_from_db()
        self.assertIsNotNone(row.checked_in_at)
        self.assertIsNone(row.checked_out_at)


class TeamRotaTests(TestCase):
    def test_manager_opens_on_a_department_that_has_types(self):
        make_department("Aaa empty team")
        filled = make_department("Zzz filled team")
        make_shift_type(filled, name="Team night")
        role = make_role("Team default", can_manage_shifts=True)
        manager = make_user("team-default", "support", filled, role)
        self.client.force_login(manager)
        response = self.client.get(reverse("shifts_team"))
        self.assertContains(response, f'value="{filled.pk}" selected')
        self.assertNotIn(">Show<", response.content.decode())

    def test_support_cannot_switch_department(self):
        home = make_department("Home team")
        other = make_department("Other team")
        agent = make_user("team-agent", "support", home)
        self.client.force_login(agent)
        response = self.client.get(reverse("shifts_team"), {"department": other.pk, "start": "2026-10-03"})
        self.assertEqual(response.status_code, 403)

    def test_branch_forbidden(self):
        self.client.force_login(make_user("team-branch", "branch"))
        self.assertEqual(self.client.get(reverse("shifts_team")).status_code, 403)


class AvailableNowTests(TestCase):
    """Available now draws today's shifts as a timeline, one bar per shift."""

    def _states(self, response, name):
        """Bar states on the row for ``name``: on, late, later, done or missed."""
        import re

        content = response.content.decode()
        row = re.search(r'data-filter-item="%s">([\s\S]*?)</div>\s*</div>' % re.escape(name), content)
        self.assertIsNotNone(row, f"no timeline row for {name}")
        return re.findall(r"shifts-tl-bar--(on|late|later|done|missed)\b", row.group(1))

    def _board(self, moment, **headers):
        with patch("django.utils.timezone.now", return_value=moment),              patch("django.utils.timezone.localdate", return_value=timezone.localtime(moment).date()):
            return self.client.get(reverse("shifts_available_board"), **headers)

    def test_night_shift_needs_check_in_until_end(self):
        department = make_department("Now")
        agent = make_user("now-agent", "support", department, make_role("Now agent", can_check_in=True))
        night = make_shift_type(department, name="Now night")
        row = make_assignment(agent, night, date(2026, 6, 1))
        self.client.force_login(make_user("now-branch", "branch"))
        self.assertEqual(self._states(self._board(aware(2026, 6, 2, 5, 0)), "now-agent"), ["late"])
        row.checked_in_at = aware(2026, 6, 1, 22, 10)
        row.save(update_fields=["checked_in_at"])
        shown = self._board(aware(2026, 6, 2, 5, 0))
        self.assertEqual(self._states(shown, "now-agent"), ["on"])
        self.assertContains(shown, "1 person on shift")
        self.assertNotContains(shown, "Now night")
        # At 6 AM the night is over and the board moves on to the day ahead.
        self.assertNotContains(self._board(aware(2026, 6, 2, 6, 0)), "now-agent")

    def test_day_shift_needs_check_in_and_etag_changes(self):
        department = make_department("Now day")
        agent = make_user("day-agent", "support", department, make_role("Day agent", can_check_in=True))
        day = make_shift_type(department, name="Now day", start=time(9, 0), end=time(17, 0))
        row = make_assignment(agent, day, date(2026, 6, 2))
        self.client.force_login(agent)
        noon = aware(2026, 6, 2, 12, 0)
        before = self._board(noon, HTTP_HX_REQUEST="true")
        self.assertEqual(self._states(before, "day-agent"), ["late"])
        self.assertContains(before, "0 people on shift")
        row.checked_in_at = noon
        row.save()
        after = self._board(noon, HTTP_HX_REQUEST="true")
        self.assertEqual(self._states(after, "day-agent"), ["on"])
        self.assertNotEqual(after["ETag"], before["ETag"])
        cached = self._board(noon, HTTP_HX_REQUEST="true", HTTP_IF_NONE_MATCH=after["ETag"])
        self.assertEqual(cached.status_code, 304)

    def test_early_check_in_stays_upcoming_until_start(self):
        department = make_department("Early")
        agent = make_user("early-agent", "support", department, make_role("Early agent", can_check_in=True))
        day = make_shift_type(department, name="Early day", start=time(9, 0), end=time(17, 0))
        row = make_assignment(agent, day, date(2026, 6, 2))
        self.client.force_login(agent)
        row.checked_in_at = aware(2026, 6, 2, 8, 40)
        row.save()
        during = self._board(aware(2026, 6, 2, 8, 40), HTTP_HX_REQUEST="true")
        self.assertEqual(self._states(during, "early-agent"), ["later"])
        self.assertContains(during, "0 people on shift")
        started = self._board(aware(2026, 6, 2, 9, 0), HTTP_HX_REQUEST="true")
        self.assertEqual(self._states(started, "early-agent"), ["on"])
        self.assertNotEqual(started["ETag"], during["ETag"])
        self.assertEqual(self._states(self._board(aware(2026, 6, 2, 17, 0)), "early-agent"), ["done"])

    def test_window_runs_from_9_am_to_6_am_and_clips_at_the_edges(self):
        department = make_department("Bars")
        night_agent = make_user("bar-night", "support", department)
        early_agent = make_user("bar-early", "support", department)
        late_agent = make_user("bar-late", "support", department)
        night = make_shift_type(department, name="Bar night")  # 10 PM to 6 AM
        long_night = make_shift_type(department, name="Bar long", start=time(22, 0), end=time(10, 0))
        late = make_shift_type(department, name="Bar late", start=time(23, 0), end=time(8, 0))
        make_assignment(night_agent, night, date(2026, 6, 1))  # ends 6 AM, before the window
        make_assignment(night_agent, night, date(2026, 6, 2))
        make_assignment(early_agent, long_night, date(2026, 6, 1))  # runs into the window
        make_assignment(late_agent, late, date(2026, 6, 2))
        self.client.force_login(night_agent)
        self.client.cookies[settings.LANGUAGE_COOKIE_NAME] = "ar"
        content = self._board(aware(2026, 6, 2, 12, 0), HTTP_ACCEPT_LANGUAGE="ar").content.decode()
        # Positions stay plain numbers in Arabic ("4.762", not a localized "4٫762").
        self.assertIn("مكتملة", content)  # rendered in Arabic
        # 21 hours, 9 AM to 6 AM: last night's 10 PM-10 AM is cut at 9 AM (1 h
        # shown); 10 PM-6 AM fills the end; 11 PM-8 AM runs past 6 AM.
        self.assertIn("inset-inline-start: 0.0%; inline-size: 4.762%;", content)
        self.assertIn("inset-inline-start: 61.905%; inline-size: 38.095%;", content)
        self.assertIn("inset-inline-start: 66.667%; inline-size: 33.333%;", content)
        self.assertEqual(content.count("shifts-tl-bar--from-yesterday"), 1)
        self.assertEqual(content.count("shifts-tl-bar--to-tomorrow"), 1)
        self.assertEqual(self._states(self._board(aware(2026, 6, 2, 12, 0)), "bar-night"), ["later"])
        self.assertIn("--tl-now: 14.286%", content)  # noon is 3 of 21 hours in
        self.assertIn("--tl-step: 14.2857%", content)
        self.assertIn('style="inset-inline-start: 0.0%"><span class="shifts-tl-tick-full">9:00', content)
        self.assertNotIn("data-now-outside", content)

    def test_a_shift_before_9_am_moves_the_start_earlier(self):
        department = make_department("Early bird")
        agent = make_user("bird-agent", "support", department)
        early = make_shift_type(department, name="Bird early", start=time(8, 0), end=time(16, 0))
        make_assignment(agent, early, date(2026, 6, 2))
        self.client.force_login(agent)
        content = self._board(aware(2026, 6, 2, 12, 0)).content.decode()
        # 8 AM to 6 AM is 22 hours; the 8 AM shift is drawn whole, not cut at 9.
        self.assertIn("inset-inline-start: 0.0%; inline-size: 36.364%;", content)
        self.assertNotIn("shifts-tl-bar--from-yesterday", content)
        self.assertIn('style="inset-inline-start: 0.0%"><span class="shifts-tl-tick-full">8:00', content)
        self.assertIn("--tl-now: 18.182%", content)  # noon is 4 of 22 hours in

    def test_six_to_nine_am_shows_the_day_about_to_start(self):
        department = make_department("Gap")
        agent = make_user("gap-agent", "support", department)
        day = make_shift_type(department, name="Gap day", start=time(9, 0), end=time(17, 0))
        make_assignment(agent, day, date(2026, 6, 2))
        self.client.force_login(agent)
        response = self._board(aware(2026, 6, 2, 7, 30))
        self.assertEqual(self._states(response, "gap-agent"), ["later"])
        self.assertContains(response, "data-now-outside")
        self.assertContains(response, "--tl-now: 0.0%")

    def test_empty_department_and_no_type_name_for_branch(self):
        empty = make_department("Empty desk")
        staffed = make_department("Staffed")
        agent = make_user("staffed-agent", "support", staffed, first_name="", last_name="")
        night = make_shift_type(staffed, name="Secret type")
        row = make_assignment(agent, night, date(2026, 6, 1))
        row.checked_in_at = aware(2026, 6, 1, 22, 10)
        row.save(update_fields=["checked_in_at"])
        self.client.force_login(make_user("board-branch", "branch"))
        response = self._board(aware(2026, 6, 2, 5, 0))
        self.assertContains(response, "Nobody scheduled today:")
        self.assertEqual(self._states(response, "staffed-agent"), ["on"])
        self.assertNotContains(response, "Secret type")
        self.assertNotContains(response, "--shift-colour")
        self.assertContains(response, "6:00 AM")
        self.assertContains(response, "1 person on shift")
        # Support staff see the shift type; staffed departments come before empty ones.
        self.client.force_login(make_user("board-support", "support", empty))
        support_view = self._board(aware(2026, 6, 2, 5, 0))
        self.assertContains(support_view, "Secret type")
        self.assertNotEqual(support_view["ETag"], response["ETag"])
        content = support_view.content.decode()
        self.assertLess(content.index("Staffed"), content.index("Empty desk"))

    def test_history_restore_is_not_304(self):
        department = make_department("Hist")
        self.client.force_login(make_user("hist-branch", "branch"))
        with patch("django.utils.timezone.now", return_value=aware(2026, 6, 2, 5, 0)), \
             patch("django.utils.timezone.localdate", return_value=date(2026, 6, 2)):
            first = self.client.get(reverse("shifts_available_board"), HTTP_HX_REQUEST="true")
            restored = self.client.get(
                reverse("shifts_available_board"),
                HTTP_HX_REQUEST="true",
                HTTP_IF_NONE_MATCH=first["ETag"],
                HTTP_HX_HISTORY_RESTORE_REQUEST="true",
            )
        self.assertEqual(restored.status_code, 200)

    def test_board_query_budget(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        department = make_department("Budget board")
        night = make_shift_type(department, name="Budget night")
        for index in range(6):
            person = make_user(f"board-{index}", "support", department)
            make_assignment(person, night, date(2026, 6, 1))
        viewer = make_user("budget-branch", "branch")
        # Steady state: the one-time save of a new user's language isn't counted.
        viewer.language = "en"
        viewer.save(update_fields=["language"])
        self.client.force_login(viewer)
        with patch("django.utils.timezone.now", return_value=aware(2026, 6, 2, 5, 0)), \
             patch("django.utils.timezone.localdate", return_value=date(2026, 6, 2)), \
             CaptureQueriesContext(connection) as captured:
            self.client.get(reverse("shifts_available_board"))
        self.assertLessEqual(len(captured), 8)

    def test_first_page_lists_worker_without_waiting_for_poll(self):
        department = make_department("First paint")
        agent = make_user("paint-agent", "support", department, first_name="Paint", last_name="Agent")
        night = make_shift_type(department, name="Paint night")
        row = make_assignment(agent, night, date(2026, 6, 1))
        row.checked_in_at = aware(2026, 6, 1, 22, 10)
        row.save(update_fields=["checked_in_at"])
        self.client.force_login(make_user("paint-branch", "branch"))
        with patch("django.utils.timezone.now", return_value=aware(2026, 6, 2, 5, 0)), \
             patch("django.utils.timezone.localdate", return_value=date(2026, 6, 2)):
            response = self.client.get(reverse("shifts_available"))
        self.assertContains(response, "Paint Agent")
        self.assertContains(response, 'hx-trigger="every 20s [!document.hidden], refresh"')
        self.assertContains(response, "data-no-progress")


class PermissionMatrixTests(TestCase):
    def setUp(self):
        self.department = make_department("Perm")
        self.other = make_department("Perm other")
        self.support = make_user(
            "perm-agent",
            "support",
            self.department,
            make_role("Perm agent", can_check_in=True),
        )
        self.outsider = make_user("perm-out", "support", self.other)
        self.branch = make_user("perm-branch", "branch")
        self.manager = make_user("perm-lead", "support", self.department, make_role("Perm lead", can_manage_shifts=True))
        self.superuser = make_user("perm-root", "support", self.department, superuser=True)
        self.night = make_shift_type(self.department, name="Perm night")
        self.day = make_shift_type(self.department, name="Perm day", start=time(0, 0), end=time(23, 59))
        self.row = make_assignment(self.support, self.night, date(2026, 1, 10))
        self.today_row = make_assignment(self.support, self.day, timezone.localdate())

    def _assert_login(self, response):
        self.assertEqual(response.status_code, 302)
        self.assertIn("/accounts/login/", response["Location"])

    def test_anonymous_redirects_every_named_url(self):
        gets = [
            reverse("shifts_home"),
            reverse("shifts_rota"),
            reverse("shifts_rota_grid"),
            reverse("shifts_cell"),
            reverse("shifts_types"),
            reverse("shifts_type_add"),
            reverse("shifts_type_edit", args=[self.night.pk]),
            reverse("shifts_calculator"),
            reverse("shifts_calc_length"),
            reverse("shifts_calc_work"),
            reverse("shifts_calc_work_export"),
            reverse("shifts_available"),
            reverse("shifts_available_board"),
            reverse("shifts_mine"),
            reverse("shifts_team"),
        ]
        for url in gets:
            self._assert_login(self.client.get(url))
        posts = [
            reverse("shifts_cell"),
            reverse("shifts_copy_week"),
            reverse("shifts_repeat_week"),
            reverse("shifts_type_archive", args=[self.night.pk]),
            reverse("shifts_calc_rotation"),
            reverse("shifts_mine_hours", args=[self.row.pk]),
            reverse("shifts_check_in", args=[self.today_row.pk]),
        ]
        for url in posts:
            self._assert_login(self.client.post(url, {}))

    def test_branch_support_and_manager_statuses(self):
        manager_gets = [
            "shifts_rota", "shifts_rota_grid", "shifts_types", "shifts_type_add",
            "shifts_calculator", "shifts_calc_length", "shifts_calc_work",
        ]
        shared_gets = ["shifts_available", "shifts_available_board"]
        support_gets = ["shifts_mine", "shifts_team"]
        for user, forbidden in (
            (self.branch, manager_gets + support_gets),
            (self.support, manager_gets),
        ):
            self.client.force_login(user)
            home = self.client.get(reverse("shifts_home"))
            self.assertEqual(home.status_code, 302, user.username)
            for name in shared_gets:
                self.assertEqual(self.client.get(reverse(name)).status_code, 200, name)
            for name in forbidden:
                response = self.client.get(reverse(name))
                self.assertEqual(response.status_code, 403, f"{user.username} {name}")
                self.assertNotContains(response, "2099-01-01", status_code=403)
        self.client.force_login(self.support)
        self.assertEqual(self.client.get(reverse("shifts_mine")).status_code, 200)
        self.assertEqual(self.client.get(reverse("shifts_team")).status_code, 200)
        self.assertEqual(self.client.get(reverse("shifts_team"), {"department": self.other.pk}).status_code, 403)
        self.assertEqual(self.client.get(reverse("shifts_type_edit", args=[self.night.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse("shifts_cell"), {"user": self.support.pk, "date": "2026-10-06", "department": self.department.pk}).status_code, 403)
        self.assertEqual(self.client.post(reverse("shifts_mine_hours", args=[self.row.pk]), {"action": "reset"}).status_code, 200)
        self.assertEqual(self.client.post(reverse("shifts_mine_hours", args=[make_assignment(self.outsider, self.night, date(2026, 10, 6)).pk]), {"action": "reset"}).status_code, 403)
        self.assertEqual(self.client.post(reverse("shifts_check_in", args=[self.today_row.pk])).status_code, 200)
        for name, payload in (
            ("shifts_cell", {"user": self.support.pk, "date": "2026-10-06", "department": self.department.pk}),
            ("shifts_copy_week", {"department": self.department.pk, "start": "2026-10-03"}),
            ("shifts_repeat_week", {"department": self.department.pk, "start": "2026-10-03", "weeks": "1"}),
            ("shifts_calc_rotation", {"action": "apply", "department": self.department.pk, "shift_type": self.night.pk, "start": "2026-10-05", "end": "2026-10-05"}),
        ):
            self.assertEqual(self.client.post(reverse(name), payload).status_code, 403, name)
        self.assertEqual(self.client.post(reverse("shifts_type_archive", args=[self.night.pk])).status_code, 403)
        self.client.force_login(self.branch)
        self.assertEqual(self.client.post(reverse("shifts_mine_hours", args=[self.row.pk]), {"action": "reset"}).status_code, 403)
        self.assertEqual(self.client.post(reverse("shifts_check_in", args=[self.today_row.pk])).status_code, 403)
        for user in (self.manager, self.superuser):
            self.client.force_login(user)
            self.assertEqual(self.client.get(reverse("shifts_home")).status_code, 302)
            for name in manager_gets + shared_gets + ["shifts_team"]:
                self.assertEqual(self.client.get(reverse(name)).status_code, 200, f"{user.username} {name}")
            self.assertEqual(self.client.get(reverse("shifts_mine")).status_code, 200, user.username)
            self.assertEqual(self.client.get(reverse("shifts_type_edit", args=[self.night.pk])).status_code, 200)
            self.assertEqual(self.client.get(reverse("shifts_cell"), {"user": self.support.pk, "date": "2026-10-06", "department": self.department.pk}, HTTP_HX_REQUEST="true").status_code, 200)
            self.assertEqual(self.client.post(reverse("shifts_cell"), {"user": self.support.pk, "date": "2026-10-07", "department": self.department.pk, "shift_type": self.night.pk}).status_code, 200)
            self.assertEqual(self.client.post(reverse("shifts_copy_week"), {"department": self.department.pk, "start": "2026-10-03"}).status_code, 302)
            self.assertEqual(self.client.post(reverse("shifts_repeat_week"), {"department": self.department.pk, "start": "2026-10-03", "weeks": "13"}).status_code, 302)
            self.assertEqual(self.client.post(reverse("shifts_calc_rotation"), {"action": "preview", "department": self.department.pk, "shift_type": self.night.pk, "start": "2026-10-05", "end": "2026-10-05", "users": [self.support.pk]}).status_code, 200)
            self.assertEqual(self.client.post(reverse("shifts_mine_hours", args=[self.row.pk]), {"action": "reset"}).status_code, 403)
            self.assertEqual(self.client.post(reverse("shifts_check_in", args=[self.today_row.pk])).status_code, 403)


class SameDayShiftTests(TestCase):
    def setUp(self):
        self.department = make_department("Double")
        self.manager = make_user(
            "double-lead",
            "support",
            self.department,
            make_role("Double lead", can_manage_shifts=True),
        )
        self.agent = make_user(
            "double-agent",
            "support",
            self.department,
            make_role("Double agent", can_check_in=True),
        )
        self.day = make_shift_type(self.department, name="Day desk", start=time(8, 0), end=time(16, 0))
        self.night = make_shift_type(self.department, name="Night desk", start=time(22, 0), end=time(6, 0))
        self.client.force_login(self.manager)

    def test_two_shifts_stack_and_overlap_is_rejected(self):
        on = date(2026, 10, 6)
        self.client.post(reverse("shifts_cell"), {
            "user": self.agent.pk,
            "date": "2026-10-06",
            "department": self.department.pk,
            "shift_type": self.day.pk,
        })
        added = self.client.post(reverse("shifts_cell"), {
            "user": self.agent.pk,
            "date": "2026-10-06",
            "department": self.department.pk,
            "shift_type": self.night.pk,
        })
        self.assertEqual(added.status_code, 200)
        self.assertEqual(ShiftAssignment.objects.filter(user=self.agent, date=on).count(), 2)
        grid = self.client.get(reverse("shifts_rota_grid"), {
            "department": self.department.pk,
            "view": "week",
            "start": "2026-10-03",
        })
        self.assertContains(grid, "Day desk")
        self.assertContains(grid, "Night desk")
        self.assertContains(grid, 'data-count="2"')
        self.assertContains(grid, "shifts-cell--filled", count=2)
        clash = self.client.post(reverse("shifts_cell"), {
            "user": self.agent.pk,
            "date": "2026-10-06",
            "department": self.department.pk,
            "shift_type": make_shift_type(self.department, name="Late morning", start=time(9, 0), end=time(12, 0)).pk,
        })
        self.assertContains(clash, "overlaps")
        self.assertEqual(ShiftAssignment.objects.filter(user=self.agent, date=on).count(), 2)
        day_row = ShiftAssignment.objects.get(user=self.agent, date=on, shift_type=self.day)
        removed = self.client.post(reverse("shifts_cell"), {
            "user": self.agent.pk,
            "date": "2026-10-06",
            "department": self.department.pk,
            "assignment": day_row.pk,
            "shift_type": "",
        })
        self.assertEqual(removed.status_code, 200)
        self.assertFalse(ShiftAssignment.objects.filter(pk=day_row.pk).exists())
        self.assertTrue(ShiftAssignment.objects.filter(user=self.agent, date=on, shift_type=self.night).exists())

    def test_copy_keeps_both_shifts(self):
        on = date(2026, 10, 3)
        make_assignment(self.agent, self.day, on)
        make_assignment(self.agent, self.night, on)
        self.client.post(reverse("shifts_copy_week"), {
            "department": self.department.pk,
            "start": "2026-10-10",
        })
        copied = ShiftAssignment.objects.filter(user=self.agent, date=date(2026, 10, 10))
        self.assertEqual(set(copied.values_list("shift_type_id", flat=True)), {self.day.pk, self.night.pk})

    def test_each_shift_checks_in_on_its_own(self):
        on = date(2026, 6, 2)
        day_row = make_assignment(self.agent, self.day, on)
        night_row = make_assignment(self.agent, self.night, on)
        self.client.force_login(self.agent)
        with _both_clocks(aware(2026, 6, 2, 9, 0), on):
            mine = self.client.get(reverse("shifts_mine"))
            day_in = self.client.post(reverse("shifts_check_in", args=[day_row.pk]))
        self.assertContains(mine, f'id="shifts-check-in-{day_row.pk}"')
        self.assertNotContains(mine, f'id="shifts-check-in-{night_row.pk}"')
        self.assertContains(day_in, "Checked in at")
        day_row.refresh_from_db()
        night_row.refresh_from_db()
        self.assertIsNotNone(day_row.checked_in_at)
        self.assertIsNone(night_row.checked_in_at)
        with _both_clocks(aware(2026, 6, 2, 22, 5), on):
            night_page = self.client.get(reverse("shifts_mine"))
            night_in = self.client.post(reverse("shifts_check_in", args=[night_row.pk]))
        self.assertContains(night_page, f'id="shifts-check-in-{night_row.pk}"')
        self.assertNotContains(night_page, f'id="shifts-check-in-{day_row.pk}"')
        self.assertContains(night_in, "Checked in at")
        day_row.refresh_from_db()
        night_row.refresh_from_db()
        self.assertIsNone(day_row.checked_in_at)
        self.assertIsNotNone(day_row.checked_out_at)
        self.assertIsNotNone(night_row.checked_in_at)
        self.assertIsNone(night_row.checked_out_at)


class ReviewFixTests(TestCase):
    """Regressions from the shifts review: bad input, tooltips, sticky forms, tabs."""

    def setUp(self):
        self.department = make_department("Review")
        self.manager = make_user(
            "review-lead",
            "support",
            self.department,
            make_role("Review lead", can_manage_shifts=True),
        )
        self.agent = make_user(
            "review-agent",
            "support",
            self.department,
            make_role("Review agent", can_check_in=True),
        )
        self.day = make_shift_type(self.department, name="Review day", start=time(8, 0), end=time(16, 0))
        self.night = make_shift_type(self.department, name="Review night", start=time(22, 0), end=time(6, 0))
        self.on = date(2026, 6, 2)
        self.client.force_login(self.manager)

    def test_bad_input_never_500s(self):
        cell = reverse("shifts_cell")
        base = {"department": self.department.pk, "user": self.agent.pk}
        for url, params in [
            (reverse("shifts_rota"), {"start": "2026-02-30"}),
            (reverse("shifts_rota"), {"department": "abc"}),
            (reverse("shifts_team"), {"start": "bad"}),
            (reverse("shifts_rota_grid"), {"start": "nope"}),
            (reverse("shifts_calc_work"), {"person": "x", "start": "x", "end": "y"}),
            (reverse("shifts_calc_work"), {"person": self.agent.pk, "start": "2026-02-30"}),
            (cell, {**base}),
            (cell, {**base, "date": "garbage"}),
            (cell, {**base, "date": "2026-06-02", "assignment": "zz"}),
            (cell, {"department": self.department.pk, "user": "zz", "date": "2026-06-02"}),
        ]:
            response = self.client.get(url, params)
            self.assertIn(response.status_code, (200, 404), f"{url} {params}")
        posted = self.client.post(cell, {
            **base, "date": "2026-06-02", "shift_type": self.day.pk,
            "start_time_override": "99:00", "end_time_override": "10:00",
        })
        self.assertEqual(posted.status_code, 200)
        self.assertContains(posted, "Enter a valid time.")
        self.assertFalse(ShiftAssignment.objects.filter(user=self.agent).exists())
        copied = self.client.post(reverse("shifts_copy_week"), {"department": self.department.pk, "start": "bad"})
        self.assertIn(copied.status_code, (200, 302))

    def test_bad_hours_shows_error_and_keeps_typed_values(self):
        row = make_assignment(self.agent, self.night, self.on)
        self.client.force_login(self.agent)
        with _both_clocks(aware(2026, 6, 2, 12, 0), self.on):
            bad = self.client.post(
                reverse("shifts_mine_hours", args=[row.pk]),
                {"action": "save", "start_time": "25:99", "end_time": "05:00"},
                HTTP_HX_REQUEST="true",
            )
            long = self.client.post(
                reverse("shifts_mine_hours", args=[row.pk]),
                {"action": "save", "start_time": "06:00", "end_time": "23:30"},
                HTTP_HX_REQUEST="true",
            )
            saved = self.client.post(
                reverse("shifts_mine_hours", args=[row.pk]),
                {"action": "save", "start_time": "21:00", "end_time": "05:00"},
                HTTP_HX_REQUEST="true",
            )
        self.assertContains(bad, "Enter a valid time.")
        self.assertContains(bad, 'value="25:99"')
        self.assertContains(long, 'value="06:00"')
        self.assertContains(long, 'value="23:30"')
        self.assertTemplateUsed(long, "shifts/mine_hours_form.html")
        self.assertNotIn("HX-Trigger", long)
        self.assertEqual(saved["HX-Trigger"], "closeModal, shiftsChanged")
        self.assertEqual(saved["HX-Retarget"], f"#shift-mine-{row.pk}")
        self.assertTemplateUsed(saved, "shifts/mine_row.html")
        row.refresh_from_db()
        self.assertEqual(row.start_time_override, time(21, 0))

    def test_overlap_error_keeps_chosen_type_and_times(self):
        make_assignment(self.agent, self.day, self.on)
        clash = make_shift_type(self.department, name="Review clash", start=time(9, 0), end=time(12, 0))
        response = self.client.post(reverse("shifts_cell"), {
            "department": self.department.pk,
            "user": self.agent.pk,
            "date": "2026-06-02",
            "shift_type": clash.pk,
            "start_time_override": "09:30",
            "end_time_override": "11:00",
        })
        self.assertContains(response, "overlaps")
        self.assertContains(response, f'<option value="{clash.pk}" selected>')
        self.assertContains(response, 'value="09:30"')
        self.assertContains(response, 'value="11:00"')

    def test_checked_in_tooltip_is_plain_text(self):
        row = make_assignment(self.agent, self.day, self.on)
        row.checked_in_at = aware(2026, 6, 2, 8, 5)
        row.save()
        with _both_clocks(aware(2026, 6, 2, 9, 0), self.on):
            grid = self.client.get(reverse("shifts_rota"), {
                "department": self.department.pk, "start": "2026-06-02",
            })
        html = grid.content.decode()
        self.assertIn('title="Checked in at 8:05 AM"', html)
        self.assertNotIn('title="Checked in at <bdi', html)

    def test_check_in_swaps_whole_list_for_htmx(self):
        row = make_assignment(self.agent, self.day, self.on)
        self.client.force_login(self.agent)
        with _both_clocks(aware(2026, 6, 2, 8, 40), self.on):
            response = self.client.post(reverse("shifts_check_in", args=[row.pk]), HTTP_HX_REQUEST="true")
        self.assertEqual(response["HX-Retarget"], "#shifts-check-in")
        self.assertContains(response, 'id="shifts-check-in"')
        self.assertContains(response, "Checked in at")
        self.assertContains(response, "You are on shift until")
        self.assertEqual(response.content.decode().count('type="submit"'), 0)

    def test_check_in_list_hidden_when_nothing_to_show(self):
        make_assignment(self.agent, self.day, self.on)
        self.client.force_login(self.agent)
        with _both_clocks(aware(2026, 6, 2, 6, 0), self.on):
            mine = self.client.get(reverse("shifts_mine"))
        self.assertContains(mine, 'class="shifts-checkin-list" data-sig="" hidden')

class MyShiftsLayoutTests(TestCase):
    def setUp(self):
        self.department = make_department("Layout")
        self.agent = make_user(
            "layout-agent",
            "support",
            self.department,
            make_role("Layout agent", can_check_in=True),
        )
        self.day = make_shift_type(self.department, name="Layout day", start=time(9, 0), end=time(17, 0))
        self.night = make_shift_type(self.department, name="Layout night", start=time(22, 0), end=time(6, 0))
        # Tuesday; the rota week runs Saturday 2026-05-30 to Friday 2026-06-05.
        self.today = date(2026, 6, 2)
        self.client.force_login(self.agent)

    def _get(self, hh, mm=0):
        with _both_clocks(aware(2026, 6, 2, hh, mm), self.today):
            return self.client.get(reverse("shifts_mine"))

    def test_groups_this_week_later_and_recent(self):
        make_assignment(self.agent, self.day, self.today)
        make_assignment(self.agent, self.day, date(2026, 6, 10))
        make_assignment(self.agent, self.night, date(2026, 5, 29))
        response = self._get(7)
        self.assertEqual([row.date for row in response.context["this_week"]], [self.today])
        self.assertEqual([row.date for row in response.context["later"]], [date(2026, 6, 10)])
        self.assertEqual([row.date for row in response.context["recent"]], [date(2026, 5, 29)])
        self.assertContains(response, "This week")
        self.assertContains(response, "Later")
        self.assertContains(response, '<details class="panel shifts-mine-section shifts-recent">')

    def test_running_night_from_yesterday_stays_in_this_week(self):
        make_assignment(self.agent, self.night, date(2026, 6, 1))
        response = self._get(2)
        self.assertEqual([row.date for row in response.context["this_week"]], [date(2026, 6, 1)])
        self.assertEqual(response.context["recent"], [])
        self.assertContains(response, "Yesterday")
        self.assertContains(response, "Next day")

    def test_status_badges(self):
        row = make_assignment(self.agent, self.day, self.today)
        self.assertContains(self._get(7), "Upcoming")
        self.assertContains(self._get(8, 40), "badge-waiting_for_branch")
        row.checked_in_at = aware(2026, 6, 2, 8, 45)
        row.save()
        on_shift = self._get(10)
        self.assertContains(on_shift, '<span class="badge badge-open">On shift</span>')
        row.checked_in_at = None
        row.save()
        self.assertContains(self._get(18), '<span class="badge badge-expired">Missed</span>')

    def test_today_row_is_marked_and_labelled(self):
        make_assignment(self.agent, self.day, self.today)
        response = self._get(7)
        self.assertContains(response, "shifts-day-pill--today")
        self.assertContains(response, "shifts-day--today")
        self.assertContains(response, "8 hours")

    def test_edit_hours_opens_a_dialog_for_night_rows_only(self):
        day_row = make_assignment(self.agent, self.day, self.today)
        night_row = make_assignment(self.agent, self.night, self.today)
        response = self._get(7)
        self.assertContains(response, reverse("shifts_mine_hours", args=[night_row.pk]))
        self.assertNotContains(response, reverse("shifts_mine_hours", args=[day_row.pk]))
        self.assertNotContains(response, 'name="start_time"')
        with _both_clocks(aware(2026, 6, 2, 7, 0), self.today):
            dialog = self.client.get(reverse("shifts_mine_hours", args=[night_row.pk]), HTTP_HX_REQUEST="true")
        self.assertContains(dialog, 'class="modal-header"')
        self.assertContains(dialog, 'value="22:00"')
        self.assertContains(dialog, "Save hours")
        self.assertNotContains(dialog, "Reset to shift default")

    def test_check_in_panel_has_one_check_in_control(self):
        make_assignment(self.agent, self.day, self.today)
        html = self._get(8, 40).content.decode()
        panel = html[html.index('id="shifts-check-in"'):html.index("shifts-mine-section")]
        self.assertEqual(panel.count(">Check in<"), 1)
        self.assertIn("Check-in open", panel)

    def test_empty_state(self):
        response = self._get(7)
        self.assertContains(response, 'class="empty-state"')
        self.assertContains(response, "You have no upcoming shifts.")


class RowRefreshAndDateTests(TestCase):
    def setUp(self):
        self.department = make_department("Refresh")
        self.agent = make_user(
            "refresh-agent",
            "support",
            self.department,
            make_role("Refresh agent", can_check_in=True),
        )
        self.day = make_shift_type(self.department, name="Refresh day", start=time(9, 0), end=time(17, 0))
        self.night = make_shift_type(self.department, name="Refresh night", start=time(22, 0), end=time(6, 0))
        self.today = date(2026, 10, 12)  # a Monday
        self.client.force_login(self.agent)

    def test_nearby_rows_reload_after_check_in(self):
        row = make_assignment(self.agent, self.day, self.today)
        later = make_assignment(self.agent, self.day, date(2026, 10, 20))
        with _both_clocks(aware(2026, 10, 12, 8, 40), self.today):
            page = self.client.get(reverse("shifts_mine"))
            self.client.post(reverse("shifts_check_in", args=[row.pk]), HTTP_HX_REQUEST="true")
            refreshed = self.client.get(reverse("shifts_mine_row", args=[row.pk]))
        self.assertContains(page, f'data-sig="{row.pk}:open"')
        self.assertContains(page, f'hx-get="{reverse("shifts_mine_row", args=[row.pk])}"')
        self.assertNotContains(page, reverse("shifts_mine_row", args=[later.pk]))
        self.assertContains(refreshed, '<span class="badge badge-open">On shift</span>')
        self.assertNotContains(refreshed, "Check-in open")

    def test_row_endpoint_is_owner_only(self):
        other = make_user("refresh-other", "support", self.department)
        row = make_assignment(other, self.day, self.today)
        self.assertEqual(self.client.get(reverse("shifts_mine_row", args=[row.pk])).status_code, 403)

    def test_arabic_dates_use_full_weekday_and_arabic_comma(self):
        make_assignment(self.agent, self.day, self.today)
        self.client.cookies["django_language"] = "ar"
        with _both_clocks(aware(2026, 10, 12, 7, 0), self.today):
            html = self.client.get(reverse("shifts_mine")).content.decode()
        self.assertIn("الاثنين، 12 أكتوبر 2026", html)
        self.assertIn("الاثنين 12 أكتوبر", html)
        self.assertNotIn("إثنين", html)
        self.assertNotIn("الاثنين, ", html)

    def test_english_dates(self):
        make_assignment(self.agent, self.day, self.today)
        with _both_clocks(aware(2026, 10, 12, 7, 0), self.today):
            html = self.client.get(reverse("shifts_mine")).content.decode()
        self.assertIn("Monday, 12 October 2026", html)
        self.assertIn("Mon 12 Oct", html)

    def test_two_shifts_are_two_entries(self):
        manager = make_user(
            "refresh-lead", "support", self.department, make_role("Refresh lead", can_manage_shifts=True),
        )
        make_assignment(self.agent, self.day, self.today)
        make_assignment(self.agent, self.night, self.today)
        self.client.force_login(manager)
        grid = self.client.get(reverse("shifts_rota_grid"), {
            "department": self.department.pk, "view": "week", "start": "2026-10-12",
        })
        self.assertContains(grid, 'data-count="2"', count=1)
        self.assertContains(grid, "shifts-cell--filled", count=2)
        self.assertContains(grid, '<span class="shifts-moon" aria-hidden="true"></span>', count=1)
        self.assertNotContains(grid, "shifts-split")


class NightTickBoxTests(TestCase):
    """A shift type is a night shift when the manager ticks Night shift."""

    def setUp(self):
        self.department = make_department("Window")

    def _type(self, name, start, end, is_night=False):
        return make_shift_type(self.department, name=name, start=start, end=end, is_night=is_night)

    def test_type_form_ticks_and_unticks_night(self):
        manager = make_user(
            "window-lead2", "support", self.department, make_role("Window lead2", can_manage_shifts=True),
        )
        self.client.force_login(manager)
        form = self.client.get(reverse("shifts_type_add"), {"department": self.department.pk}, HTTP_HX_REQUEST="true")
        self.assertContains(form, 'type="checkbox" name="is_night"')
        self.assertNotContains(form, "night_mode")
        page = self.client.get(reverse("shifts_types"), {"department": self.department.pk})
        self.assertNotContains(page, "night_starts_at")
        self.assertNotContains(page, "Set manually")
        payload = {"department": self.department.pk, "name": "Late", "start_time": "14:00",
                   "end_time": "22:00", "colour": "#112233"}
        self.client.post(reverse("shifts_type_add"), {**payload, "is_night": "on"})
        row = ShiftType.objects.get(department=self.department, name="Late")
        self.assertTrue(row.is_night)
        self.client.post(reverse("shifts_type_edit", args=[row.pk]), payload)
        row.refresh_from_db()
        self.assertFalse(row.is_night)
        # An overnight type is a day shift until ticked.
        self.client.post(reverse("shifts_type_add"), {**payload, "name": "Overnight",
                                                      "start_time": "22:00", "end_time": "06:00"})
        self.assertFalse(ShiftType.objects.get(department=self.department, name="Overnight").is_night)

    def test_no_next_day_badge_when_ending_at_midnight(self):
        # Dates stay before today: a clock ahead of real time trips session expiry.
        agent = make_user("window-agent2", "support", self.department)
        noon = self._type("Noon to midnight", time(12, 0), time(0, 0))
        overnight = self._type("Overnight", time(22, 0), time(6, 0), is_night=True)
        make_assignment(agent, noon, date(2026, 10, 9))
        self.client.force_login(agent)
        with _both_clocks(aware(2026, 10, 9, 9, 0), date(2026, 10, 9)):
            page = self.client.get(reverse("shifts_mine"))
        self.assertNotContains(page, "Next day")
        make_assignment(agent, overnight, date(2026, 10, 10))
        with _both_clocks(aware(2026, 10, 9, 9, 0), date(2026, 10, 9)):
            page = self.client.get(reverse("shifts_mine"))
        self.assertContains(page, "Next day", count=1)


class DirectVisitAndReadOnlyTests(TestCase):
    """Dialog URLs opened directly and the read-only row reload."""

    def setUp(self):
        self.department = make_department("Direct")
        self.manager = make_user(
            "direct-lead", "support", self.department, make_role("Direct lead", can_manage_shifts=True),
        )
        self.agent = make_user(
            "direct-agent", "support", self.department, make_role("Direct agent", can_check_in=True),
        )
        self.day = make_shift_type(self.department, name="Direct day", start=time(9, 0), end=time(17, 0))
        self.night = make_shift_type(self.department, name="Direct night")

    def test_type_edit_opened_directly_is_a_full_page(self):
        self.client.force_login(self.manager)
        page = self.client.get(reverse("shifts_type_edit", args=[self.day.pk]))
        self.assertTemplateUsed(page, "shifts/type_form_page.html")
        self.assertContains(page, 'id="shifts-shell-pane"')
        self.assertContains(page, 'class="settings-tab active"')
        self.assertNotContains(page, 'class="modal-header"')
        dialog = self.client.get(reverse("shifts_type_edit", args=[self.day.pk]), HTTP_HX_REQUEST="true")
        self.assertTemplateUsed(dialog, "shifts/type_form.html")
        self.assertNotContains(dialog, 'id="shifts-shell-pane"')
        saved = self.client.post(reverse("shifts_type_edit", args=[self.day.pk]), {
            "department": self.department.pk, "name": "Direct day", "start_time": "09:00",
            "end_time": "17:00", "colour": "#336699",
        })
        self.assertRedirects(saved, f"{reverse('shifts_types')}?department={self.department.pk}")
        bad = self.client.post(reverse("shifts_type_add"), {"department": self.department.pk, "name": "x"})
        self.assertTemplateUsed(bad, "shifts/type_form_page.html")

    def test_dialog_only_urls_redirect_when_opened_directly(self):
        self.client.force_login(self.manager)
        cell = self.client.get(reverse("shifts_cell"), {
            "user": self.agent.pk, "date": "2026-06-02", "department": self.department.pk,
        })
        self.assertEqual(cell.status_code, 302)
        self.assertIn(reverse("shifts_rota"), cell["Location"])
        self.assertIn("start=2026-05-30", cell["Location"])
        row = make_assignment(self.agent, self.night, timezone.localdate())
        self.client.force_login(self.agent)
        hours = self.client.get(reverse("shifts_mine_hours", args=[row.pk]))
        self.assertRedirects(hours, reverse("shifts_mine"))

    def test_row_reload_does_not_write(self):
        on = date(2026, 6, 2)
        row = make_assignment(self.agent, self.day, on)
        row.checked_in_at = aware(2026, 6, 2, 8, 50)
        row.save()
        self.client.force_login(self.agent)
        with _both_clocks(aware(2026, 6, 2, 18, 0), on):
            response = self.client.get(reverse("shifts_mine_row", args=[row.pk]))
        self.assertContains(response, '<span class="badge badge-closed">Completed</span>')
        row.refresh_from_db()
        self.assertIsNotNone(row.checked_in_at)
        self.assertIsNone(row.checked_out_at)


class PhoneCardLabelTests(TestCase):
    def test_shift_types_cells_carry_labels_for_phone_cards(self):
        department = make_department("Labels")
        manager = make_user("labels-lead", "support", department, make_role("Labels lead", can_manage_shifts=True))
        make_shift_type(department, name="Labels day", start=time(9, 0), end=time(17, 0))
        self.client.force_login(manager)
        page = self.client.get(reverse("shifts_types"), {"department": department.pk})
        for label in ("Name", "Times", "Night", "Colour", "Actions"):
            self.assertContains(page, f'data-label="{label}"')


class WorkHoursReportTests(TestCase):
    """The Calculator's work hours card and its Excel export."""

    def setUp(self):
        self.department = make_department("Report")
        self.manager = make_user(
            "report-lead", "support", self.department, make_role("Report lead", can_manage_shifts=True),
        )
        self.agent = make_user(
            "report-agent", "support", self.department, make_role("Report agent", can_check_in=True),
            first_name="Omar", last_name="Elrefaey",
        )
        self.evening = make_shift_type(self.department, name="Evening", start=time(18, 0), end=time(23, 0))
        self.long_day = make_shift_type(self.department, name="Long day", start=time(12, 0), end=time(0, 0))
        self.client.force_login(self.manager)

    def _check_in(self, row):
        _start, end = effective_interval(row)
        row.checked_in_at = end - timedelta(hours=1)
        row.checked_out_at = None
        row.save()

    def _september(self):
        """Ten shifts like the paper statement: eight evenings and two 12-hour Fridays."""
        rows = []
        for day in (2, 7, 10, 13, 16, 19, 22, 29):
            rows.append(make_assignment(self.agent, self.evening, date(2026, 9, day)))
        for day in (4, 25):
            rows.append(make_assignment(self.agent, self.long_day, date(2026, 9, day)))
        for row in rows:
            self._check_in(row)
        return rows

    def _get(self, **params):
        query = {"person": self.agent.pk, "start": "2026-09-01", "end": "2026-09-30", "shifts": "all", **params}
        with _both_clocks(aware(2026, 10, 5, 12, 0), date(2026, 10, 5)):
            return self.client.get(reverse("shifts_calc_work"), query)

    def test_statement_matches_the_paper_version(self):
        self._september()
        response = self._get()
        report = response.context["report"]
        self.assertEqual(len(report.rows), 10)
        self.assertEqual(report.days_present, 10)
        self.assertEqual(report.total_hours, 64)
        self.assertEqual([row.number for row in report.rows], list(range(1, 11)))
        self.assertEqual(report.rows[1].day, date(2026, 9, 4))
        self.assertEqual(report.rows[1].hours, 12)
        self.assertEqual((report.rows[1].starts, report.rows[1].ends), ("12:00 PM", "12:00 AM"))
        self.assertContains(response, "Omar Elrefaey")
        self.assertContains(response, "04/09/2026")
        self.assertContains(response, '<span class="badge badge-open">Present</span>', count=10)
        self.assertContains(response, reverse("shifts_calc_work_export"))

    def test_status_and_totals(self):
        present, absent, upcoming = (
            make_assignment(self.agent, self.evening, date(2026, 10, day)) for day in (1, 2, 6)
        )
        self._check_in(present)
        cleared = make_assignment(self.agent, self.evening, date(2026, 10, 3))
        cleared.checked_out_at = aware(2026, 10, 3, 18, 30)  # a manager cleared it early
        cleared.save()
        report = self._get(start="2026-10-01", end="2026-10-31").context["report"]
        self.assertEqual(
            [row.status for row in report.rows], ["present", "absent", "absent", "upcoming"],
        )
        self.assertEqual((report.days_present, report.total_hours), (1, 5))
        self.assertEqual([row.hours for row in report.rows], [5, 0, 0, 5])

    def test_present_shift_starts_when_the_agent_checked_in(self):
        row = make_assignment(self.agent, self.evening, date(2026, 10, 1))
        self.client.force_login(self.agent)
        with _both_clocks(aware(2026, 10, 1, 18, 30), date(2026, 10, 1)):
            self.client.post(reverse("shifts_check_in", args=[row.pk]))
        self.client.force_login(self.manager)
        from shifts.services import settle_check_in

        row.refresh_from_db()
        settle_check_in(row, aware(2026, 10, 2, 0, 0))
        row.refresh_from_db()
        # The shift has ended and the check-in is cleared, but the press time stays.
        self.assertIsNone(row.checked_in_at)
        self.assertEqual(row.arrived_at, aware(2026, 10, 1, 18, 30))
        report = self._get(start="2026-10-01", end="2026-10-31").context["report"]
        self.assertEqual((report.rows[0].starts, report.rows[0].ends), ("6:30 PM", "11:00 PM"))
        self.assertEqual((report.rows[0].hours, report.total_hours), (Decimal("4.5"), Decimal("4.5")))

    def test_early_check_in_counts_from_the_shift_start(self):
        day = make_shift_type(self.department, name="Report day", start=time(9, 0), end=time(17, 0))
        row = make_assignment(self.agent, day, date(2026, 10, 1))
        row.arrived_at = aware(2026, 10, 1, 8, 40)
        row.checked_out_at = aware(2026, 10, 1, 17, 0)
        row.save()
        report = self._get(start="2026-10-01", end="2026-10-31").context["report"]
        self.assertEqual(report.rows[0].starts, "8:40 AM")  # the press time still shows
        self.assertEqual((report.rows[0].hours, report.total_hours), (8, 8))

    def test_running_shift_counts_only_up_to_now(self):
        row = make_assignment(self.agent, self.evening, date(2026, 10, 1))
        row.checked_in_at = row.arrived_at = aware(2026, 10, 1, 18, 0)
        row.save()
        query = {"person": self.agent.pk, "start": "2026-10-01", "end": "2026-10-31", "shifts": "all"}
        with _both_clocks(aware(2026, 10, 1, 19, 30), date(2026, 10, 1)):
            report = self.client.get(reverse("shifts_calc_work"), query).context["report"]
        self.assertEqual(report.rows[0].status, "present")
        self.assertEqual((report.rows[0].hours, report.total_hours), (Decimal("1.5"), Decimal("1.5")))

    def test_people_who_do_not_check_in_show_scheduled_and_count(self):
        plain = make_user("report-plain", "support", self.department)
        make_assignment(plain, self.evening, date(2026, 10, 1))
        report = self._get(person=plain.pk, start="2026-10-01", end="2026-10-31").context["report"]
        self.assertEqual([row.status for row in report.rows], ["scheduled"])
        self.assertEqual((report.days_present, report.total_hours), (1, 5))

    def test_night_only_filter(self):
        self._september()
        report = self._get(shifts="night").context["report"]
        # Evening (6-11 PM) is all night; Long day (noon-midnight) is half, so not a night shift.
        self.assertEqual(len(report.rows), 8)
        self.assertEqual(report.total_hours, 40)
        self.assertIn("night shifts", report.title)

    def test_picked_shift_types_filter(self):
        self._september()
        response = self._get(shift_type=[self.long_day.pk])
        report = response.context["report"]
        self.assertEqual(len(report.rows), 2)
        self.assertEqual(report.total_hours, 24)
        self.assertEqual(report.title, "Statement of work hours: Long day")
        # The Excel link carries the same pick.
        self.assertContains(response, f"shift_type={self.long_day.pk}")
        both = self._get(shift_type=[self.evening.pk, self.long_day.pk]).context["report"]
        self.assertEqual((len(both.rows), both.total_hours), (10, 64))
        # Every type the department offers ticked is all shifts, as on the page.
        self.assertEqual(both.title, "Statement of work hours")
        make_shift_type(self.department, name="Early", start=time(6, 0), end=time(12, 0))
        two = self._get(shift_type=[self.evening.pk, self.long_day.pk]).context["report"]
        self.assertEqual(two.title, "Statement of work hours: Long day and Evening")  # start-time order

    def test_night_only_wins_over_picked_types(self):
        self._september()
        report = self._get(shifts="night", shift_type=[self.long_day.pk]).context["report"]
        self.assertEqual(len(report.rows), 8)
        self.assertIn("night shifts", report.title)

    def test_junk_shift_type_ids_mean_all_shifts(self):
        self._september()
        report = self._get(shift_type=["x", "99999"]).context["report"]
        self.assertEqual(len(report.rows), 10)
        self.assertEqual(report.title, "Statement of work hours")

    def test_calculator_shows_a_switch_per_shift_type(self):
        page = self.client.get(reverse("shifts_calculator"), {"department": self.department.pk})
        content = page.content.decode()
        self.assertIn('id="work-shifts-field"', content)
        for shift_type in (self.evening, self.long_day):
            self.assertIn(f'name="shift_type" value="{shift_type.pk}"', content)
        self.assertIn("data-multi-night", content)  # Evening is a night shift
        self.assertNotIn('hx-swap-oob="true"', content[content.index('id="work-shifts-field"') - 60:])
        # Changing department swaps in that department's switches.
        other = make_department("Report other")
        mine = make_shift_type(other, name="Other early", start=time(6, 0), end=time(14, 0))
        part = self.client.get(reverse("shifts_calculator"), {"department": other.pk, "part": "work_people"})
        part_content = part.content.decode()
        self.assertIn('id="work-shifts-field" hx-swap-oob="true"', part_content)
        self.assertIn(f'value="{mine.pk}"', part_content)
        self.assertNotIn(f'value="{self.evening.pk}"', part_content)
        self.assertNotIn("data-multi-night", part_content)  # no night types there

    def test_report_is_read_only(self):
        row = make_assignment(self.agent, self.evening, date(2026, 10, 1))
        row.checked_in_at = aware(2026, 10, 1, 18, 5)
        row.save()
        self._get(start="2026-10-01", end="2026-10-31")
        row.refresh_from_db()
        self.assertIsNotNone(row.checked_in_at)

    def test_missing_person_and_empty_period(self):
        self.assertContains(self._get(person=""), "Choose a person.")
        self.assertContains(self._get(), "No shifts in this period.")

    def test_excel_export(self):
        from io import BytesIO

        from openpyxl import load_workbook

        self._september()
        with _both_clocks(aware(2026, 10, 5, 12, 0), date(2026, 10, 5)):
            response = self.client.get(reverse("shifts_calc_work_export"), {
                "person": self.agent.pk, "start": "2026-09-01", "end": "2026-09-30", "shifts": "all",
            })
        self.assertEqual(
            response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn('filename="work-hours_report-agent_2026-09-01_2026-09-30.xlsx"', response["Content-Disposition"])
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet["A1"].value, "Name: Omar Elrefaey")
        self.assertEqual(
            [cell.value for cell in sheet[5]],
            ["No.", "Day", "Date", "Status", "From", "To", "Work hours"],
        )
        self.assertEqual(sheet["A7"].value, 2)
        self.assertEqual(sheet["B7"].value, "Friday")
        self.assertEqual(sheet["C7"].value.date(), date(2026, 9, 4))
        self.assertEqual((sheet["E7"].value, sheet["F7"].value, sheet["G7"].value), ("12:00 PM", "12:00 AM", 12))
        self.assertEqual(sheet["D7"].value, "Present")
        self.assertEqual((sheet["F17"].value, sheet["G17"].value), (10, "Days present"))
        self.assertEqual((sheet["F18"].value, sheet["G18"].value), (64, "Sum of work hours"))
        self.assertFalse(sheet.sheet_view.rightToLeft)

    def test_arabic_export_is_right_to_left(self):
        from io import BytesIO

        from openpyxl import load_workbook

        self._september()
        self.client.cookies["django_language"] = "ar"
        with _both_clocks(aware(2026, 10, 5, 12, 0), date(2026, 10, 5)):
            response = self.client.get(reverse("shifts_calc_work_export"), {
                "person": self.agent.pk, "start": "2026-09-01", "end": "2026-09-30",
            })
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertEqual(sheet["D6"].value, "حاضر")
        self.assertEqual(sheet["G17"].value, "عدد أيام الحضور")
        self.assertEqual(sheet["B5"].value, "اليوم")

    def test_support_users_cannot_open_the_report(self):
        self.client.force_login(self.agent)
        self.assertEqual(self.client.get(reverse("shifts_calc_work"), {"person": self.agent.pk}).status_code, 403)
        self.assertEqual(
            self.client.get(reverse("shifts_calc_work_export"), {"person": self.agent.pk}).status_code, 403,
        )

    def test_calculator_shows_three_cards(self):
        response = self.client.get(reverse("shifts_calculator"))
        self.assertContains(response, 'id="shifts-length"')
        self.assertContains(response, 'id="shifts-rotation"')
        self.assertContains(response, 'id="shifts-work-hours"')
        self.assertNotContains(response, 'id="shifts-hours"')
        self.assertNotContains(response, 'id="shifts-coverage"')
        self.assertContains(response, "Omar Elrefaey")

    def test_department_picker_reloads_only_its_people(self):
        other = make_department("Report other")
        stranger = make_user("report-stranger", "support", other, first_name="Sara", last_name="Ali")
        page = self.client.get(reverse("shifts_calculator"))
        self.assertContains(page, "hx-trigger=\"change[target.name!=='department']\"")
        self.assertContains(page, "Omar Elrefaey")
        self.assertNotContains(page, "Sara Ali")
        swapped = self.client.get(reverse("shifts_calculator"), {"part": "work_people", "department": other.pk})
        self.assertContains(swapped, 'id="work-person-field"')
        self.assertContains(swapped, f'<option value="{stranger.pk}">Sara Ali</option>')
        self.assertNotContains(swapped, "Omar Elrefaey")
        self.assertContains(swapped, 'id="shifts-work-result" hx-swap-oob="innerHTML"')
        empty = make_department("Report empty")
        nobody = self.client.get(reverse("shifts_calculator"), {"part": "work_people", "department": empty.pk})
        self.assertContains(nobody, "No support staff in this department")
        self.assertContains(nobody, "disabled")


class MyShiftsTabTests(TestCase):
    """My shifts is for every support-type account, superusers included."""

    def setUp(self):
        self.department = make_department("Tab")
        self.lead = make_user(
            "tab-lead", "support", self.department, make_role("Tab lead", can_manage_shifts=True),
        )
        self.admin = make_user("tab-admin", "support", superuser=True)
        self.branch_admin = make_user("tab-branch-admin", "branch", superuser=True)

    def test_support_superuser_has_my_shifts_tab(self):
        self.client.force_login(self.admin)
        rota = self.client.get(reverse("shifts_rota"))
        self.assertContains(rota, f'href="{reverse("shifts_mine")}"')
        self.assertContains(rota, f'href="{reverse("shifts_team")}"')
        self.assertEqual(self.client.get(reverse("shifts_mine")).status_code, 200)
        self.assertEqual(self.client.get(reverse("shifts_check_in_panels")).status_code, 200)

    def test_branch_superuser_has_no_my_shifts_tab(self):
        self.client.force_login(self.branch_admin)
        rota = self.client.get(reverse("shifts_rota"))
        self.assertNotContains(rota, f'href="{reverse("shifts_mine")}"')
        self.assertEqual(self.client.get(reverse("shifts_check_in_panels")).status_code, 403)

    def test_team_lead_keeps_my_shifts(self):
        self.client.force_login(self.lead)
        rota = self.client.get(reverse("shifts_rota"))
        self.assertContains(rota, f'href="{reverse("shifts_mine")}"')
        self.assertEqual(self.client.get(reverse("shifts_mine")).status_code, 200)


@override_settings(TIME_ZONE="Africa/Cairo")
class ShiftStatusConsistencyTests(TestCase):
    """Every ended shift shows one status, the same on the rota, My shifts and the report."""

    def setUp(self):
        self.department = make_department("Status")
        role = make_role("Status agent", can_check_in=True)
        self.present = make_user("status-present", "support", self.department, role)
        self.absent = make_user("status-absent", "support", self.department, role)
        self.no_check_in = make_user("status-plain", "support", self.department, make_role("Plain"))
        self.day = make_shift_type(self.department, name="Status day", start=time(9, 0), end=time(17, 0))
        self.night = make_shift_type(self.department, name="Status night")
        self.today = date(2026, 6, 2)
        self.present_row = make_assignment(self.present, self.day, self.today)
        self.present_row.checked_in_at = aware(2026, 6, 2, 9, 5)
        self.present_row.save(update_fields=["checked_in_at"])
        self.absent_row = make_assignment(self.absent, self.day, self.today)
        self.plain_row = make_assignment(self.no_check_in, self.day, self.today)

    def _bundle(self, now, dates):
        from shifts.views import _rota_bundle

        with _both_clocks(now, self.today):
            return _rota_bundle(self.department, dates)

    def _markers(self, now):
        dates = [self.today - timedelta(days=1), self.today]
        bundle = self._bundle(now, dates)
        return {row.pk: row.check_marker for row in bundle["assignments"] if row.date in dates}

    def _mine_status(self, user, row, now):
        from shifts.access import can_check_in
        from shifts.services import settle_check_in
        from shifts.views import _decorate_mine

        row = ShiftAssignment.objects.select_related("shift_type", "times_set_by").get(pk=row.pk)
        settle_check_in(row, now, save=False)
        return _decorate_mine(row, now, self.today, can_check_in(user)).status

    def _report_status(self, user, now):
        from shifts.reports import build_report

        with _both_clocks(now, self.today):
            report = build_report(user, self.today, self.today, now=now)
        return report.rows[0].status

    def test_ended_shifts_agree_across_pages(self):
        after = aware(2026, 6, 2, 18, 0)
        markers = self._markers(after)
        self.assertEqual(markers[self.present_row.pk], "completed")
        self.assertEqual(markers[self.absent_row.pk], "missed")
        self.assertEqual(markers[self.plain_row.pk], "completed")

        self.assertEqual(self._mine_status(self.present, self.present_row, after), "completed")
        self.assertEqual(self._mine_status(self.absent, self.absent_row, after), "missed")
        self.assertEqual(self._mine_status(self.no_check_in, self.plain_row, after), "completed")

        self.assertEqual(self._report_status(self.present, after), "present")
        self.assertEqual(self._report_status(self.absent, after), "absent")
        self.assertEqual(self._report_status(self.no_check_in, after), "scheduled")

    def test_running_shifts_agree_across_pages(self):
        # The report used to call these "Upcoming" while they were running.
        noon = aware(2026, 6, 2, 12, 0)
        self.assertEqual(self._mine_status(self.no_check_in, self.plain_row, noon), "on_shift")
        self.assertEqual(self._report_status(self.no_check_in, noon), "on_shift")
        self.assertEqual(self._report_status(self.present, noon), "present")
        self.assertEqual(self._report_status(self.absent, noon), "not_checked_in")
        from shifts.reports import build_report

        with _both_clocks(noon, self.today):
            row = build_report(self.no_check_in, self.today, self.today, now=noon).rows[0]
        self.assertEqual(row.status_text, "On shift")
        self.assertTrue(row.counted)
        self.assertEqual(row.hours, 3)  # 9 AM to noon so far, like a checked-in shift

    def test_rota_labels_every_ended_shift(self):
        manager = make_user(
            "status-lead", "support", self.department, make_role("Status lead", can_manage_shifts=True),
        )
        self.client.force_login(manager)
        with _both_clocks(aware(2026, 6, 2, 18, 0), self.today):
            rota = self.client.get(reverse("shifts_rota"), {
                "department": self.department.pk, "view": "week", "start": "2026-05-30",
            })
        content = rota.content.decode()
        self.assertEqual(content.count('shifts-st shifts-st--muted">Completed<'), 2)
        self.assertEqual(content.count('shifts-st shifts-st--miss">Missed<'), 1)

    def test_running_night_from_yesterday_is_not_missed(self):
        night = make_assignment(self.absent, self.night, self.today - timedelta(days=1))
        self.assertEqual(self._markers(aware(2026, 6, 2, 3, 0))[night.pk], "not")
        self.assertEqual(self._markers(aware(2026, 6, 2, 7, 0))[night.pk], "missed")

    def test_rota_etag_changes_when_a_shift_ends(self):
        from shifts.views import _rota_etag

        during = self._bundle(aware(2026, 6, 2, 16, 0), [self.today])
        after = self._bundle(aware(2026, 6, 2, 18, 0), [self.today])
        self.assertNotEqual(
            _rota_etag(self.department, "week", self.today, during["assignments"]),
            _rota_etag(self.department, "week", self.today, after["assignments"]),
        )


@override_settings(TIME_ZONE="Africa/Cairo")
class SuperuserCheckInTests(TestCase):
    def test_support_superuser_with_a_shift_can_check_in(self):
        from shifts.access import can_check_in

        department = make_department("Admin desk")
        # No role flag needed: a superuser has every permission.
        admin = make_user("rota-admin", "support", department, make_role("Admin role"), superuser=True)
        day = make_shift_type(department, name="Admin day", start=time(9, 0), end=time(17, 0))
        row = make_assignment(admin, day, date(2026, 6, 2))
        self.assertTrue(can_check_in(admin))
        self.client.force_login(admin)
        with _both_clocks(aware(2026, 6, 2, 10, 0), date(2026, 6, 2)):
            posted = self.client.post(reverse("shifts_check_in", args=[row.pk]))
        self.assertEqual(posted.status_code, 200)
        row.refresh_from_db()
        self.assertIsNotNone(row.checked_in_at)

    def test_branch_superuser_cannot_check_in(self):
        from shifts.access import can_check_in

        self.assertFalse(can_check_in(make_user("branch-root", "branch", superuser=True)))


@override_settings(TIME_ZONE="Africa/Cairo")
class StaleMyShiftsPageTests(TestCase):
    """A My shifts tab left open across logout/login reloads instead of 403-ing."""

    def setUp(self):
        department = make_department("Stale")
        role = make_role("Stale agent", can_check_in=True)
        self.first = make_user("stale-first", "support", department, role)
        self.second = make_user("stale-second", "support", department, role)
        day = make_shift_type(department, name="Stale day", start=time(9, 0), end=time(17, 0))
        self.row = make_assignment(self.first, day, date(2026, 6, 2))
        self.client.force_login(self.second)

    def test_row_refresh_for_previous_user_reloads_page(self):
        url = reverse("shifts_mine_row", args=[self.row.pk])
        with _both_clocks(aware(2026, 6, 2, 10, 0), date(2026, 6, 2)):
            background = self.client.get(url, HTTP_HX_REQUEST="true")
            direct = self.client.get(url)
        self.assertEqual(background.status_code, 204)
        self.assertEqual(background["HX-Refresh"], "true")
        self.assertEqual(direct.status_code, 403)

    def test_panel_poll_for_previous_user_reloads_page(self):
        url = reverse("shifts_check_in_panels")
        with _both_clocks(aware(2026, 6, 2, 10, 0), date(2026, 6, 2)):
            stale = self.client.get(url, {"u": self.first.pk}, HTTP_HX_REQUEST="true")
            current = self.client.get(url, {"u": self.second.pk}, HTTP_HX_REQUEST="true")
        self.assertEqual(stale.status_code, 204)
        self.assertEqual(stale["HX-Refresh"], "true")
        self.assertEqual(current.status_code, 200)
        self.assertFalse(current.has_header("HX-Refresh"))

    def test_mine_page_polls_with_its_user(self):
        self.client.force_login(self.first)
        with _both_clocks(aware(2026, 6, 2, 8, 45), date(2026, 6, 2)):
            page = self.client.get(reverse("shifts_mine"))
        self.assertContains(page, f'?u={self.first.pk}"')


@patch("notifications.email_jobs.send_with_retries", return_value=True)
class ShiftReminderTests(TestCase):
    """Agents are reminded once, when check-in opens 30 minutes before the shift."""

    def setUp(self):
        from core.models import EmailSetting

        self.department = make_department("Remind")
        self.role = make_role("Remind agent", can_check_in=True)
        self.agent = make_user("remind-agent", "support", self.department, self.role)
        self.day = make_shift_type(self.department, name="Remind day", start=time(9, 0), end=time(17, 0))
        self.row = make_assignment(self.agent, self.day, date(2026, 6, 2))
        self.email = EmailSetting.objects.create(
            smtp_host="smtp.test.local", smtp_port=587, smtp_email="noreply@test.local",
            smtp_password="secret", encryption="tls", from_name="Test", from_email="noreply@test.local",
            is_active=True, notify_shift_reminder=True,
        )

    def _send(self, hh, mm=0):
        from shifts.reminders import send_due_reminders

        return send_due_reminders(aware(2026, 6, 2, hh, mm))

    def _reminders(self, user=None):
        from notifications.models import InAppNotification

        return InAppNotification.objects.filter(recipient=user or self.agent, notification_type="shift")

    def test_sent_once_when_check_in_opens(self, send_mail):
        self.assertEqual(self._send(8, 20), 0)
        self.assertFalse(self._reminders().exists())
        self.assertEqual(self._send(8, 30), 1)
        note = self._reminders().get()
        self.assertEqual(note.title, "Your shift starts at 9:00 AM")
        self.assertEqual(note.message, "Remind day · Remind. You can check in now.")
        self.assertIn("9:00 ص", note.title_ar)
        self.assertEqual(note.link, "/shifts/mine/#shifts-check-in")
        send_mail.assert_called_once()
        subject, _body, recipients = send_mail.call_args.args[:3]
        self.assertEqual(recipients, ["remind-agent@example.com"])
        self.assertIn("Your shift starts at 9:00 AM", subject)
        self.assertIn("/shifts/mine/", send_mail.call_args.kwargs["html_body"])
        self.row.refresh_from_db()
        self.assertEqual(self.row.reminded_for, aware(2026, 6, 2, 9, 0))
        # Later ticks inside the window do not repeat it.
        self.assertEqual(self._send(8, 45), 0)
        self.assertEqual(self._reminders().count(), 1)

    def test_moved_start_is_reminded_again(self, send_mail):
        self._send(8, 30)
        self.row.start_time_override = time(10, 0)
        self.row.end_time_override = time(17, 0)
        self.row.save()
        self.assertEqual(self._send(9, 15), 0)
        self.assertEqual(self._send(9, 30), 1)
        self.assertEqual(self._reminders().count(), 2)

    def test_skipped_when_checked_in_started_inactive_or_branch(self, send_mail):
        self.row.checked_in_at = aware(2026, 6, 2, 8, 31)
        self.row.save()
        self.assertEqual(self._send(8, 40), 0)
        self.row.checked_in_at = None
        self.row.save()
        self.assertEqual(self._send(9, 0), 0)  # the shift has started: too late to remind
        branch = make_user("remind-branch", "branch", self.department)
        make_assignment(branch, self.day, date(2026, 6, 3))
        gone = make_user("remind-gone", "support", self.department, self.role)
        gone.status = User.Status.INACTIVE
        gone.save()
        make_assignment(gone, self.day, date(2026, 6, 3))
        from shifts.reminders import send_due_reminders

        self.assertEqual(send_due_reminders(aware(2026, 6, 3, 8, 40)), 0)
        send_mail.assert_not_called()

    def test_email_checkbox_off_keeps_the_bell_only(self, send_mail):
        from notifications.email_service import clear_email_setting_cache

        self.email.notify_shift_reminder = False
        self.email.save()
        clear_email_setting_cache()
        self.assertEqual(self._send(8, 30), 1)
        self.assertTrue(self._reminders().exists())
        send_mail.assert_not_called()

    def test_agents_without_check_in_get_a_plain_reminder(self, send_mail):
        plain = make_user("remind-plain", "support", self.department, make_role("Remind plain"))
        make_assignment(plain, self.day, date(2026, 6, 2))
        self._send(8, 30)
        self.assertEqual(self._reminders(plain).get().message, "Remind day · Remind. Have a good shift.")
        # The email follows the role too: no "check in" line or button.
        emails = {call.args[2][0]: call for call in send_mail.call_args_list}
        plain_text, plain_html = emails["remind-plain@example.com"].args[1], emails["remind-plain@example.com"].kwargs["html_body"]
        self.assertIn("Have a good shift.", plain_text)
        self.assertNotIn("Check-in is open now", plain_text)
        self.assertIn("View my shifts", plain_html)
        self.assertNotIn(">Check in<", plain_html)
        self.assertIn("Check-in is open now.", emails["remind-agent@example.com"].args[1])
        self.assertIn(">Check in<", emails["remind-agent@example.com"].kwargs["html_body"])

    def test_management_command_sends_the_same(self, send_mail):
        from io import StringIO

        from django.core.management import call_command

        out = StringIO()
        with patch("django.utils.timezone.now", return_value=aware(2026, 6, 2, 8, 35)):
            call_command("send_shift_reminders", stdout=out)
        self.assertIn("Sent 1 shift reminder(s).", out.getvalue())
        self.assertEqual(self._reminders().count(), 1)


    def test_email_settings_form_saves_the_shift_reminder_checkbox(self, send_mail):
        from core.forms import EmailSettingForm

        data = {
            "smtp_host": "smtp.test.local", "smtp_port": 587, "smtp_email": "noreply@test.local",
            "smtp_password": "secret", "encryption": "tls", "from_name": "Test",
            "from_email": "noreply@test.local", "is_active": "on",
        }
        form = EmailSettingForm(data, instance=self.email)
        self.assertIn("notify_shift_reminder", form.fields)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertFalse(form.save().notify_shift_reminder)  # unticked
        form = EmailSettingForm({**data, "notify_shift_reminder": "on"}, instance=self.email)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertTrue(form.save().notify_shift_reminder)


class WhoAvailableTests(TestCase):
    """The ticket chat's "Who's available" popup and the on-shift marks in Transfer."""

    def setUp(self):
        from core.models import Branch, Category
        from tickets.models import Ticket

        self.department = make_department("Help desk")
        other = make_department("Elsewhere")
        role = make_role("Avail agent", can_check_in=True)
        self.assignee = make_user("avail-owner", "support", self.department, role)
        self.on_agent = make_user("avail-on", "support", self.department, role, first_name="Mona", last_name="Adel")
        self.late_agent = make_user("avail-late", "support", self.department, role)
        self.evening_agent = make_user("avail-evening", "support", self.department, role)
        self.other_agent = make_user("avail-other", "support", other, role)
        day = make_shift_type(self.department, name="Avail day", start=time(9, 0), end=time(17, 0))
        evening = make_shift_type(self.department, name="Avail evening", start=time(18, 0), end=time(23, 0))
        other_day = make_shift_type(other, name="Other day", start=time(9, 0), end=time(17, 0))
        today = date(2026, 6, 2)
        for agent, shift_type, checked in (
            (self.on_agent, day, True),
            (self.late_agent, day, False),
            (self.evening_agent, evening, False),
            (self.other_agent, other_day, True),
        ):
            row = make_assignment(agent, shift_type, today)
            if checked:
                row.checked_in_at = aware(2026, 6, 2, 8, 55)
                row.save()
        branch = Branch.objects.create(code="AV", name="Avail branch")
        category = Category.objects.create(department=self.department, name="Avail category")
        self.branch_user = make_user("avail-branch", "branch")
        self.branch_user.branch = branch
        self.branch_user.save()
        self.ticket = Ticket.objects.create(
            ticket_number="TK-AV1", title="Printer", description="Offline", branch=branch,
            department=self.department, category=category, created_by=self.branch_user,
            assigned_to=self.assignee, client_name="Desk", client_phone="123",
        )

    def _popup(self, viewer):
        self.client.force_login(viewer)
        url = reverse("shifts_available_department", args=[self.department.pk])
        with _both_clocks(aware(2026, 6, 2, 12, 0), date(2026, 6, 2)):
            return self.client.get(url, {"ticket": self.ticket.pk})

    def test_assignee_sees_who_is_on_and_can_transfer(self):
        popup = self._popup(self.assignee)
        content = popup.content.decode()
        self.assertContains(popup, "Mona Adel")
        self.assertContains(popup, "Avail day")
        self.assertContains(popup, "5:00 PM")
        self.assertContains(popup, f'data-who-transfer="{self.on_agent.pk}"')
        # Not checked in is listed (as on the board) but not offered for Transfer;
        # the evening shift is later today.
        self.assertLess(content.index("Not checked in"), content.index("avail-late"))
        self.assertLess(content.index("avail-late"), content.index("Later today"))
        self.assertLess(content.index("Later today"), content.index("avail-evening"))
        self.assertNotIn(f'data-who-transfer="{self.late_agent.pk}"', content)
        self.assertNotContains(popup, "avail-other")

    def test_branch_viewer_sees_names_without_shift_types_or_transfer(self):
        popup = self._popup(self.branch_user)
        self.assertContains(popup, "Mona Adel")
        self.assertNotContains(popup, "Avail day")
        self.assertNotContains(popup, "data-who-transfer")

    def test_other_agents_and_pending_transfers_get_no_transfer_button(self):
        self.assertNotContains(self._popup(self.late_agent), "data-who-transfer")
        self.ticket.pending_transfer_to = self.late_agent
        self.ticket.save()
        self.assertNotContains(self._popup(self.assignee), "data-who-transfer")

    def test_chat_has_the_button_and_transfer_lists_on_shift_first(self):
        admin = make_user("avail-admin", "support", self.department, superuser=True)
        self.client.force_login(admin)
        with _both_clocks(aware(2026, 6, 2, 12, 0), date(2026, 6, 2)):
            page = self.client.get(reverse("ticket_detail", args=[self.ticket.pk]))
        content = page.content.decode()
        self.assertIn(reverse("shifts_available_department", args=[self.department.pk]), content)
        self.assertIn('id="who-available-modal"', content)
        options = content[content.index('id="new_assignee"'):]
        self.assertIn("Mona Adel (Help desk) — On shift until 5:00 PM", options)
        self.assertLess(options.index("Mona Adel"), options.index("avail-late"))

    def _available(self, viewer, name="shifts_available", **headers):
        self.client.force_login(viewer)
        with _both_clocks(aware(2026, 6, 2, 12, 0), date(2026, 6, 2)):
            return self.client.get(reverse(name), **headers)

    def test_board_and_popup_show_phones_with_copy_buttons(self):
        self.on_agent.phone = "+201000000001"
        self.on_agent.save()
        for page in (self._available(self.branch_user), self._popup(self.branch_user)):
            self.assertContains(page, 'href="tel:+201000000001"')
            self.assertContains(page, 'data-copy-text="+201000000001"')
            # Agents without a phone get no link or button.
            self.assertContains(page, 'href="tel:', count=1)
        # The search box matches the number too.
        self.assertContains(self._available(self.branch_user), 'data-filter-item="mona adel +201000000001"')

    def test_board_etag_follows_phone_changes(self):
        first = self._available(self.branch_user, "shifts_available_board", HTTP_HX_REQUEST="true")
        self.on_agent.phone = "+201000000002"
        self.on_agent.save()
        again = self._available(
            self.branch_user, "shifts_available_board",
            HTTP_HX_REQUEST="true", HTTP_IF_NONE_MATCH=first["ETag"],
        )
        self.assertEqual(again.status_code, 200)
        self.assertContains(again, "+201000000002")

    def test_contact_us_lists_shift_managers_with_phones(self):
        self.assertNotContains(self._available(self.branch_user), "Contact us")
        lead = make_user("avail-lead", "support", self.department, make_role("Avail lead", can_manage_shifts=True),
                         first_name="Sara", last_name="Lead")
        admin = make_user("avail-boss", "support", superuser=True)
        gone = make_user("avail-gone", "support", self.department, lead.role)
        gone.status = User.Status.INACTIVE
        for person, phone in ((lead, "+201000000010"), (admin, "+201000000011"),
                              (gone, "+201000000012"), (self.late_agent, "+201000000013")):
            person.phone = phone
            person.save()
        page = self._available(self.branch_user)
        contact = page.content.decode().split('id="shifts-contact-title"', 1)[1]
        self.assertIn("Sara Lead", contact)
        self.assertIn('data-copy-text="+201000000010"', contact)
        self.assertIn('data-copy-text="+201000000011"', contact)
        self.assertNotIn("+201000000012", contact)  # inactive
        self.assertNotIn("+201000000013", contact)  # not a manager
        self.assertLess(contact.index("avail-boss"), contact.index("Sara Lead"))

    def test_copy_script_is_shared_by_every_page(self):
        # The ticket drawer used to define it inline, so copy did nothing on Shifts.
        from django.conf import settings

        shell = (settings.BASE_DIR / "static/js/app-shell.js").read_text(encoding="utf-8")
        drawer = (settings.BASE_DIR / "templates/tickets/partials/details_drawer.html").read_text(encoding="utf-8")
        self.assertIn("window.copyTicketNumber = function", shell)
        self.assertNotIn("window.copyTicketNumber", drawer)



class ShiftChangeNotificationTests(TestCase):
    """Agents hear about added, changed and removed shifts: bell and web push, one per action."""

    def setUp(self):
        self.department = make_department("Notify")
        self.manager = make_user(
            "notify-lead", "support", self.department, make_role("Notify lead", can_manage_shifts=True),
        )
        self.agent = make_user("notify-agent", "support", self.department)
        self.day = make_shift_type(self.department, name="Notify day", start=time(9, 0), end=time(17, 0))
        self.late = make_shift_type(self.department, name="Notify late", start=time(13, 0), end=time(21, 0))
        self.night = make_shift_type(self.department, name="Notify night")
        self.client.force_login(self.manager)
        self.now = aware(2026, 10, 4, 8, 0)
        self.today = date(2026, 10, 4)

    def _notes(self, user=None):
        from notifications.models import InAppNotification

        return InAppNotification.objects.filter(
            recipient=user or self.agent, notification_type="shift",
        ).order_by("pk")

    def _post(self, name, data):
        with _both_clocks(self.now, self.today), self.captureOnCommitCallbacks(execute=True):
            return self.client.post(reverse(name), data, HTTP_HX_REQUEST="true")

    def _cell(self, user, day, shift_type=None, **extra):
        return self._post("shifts_cell", {
            "user": user.pk,
            "date": day.isoformat(),
            "department": self.department.pk,
            "shift_type": shift_type.pk if shift_type else "",
            **extra,
        })

    def test_added_shift_notifies_the_agent(self):
        self._cell(self.agent, date(2026, 10, 6), self.day)
        note = self._notes().get()
        self.assertEqual(note.title, "New shift: Tue, Oct 6")
        self.assertEqual(note.message, "Notify day · 9:00 AM – 5:00 PM")
        self.assertEqual(note.link, "/shifts/mine/")
        self.assertIn("مناوبة جديدة", note.title_ar)
        self.assertIn("9:00 ص", note.message_ar)

    def test_changed_and_removed_shifts_notify_but_unchanged_saves_do_not(self):
        row = make_assignment(self.agent, self.day, date(2026, 10, 6))
        edit = {"assignment": row.pk}  # the dialog names the row it edits
        self._cell(self.agent, row.date, self.day, **edit)  # saved as it was
        self.assertFalse(self._notes().exists())
        self._cell(self.agent, row.date, self.late, **edit)
        self.assertEqual(self._notes().last().title, "Shift changed: Tue, Oct 6")
        self.assertEqual(self._notes().last().message, "Now Notify late · 1:00 PM – 9:00 PM")
        self._cell(self.agent, row.date, self.late, start_time_override="14:00", end_time_override="22:00", **edit)
        self.assertEqual(self._notes().last().message, "Now Notify late · 2:00 PM – 10:00 PM")
        self._cell(self.agent, row.date, None, **edit)
        self.assertEqual(self._notes().last().title, "Shift removed: Tue, Oct 6")
        # Quick edits share titles and links; every one arrives.
        self.assertEqual(self._notes().count(), 3)

    def test_check_in_fixes_past_shifts_and_own_shifts_are_quiet(self):
        today_row = make_assignment(self.agent, self.day, self.today)
        today_row.checked_in_at = aware(2026, 10, 4, 7, 50)
        today_row.save()
        self._cell(self.agent, self.today, self.day, clear_check_in="1")
        self._cell(self.agent, date(2026, 10, 1), self.day)  # already over
        self._cell(self.manager, date(2026, 10, 6), self.day)  # the manager's own shift
        self.assertFalse(self._notes().exists())
        self.assertFalse(self._notes(self.manager).exists())

    def test_copy_week_sends_one_summary_per_person(self):
        for day in (date(2026, 10, 3), date(2026, 10, 4), date(2026, 10, 5)):
            make_assignment(self.agent, self.day, day)
        self._post("shifts_copy_week", {"department": self.department.pk, "start": "2026-10-10"})
        note = self._notes().get()
        self.assertEqual(note.title, "Your shifts were updated")
        self.assertEqual(note.message, "Added: Oct 10, Oct 11, Oct 12.")
        self.assertEqual(note.title_ar, "تم تحديث مناوباتك")

    def test_long_summaries_are_cut_short(self):
        from notifications.services import ShiftChange, notify_shift_changes

        rows = [make_assignment(self.agent, self.day, date(2026, 10, 5 + n)) for n in range(7)]
        with _both_clocks(self.now, self.today):
            notify_shift_changes([ShiftChange.of("added", row) for row in rows], self.manager)
        self.assertEqual(self._notes().get().message, "Added: Oct 5, Oct 6, Oct 7, Oct 8, Oct 9, +2 more.")

    def test_night_rotation_apply_notifies_once(self):
        self._post("shifts_calc_rotation", {
            "department": self.department.pk,
            "start": "2026-10-05",
            "end": "2026-10-07",
            "shift_type": self.night.pk,
            "users": [self.agent.pk],
            "action": "apply",
        })
        self.assertTrue(ShiftAssignment.objects.filter(user=self.agent, shift_type=self.night).exists())
        note = self._notes().get()
        self.assertTrue(note.title.startswith("New shift") or note.title == "Your shifts were updated")

    def test_a_failing_notification_does_not_break_the_save(self):
        with patch("notifications.services.notify_shift_changes", side_effect=RuntimeError("push down")):
            response = self._cell(self.agent, date(2026, 10, 6), self.day)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(ShiftAssignment.objects.filter(user=self.agent, date=date(2026, 10, 6)).exists())


@patch("notifications.services._enqueue")
@patch("notifications.services._broadcast_notification")
class ShiftTypeTimeChangeNotifyTests(TestCase):
    """Moving a shift type's times moves people's shifts, so they're told."""

    def test_new_times_notify_upcoming_shifts_without_their_own_times(self, _broadcast, _enqueue):
        from notifications.models import InAppNotification

        department = make_department("Retime")
        manager = make_user("retime-lead", "support", department, make_role("Retime lead", can_manage_shifts=True))
        moved = make_user("retime-moved", "support", department)
        own_times = make_user("retime-own", "support", department)
        shift_type = make_shift_type(department, name="Day", start=time(10, 0), end=time(18, 0))
        soon = timezone.localdate() + timedelta(days=3)
        make_assignment(moved, shift_type, soon)
        make_assignment(own_times, shift_type, soon, start=time(9, 0), end=time(17, 0))
        make_assignment(moved, shift_type, timezone.localdate() - timedelta(days=10))  # long over

        self.client.force_login(manager)
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(reverse("shifts_type_edit", args=[shift_type.pk]), {
                "department": department.pk, "name": "Day", "start_time": "12:00", "end_time": "20:00",
                "colour": "#336699",
            })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(InAppNotification.objects.filter(recipient=moved).count(), 1)
        self.assertFalse(InAppNotification.objects.filter(recipient=own_times).exists())

    def test_renaming_only_notifies_nobody(self, _broadcast, _enqueue):
        from notifications.models import InAppNotification

        department = make_department("Rename")
        manager = make_user("rename-lead", "support", department, make_role("Rename lead", can_manage_shifts=True))
        person = make_user("rename-person", "support", department)
        shift_type = make_shift_type(department, name="Day", start=time(10, 0), end=time(18, 0))
        make_assignment(person, shift_type, timezone.localdate() + timedelta(days=2))

        self.client.force_login(manager)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("shifts_type_edit", args=[shift_type.pk]), {
                "department": department.pk, "name": "Morning", "start_time": "10:00", "end_time": "18:00",
                "colour": "#336699",
            })
        self.assertFalse(InAppNotification.objects.filter(recipient=person).exists())


class Round14RotaTests(TestCase):
    def setUp(self):
        self.department = make_department("R14 rota")
        self.manager = make_user("r14-lead", "support", self.department, make_role("R14 lead", can_manage_shifts=True))
        self.agent = make_user("r14-agent", "support", self.department)
        self.day = make_shift_type(self.department, name="Day", start=time(9, 0), end=time(17, 0))
        self.client.force_login(self.manager)

    def test_saving_a_shift_keeps_the_week_on_screen(self):
        form = self.client.get(reverse("shifts_cell"), {
            "user": self.agent.pk, "date": "2026-11-10", "department": self.department.pk,
            "view": "week", "start": "2026-11-07",
        }, HTTP_HX_REQUEST="true")
        self.assertContains(form, 'name="start" value="2026-11-07"')
        self.assertContains(form, 'name="view" value="week"')
        saved = self.client.post(reverse("shifts_cell"), {
            "user": self.agent.pk, "date": "2026-11-10", "department": self.department.pk,
            "shift_type": self.day.pk, "view": "week", "start": "2026-11-07",
        }, HTTP_HX_REQUEST="true")
        self.assertEqual(saved.context["start"], date(2026, 11, 7))

    def test_remove_is_only_the_confirmed_button(self):
        row = make_assignment(self.agent, self.day, date(2026, 11, 10))
        form = self.client.get(reverse("shifts_cell"), {
            "user": self.agent.pk, "date": "2026-11-10", "department": self.department.pk, "assignment": row.pk,
        }, HTTP_HX_REQUEST="true")
        self.assertNotContains(form, '<option value="">')
        self.assertContains(form, 'hx-confirm="Remove this shift?"')

    def test_add_type_starts_on_the_viewed_department(self):
        form = self.client.get(reverse("shifts_type_add"), {"department": self.department.pk}, HTTP_HX_REQUEST="true")
        self.assertEqual(form.context["form"]["department"].value(), self.department.pk)


class BulkAssignTests(TestCase):
    """Calculator → Bulk assign: the fair-share planner and the card."""

    def setUp(self):
        self.department = make_department("Bulk")
        self.manager = make_user("bulk-lead", "support", self.department, make_role("Bulk lead", can_manage_shifts=True))
        self.people = [
            make_user(f"bulk-{name}", "support", self.department, first_name=name.title())
            for name in ("amal", "badr", "chadi")
        ]
        self.night = make_shift_type(self.department, name="Bulk night")
        self.day = make_shift_type(self.department, name="Bulk day", start=time(9, 0), end=time(17, 0))
        # A Saturday a week or more ahead, so nothing is in the past.
        today = timezone.localdate()
        self.first = today + timedelta(days=7 + (5 - today.weekday()) % 7)
        self.client.force_login(self.manager)

    def _dates(self, count):
        return [self.first + timedelta(days=i) for i in range(count)]

    def _plan(self, people, dates, shift_type=None, existing=None, **kwargs):
        from shifts.services import plan_bulk

        existing = existing if existing is not None else {person.pk: [] for person in people}
        return plan_bulk(people, dates, shift_type or self.night, existing, **kwargs)

    @staticmethod
    def _counts(plan):
        counts = {}
        for item in plan:
            counts[item["user_id"]] = counts.get(item["user_id"], 0) + 1
        return counts

    def _post(self, **data):
        payload = {
            "department": self.department.pk,
            "shift_type": self.night.pk,
            "start": self.first.isoformat(),
            "end": (self.first + timedelta(days=6)).isoformat(),
            "users": [person.pk for person in self.people],
            "action": "preview",
            **data,
        }
        return self.client.post(reverse("shifts_calc_rotation"), payload)

    def test_covered_days_are_not_given_a_second_person(self):
        dates = self._dates(3)
        held = make_assignment(self.people[0], self.night, dates[1])
        plan, gaps = self._plan(
            self.people, dates, existing={self.people[0].pk: [held], self.people[1].pk: [], self.people[2].pk: []},
            covered={dates[1]: 1},
        )
        self.assertEqual(sorted(item["date"] for item in plan), [dates[0], dates[2]])
        self.assertEqual(gaps, [])

    def test_shifts_already_held_in_the_run_count_toward_the_share(self):
        dates = self._dates(4)
        a, b = self.people[:2]
        held = make_assignment(a, self.night, dates[0])
        plan, _gaps = self._plan([a, b], dates, existing={a.pk: [held], b.pk: []}, covered={dates[0]: 1})
        counts = self._counts(plan)
        # a already has one of the four nights: one more for a, two for b.
        self.assertEqual((counts.get(a.pk, 0) + 1, counts.get(b.pk, 0)), (2, 2))

    def test_no_catch_up_burst_for_someone_with_no_recent_nights(self):
        dates = self._dates(6)
        veteran, newcomer = self.people[:2]
        history = [make_assignment(veteran, self.night, self.first - timedelta(days=d)) for d in range(1, 8)]
        plan, _gaps = self._plan([veteran, newcomer], dates, existing={veteran.pk: history, newcomer.pk: []})
        self.assertEqual(sorted(self._counts(plan).values()), [3, 3])

    def test_avoids_a_night_right_before_one_already_held(self):
        dates = self._dates(2)
        a, b = self.people[:2]
        # a works the night after the run starts' first night.
        held = make_assignment(a, self.night, dates[1])
        plan, _gaps = self._plan([a, b], dates[:1], existing={a.pk: [held], b.pk: []})
        self.assertEqual(plan[0]["user_id"], b.pk)
        self.assertFalse(plan[0]["back_to_back"])

    def test_weekend_nights_are_shared_too(self):
        dates = self._dates(14)  # two Friday-Saturday weekends
        plan, _gaps = self._plan(self.people[:2], dates)
        from shifts.services import is_weekend

        weekend = self._counts([item for item in plan if is_weekend(item["date"])])
        self.assertEqual(sorted(weekend.values()), [2, 2])
        self.assertEqual(sorted(self._counts(plan).values()), [7, 7])

    def test_two_people_per_day_and_gaps(self):
        dates = self._dates(3)
        plan, gaps = self._plan(self.people, dates, per_day=2)
        self.assertEqual(len(plan), 6)
        self.assertEqual(sorted(self._counts(plan).values()), [2, 2, 2])
        for day in dates:
            self.assertEqual(len({item["user_id"] for item in plan if item["date"] == day}), 2)
        _plan, short = self._plan(self.people[:1], dates[:1], per_day=2)
        self.assertEqual(short, [{"date": dates[0], "missing": 1}])

    def test_everyone_mode_gives_each_free_person_every_day(self):
        dates = self._dates(3)
        busy = make_assignment(self.people[0], self.day, dates[0])
        plan, _gaps = self._plan(
            self.people, dates, shift_type=self.day,
            existing={self.people[0].pk: [busy], self.people[1].pk: [], self.people[2].pk: []},
            everyone=True,
        )
        self.assertEqual(len(plan), 8)  # 3 people x 3 days, minus the one already held

    def test_preview_skips_weekdays_dates_and_shows_a_mini_rota(self):
        friday = self.first + timedelta(days=6)
        holiday = self.first + timedelta(days=2)
        response = self._post(skip_weekday=["4"], skip_date=[holiday.isoformat()])
        content = response.content.decode()
        self.assertContains(response, "shifts-grid--mini")
        self.assertContains(response, "Skipped day")
        self.assertContains(response, "Skipped date")
        planned_days = set(re.findall(r'name="plan" value="\d+:([\d-]+)"', content))
        self.assertEqual(len(planned_days), 5)
        self.assertNotIn(friday.isoformat(), planned_days)
        self.assertNotIn(holiday.isoformat(), planned_days)
        self.assertContains(response, "Apply 5 shifts")

    def test_apply_creates_exactly_the_previewed_plan(self):
        a, b = self.people[:2]
        days = self._dates(2)
        # A plan the planner would not pick by itself: both nights to a.
        response = self._post(
            action="apply",
            end=days[1].isoformat(),
            users=[a.pk, b.pk],
            plan=[f"{a.pk}:{days[0].isoformat()}", f"{a.pk}:{days[1].isoformat()}", f"{a.pk}:2001-01-01"],
        )
        self.assertContains(response, "Added 2 shifts.")
        self.assertEqual(
            sorted(ShiftAssignment.objects.values_list("user_id", "date")), [(a.pk, days[0]), (a.pk, days[1])],
        )
        self.assertEqual(response["HX-Trigger"], "shiftsChanged")

    def test_past_days_are_skipped(self):
        today = timezone.localdate()
        response = self._post(start=(today - timedelta(days=2)).isoformat(), end=(today - timedelta(days=1)).isoformat())
        self.assertContains(response, "Past")
        self.assertNotContains(response, 'name="plan"')

    def test_long_ranges_and_empty_people_are_refused(self):
        from shifts.services import BULK_MAX_DAYS

        long_end = self.first + timedelta(days=BULK_MAX_DAYS)
        self.assertContains(self._post(end=long_end.isoformat()), f"Choose at most {BULK_MAX_DAYS} days.")
        self.assertContains(self._post(users=[]), "Tick at least one person.")
        self.assertEqual(ShiftAssignment.objects.count(), 0)


class BulkAssignRound17Tests(TestCase):
    """Round 17 findings on Bulk assign, Work hours and the rota."""

    setUp = BulkAssignTests.setUp
    _dates = BulkAssignTests._dates
    _plan = BulkAssignTests._plan
    _counts = staticmethod(BulkAssignTests._counts)
    _post = BulkAssignTests._post

    def test_stale_preview_does_not_overfill_a_day(self):
        days = self._dates(3)
        a, b, c = self.people
        plan = [f"{person.pk}:{day.isoformat()}" for person, day in zip((a, b, a), days)]
        # After the preview, someone else puts c on the night of day 3.
        make_assignment(c, self.night, days[2])
        response = self._post(action="apply", end=days[2].isoformat(), plan=plan)
        self.assertContains(response, "Added 2 shifts.")
        self.assertContains(response, "this day already has enough people on this shift")
        self.assertEqual(ShiftAssignment.objects.filter(date=days[2], shift_type=self.night).count(), 1)

    def test_everyone_mode_is_not_capped(self):
        day = self._dates(1)[0]
        plan = [f"{person.pk}:{day.isoformat()}" for person in self.people]
        response = self._post(action="apply", mode="everyone", end=day.isoformat(), plan=plan)
        self.assertContains(response, "Added 3 shifts.")

    def test_a_clash_with_another_save_is_retried_then_reported(self):
        from unittest.mock import patch

        from django.db import OperationalError

        from shifts import views

        day = self._dates(1)[0]
        plan = [f"{self.people[0].pk}:{day.isoformat()}"]
        real = views._apply_rotation_once
        calls = []

        def clash_once(*args):
            calls.append(1)
            if len(calls) == 1:
                raise OperationalError("deadlock")
            return real(*args)

        with patch.object(views, "_apply_rotation_once", side_effect=clash_once) as once:
            self.assertContains(self._post(action="apply", end=day.isoformat(), plan=plan), "Added 1 shift.")
        self.assertEqual(once.call_count, 2)
        with patch.object(views, "_apply_rotation_once", side_effect=OperationalError("deadlock")):
            response = self._post(action="apply", end=(day + timedelta(days=1)).isoformat(),
                                  plan=[f"{self.people[1].pk}:{(day + timedelta(days=1)).isoformat()}"])
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Someone else changed the rota at the same time.")

    def test_ties_rotate_between_runs(self):
        winners = set()
        for offset in range(10):
            plan, _gaps = self._plan(self.people, [self.first + timedelta(days=offset)])
            winners.add(plan[0]["user_id"])
        self.assertGreater(len(winners), 1)

    def test_weekends_are_balanced_in_a_long_run(self):
        from shifts.services import is_weekend

        people = self.people + [make_user("bulk-dina", "support", self.department)]
        plan, _gaps = self._plan(people, self._dates(28))
        weekend = self._counts([item for item in plan if is_weekend(item["date"])])
        self.assertEqual(sorted(weekend.values()), [2, 2, 2, 2])
        self.assertEqual(sorted(self._counts(plan).values()), [7, 7, 7, 7])

    def test_department_select_does_not_inherit_the_submit_selector(self):
        page = self.client.get(reverse("shifts_calculator"), {"department": self.department.pk})
        self.assertContains(page, '''hx-disabled-elt="#shifts-bulk-form button[type='submit']"''')
        self.assertNotContains(page, '''hx-disabled-elt="find ''')

    def test_rota_changes_are_broadcast_and_the_grid_listens(self):
        from unittest.mock import patch

        day = self._dates(1)[0]
        sent = []

        class Layer:
            async def group_send(self, group, message):
                sent.append((group, message))

        with patch("channels.layers.get_channel_layer", return_value=Layer()), \
                self.captureOnCommitCallbacks(execute=True):
            self._post(action="apply", end=day.isoformat(), plan=[f"{self.people[0].pk}:{day.isoformat()}"])
        self.assertIn(("shifts_rota", {"type": "rota.changed", "departments": [self.department.pk]}), sent)
        # Cell buttons inside the grid must not inherit its own swap.
        self.assertContains(
            self.client.get(reverse("shifts_rota"), {"department": self.department.pk, "start": day.isoformat()}),
            'hx-disinherit="*"',
        )
        rota = self.client.get(reverse("shifts_rota"), {"department": self.department.pk, "start": day.isoformat()})
        self.assertContains(rota, f"rotaChanged[detail.departments.indexOf({self.department.pk}) >= 0] from:body")
        grid_url = reverse("shifts_rota_grid")
        query = {"department": self.department.pk, "view": "week", "start": day.isoformat()}
        first = self.client.get(grid_url, query, HTTP_HX_REQUEST="true")
        self.assertContains(first, 'data-etag="')
        again = self.client.get(grid_url, query, HTTP_HX_REQUEST="true", HTTP_IF_NONE_MATCH=first["ETag"])
        self.assertEqual(again.status_code, 304)


class WorkHoursTypeScopeTests(TestCase):
    setUp = WorkHoursReportTests.setUp
    _check_in = WorkHoursReportTests._check_in
    _september = WorkHoursReportTests._september
    _get = WorkHoursReportTests._get
    def test_archived_and_other_department_types_are_ignored(self):
        self._september()
        archived = make_shift_type(self.department, name="Old", archived=True)
        foreign = make_shift_type(make_department("Facilities"), name="Day", start=time(9, 0), end=time(17, 0))
        report = self._get(department=self.department.pk, shift_type=[archived.pk, foreign.pk]).context["report"]
        self.assertEqual(report.title, "Statement of work hours")
        self.assertEqual(len(report.rows), 10)
        picked = self._get(department=self.department.pk, shift_type=[foreign.pk, self.long_day.pk])
        self.assertEqual(picked.context["report"].title, "Statement of work hours: Long day")
        self.assertContains(picked, f"department={self.department.pk}")

    def test_export_filename_names_the_pick(self):
        self._september()
        query = {"person": self.agent.pk, "start": "2026-09-01", "end": "2026-09-30",
                 "department": self.department.pk, "shift_type": [self.long_day.pk]}
        with _both_clocks(aware(2026, 10, 5, 12, 0), date(2026, 10, 5)):
            response = self.client.get(reverse("shifts_calc_work_export"), query)
        self.assertIn("_long-day.xlsx", response["Content-Disposition"])


class ShiftsRound19Tests(TestCase):
    """Shifts-area findings at 4cca8ad: archive, cell remove, Month → Week, export, bulk."""

    def setUp(self):
        self.department = make_department("Round19")
        self.other = make_department("Aaa first")  # sorts first, like Facilities
        self.manager = make_user("r19-lead", "support", self.department, make_role("R19 lead", can_manage_shifts=True))
        self.agent = make_user("r19-agent", "support", self.department)
        self.night = make_shift_type(self.department, name="R19 night")
        self.client.force_login(self.manager)

    def test_archive_stays_on_the_department(self):
        doomed = make_shift_type(self.department, name="R19 Test", start=time(9, 0), end=time(17, 0))
        plain = self.client.post(reverse("shifts_type_archive", args=[doomed.pk]))
        self.assertEqual(plain.status_code, 302)
        self.assertIn(f"department={self.department.pk}", plain["Location"])
        again = make_shift_type(self.department, name="R19 Test 2", start=time(9, 0), end=time(17, 0))
        hx = self.client.post(
            reverse("shifts_type_archive", args=[again.pk]), {"department": self.department.pk},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(hx.status_code, 200)
        self.assertContains(hx, f'value="{self.department.pk}" selected')
        self.assertContains(hx, "R19 Test 2")
        again.refresh_from_db()
        self.assertTrue(again.archived)

    def test_types_page_asks_before_archiving(self):
        page = self.client.get(reverse("shifts_types"), {"department": self.department.pk})
        self.assertContains(page, 'hx-confirm="Archive R19 night?')
        self.assertContains(page, f'<input type="hidden" name="department" value="{self.department.pk}">')

    def test_cell_remove_button_deletes_the_shift(self):
        on = timezone.localdate() + timedelta(days=3)
        row = make_assignment(self.agent, self.night, on)
        dialog = self.client.get(reverse("shifts_cell"), {
            "user": self.agent.pk, "date": on.isoformat(), "department": self.department.pk,
            "assignment": row.pk,
        }, HTTP_HX_REQUEST="true")
        self.assertContains(dialog, 'name="remove" value="1"')
        response = self.client.post(reverse("shifts_cell"), {
            "user": self.agent.pk, "date": on.isoformat(), "department": self.department.pk,
            "assignment": row.pk, "shift_type": self.night.pk, "remove": "1",
        }, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.assertFalse(ShiftAssignment.objects.filter(pk=row.pk).exists())

    def test_month_to_week_link_starts_on_a_saturday(self):
        from shifts.services import saturday_of

        today = timezone.localdate()
        # A month without today: its first week, from the Saturday on or before the 1st.
        past = (today.replace(day=1) - timedelta(days=40)).replace(day=1)
        outside = self.client.get(reverse("shifts_rota"), {
            "department": self.department.pk, "view": "month", "start": past.isoformat(),
        })
        self.assertContains(outside, f"view=week&start={saturday_of(past).isoformat()}")
        # This month: the week holding today.
        inside = self.client.get(reverse("shifts_rota"), {
            "department": self.department.pk, "view": "month", "start": today.replace(day=1).isoformat(),
        })
        self.assertContains(inside, f"view=week&start={saturday_of(today).isoformat()}")

    def test_work_hours_export_keeps_the_department_and_one_filter(self):
        query = {"person": self.agent.pk, "start": "2026-09-01", "end": "2026-09-30",
                 "department": self.department.pk, "shifts": "all"}
        make_assignment(self.agent, self.night, date(2026, 9, 2))
        with _both_clocks(aware(2026, 10, 5, 12, 0), date(2026, 10, 5)):
            every = self.client.get(reverse("shifts_calc_work"), query).context["export_url"]
            day = make_shift_type(self.department, name="R19 day", start=time(9, 0), end=time(17, 0))
            picked = self.client.get(
                reverse("shifts_calc_work"), {**query, "shift_type": [self.night.pk]},
            ).context["export_url"]
        self.assertIn(f"department={self.department.pk}", every)
        self.assertIn("shifts=all", every)
        self.assertIn(f"department={self.department.pk}", picked)
        self.assertIn(f"shift_type={self.night.pk}", picked)
        self.assertNotIn("shifts=", picked)
        self.assertNotIn(f"shift_type={day.pk}", picked)

    def test_bulk_card_explains_an_empty_department(self):
        page = self.client.get(reverse("shifts_calculator"), {"department": self.other.pk, "part": "rotation"})
        self.assertContains(page, "No shift types in this department yet.")
        self.assertContains(page, "No active support staff in this department.")
        self.assertNotContains(page, "shifts-select-all")


class ShiftsRound20Tests(TestCase):
    """R20 findings at 28f682a: Team rota navigation, default department, legend, remove."""

    def setUp(self):
        self.first = make_department("Aaa facilities")  # first by name, has types
        make_shift_type(self.first, name="Facilities night")
        self.department = make_department("IT")
        self.lead_role = make_role("R20 lead", can_manage_shifts=True)
        self.manager = make_user("r20-lead", "support", self.department, self.lead_role)
        self.agent = make_user("r20-agent", "support", self.department)
        self.night = make_shift_type(self.department, name="IT night")

    def test_shift_pages_open_on_the_managers_own_department(self):
        self.client.force_login(self.manager)
        for name in ("shifts_rota", "shifts_types", "shifts_team", "shifts_calculator"):
            page = self.client.get(reverse(name))
            self.assertEqual(page.context["department"], self.department, name)

    def test_own_department_without_types_falls_back_to_one_with_types(self):
        empty = make_department("Zzz empty")
        lead = make_user("r20-empty-lead", "support", empty, self.lead_role)
        self.client.force_login(lead)
        self.assertEqual(self.client.get(reverse("shifts_rota")).context["department"], self.first)

    def test_team_rota_has_navigation_for_agents(self):
        self.client.force_login(self.agent)
        page = self.client.get(reverse("shifts_team"), {"view": "week", "start": "2026-10-07"})
        self.assertEqual(page.status_code, 200)
        base = f"{reverse('shifts_team')}?department={self.department.pk}"
        self.assertContains(page, f"{base}&view=week&start=2026-09-26")  # Previous
        self.assertContains(page, f"{base}&view=week&start=2026-10-10")  # Next
        self.assertContains(page, f"{base}&view=month&start=2026-10-03")
        self.assertNotContains(page, 'id="team-department"')
        # The links an agent follows carry their own department, which is allowed.
        self.assertEqual(self.client.get(reverse("shifts_team"), {
            "department": self.department.pk, "view": "month", "start": "2026-10-01",
        }).status_code, 200)

    def test_board_legend_lists_missed(self):
        self.client.force_login(self.manager)
        page = self.client.get(reverse("shifts_available"))
        self.assertContains(page, 'shifts-tl-key shifts-tl-bar--missed')

    def test_cell_keeps_an_archived_type_selected(self):
        old = make_shift_type(self.department, name="Old IT day", start=time(9, 0), end=time(17, 0))
        row = make_assignment(self.agent, old, date(2026, 10, 1))
        old.archived = True
        old.save()
        self.client.force_login(self.manager)
        form = self.client.get(reverse("shifts_cell"), {
            "user": self.agent.pk, "date": "2026-10-01", "department": self.department.pk, "assignment": row.pk,
        }, HTTP_HX_REQUEST="true")
        self.assertContains(form, f'<option value="{old.pk}" selected>Old IT day</option>', html=False)
