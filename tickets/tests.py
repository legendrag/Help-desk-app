from pathlib import Path

from django.test import TestCase
from django.conf import settings
from accounts.models import User
from core.models import Branch, Department, Category
from tickets.models import Ticket, TicketMessage


class UnpickedTicketNoticeTests(TestCase):
    def setUp(self):
        # Create Branch, Department, Category
        self.branch = Branch.objects.create(code="TEST", name="Test Branch")
        self.department = Department.objects.create(name="Test Department")
        self.category = Category.objects.create(
            department=self.department,
            name="Test Category",
            default_priority=Ticket.Priority.MEDIUM
        )

        # Create branch user
        self.branch_user = User.objects.create_user(
            username="branch_user",
            email="branch@test.com",
            password="testpassword123",
            user_type=User.UserType.BRANCH,
            branch=self.branch
        )

        # Create support user
        self.support_user = User.objects.create_user(
            username="support_user",
            email="support@test.com",
            password="testpassword123",
            user_type=User.UserType.SUPPORT,
            department=self.department
        )

        # Create unassigned ticket
        self.ticket = Ticket.objects.create(
            ticket_number="TK-1001",
            title="Need assistance",
            description="Trouble logging in.",
            branch=self.branch,
            department=self.department,
            category=self.category,
            created_by=self.branch_user,
            client_name="Test Client",
            client_phone="123456789"
        )

    def test_first_message_no_system_message(self):
        # When branch user sends first message, count is 1. No system message is expected.
        TicketMessage.objects.create(
            ticket=self.ticket,
            sender=self.branch_user,
            message="Hello? Anyone there?"
        )
        
        system_msgs = TicketMessage.objects.filter(ticket=self.ticket, is_system_message=True)
        self.assertEqual(system_msgs.count(), 0)

    def test_second_message_creates_system_message(self):
        # First message
        TicketMessage.objects.create(
            ticket=self.ticket,
            sender=self.branch_user,
            message="Hello? Anyone there?"
        )
        
        # Second message
        TicketMessage.objects.create(
            ticket=self.ticket,
            sender=self.branch_user,
            message="I still need help."
        )

        # Verify that the system message was created
        system_msgs = TicketMessage.objects.filter(ticket=self.ticket, is_system_message=True)
        self.assertEqual(system_msgs.count(), 1)
        
        expected_text = getattr(settings, "TICKET_UNPICKED_SYSTEM_MESSAGE", "Someone will help you soon.")
        self.assertEqual(system_msgs.first().message, expected_text)

    def test_third_message_no_duplicate_system_message(self):
        # Send 3 messages
        TicketMessage.objects.create(
            ticket=self.ticket,
            sender=self.branch_user,
            message="Msg 1"
        )
        TicketMessage.objects.create(
            ticket=self.ticket,
            sender=self.branch_user,
            message="Msg 2"
        )
        TicketMessage.objects.create(
            ticket=self.ticket,
            sender=self.branch_user,
            message="Msg 3"
        )

        # Verify that ONLY ONE system message was created
        system_msgs = TicketMessage.objects.filter(ticket=self.ticket, is_system_message=True)
        self.assertEqual(system_msgs.count(), 1)

    def test_assigned_ticket_no_system_message(self):
        # Assign the ticket to support user
        self.ticket.assigned_to = self.support_user
        self.ticket.status = Ticket.Status.IN_PROGRESS
        self.ticket.save()

        # Send 2 messages from the branch user
        TicketMessage.objects.create(
            ticket=self.ticket,
            sender=self.branch_user,
            message="Msg 1"
        )
        TicketMessage.objects.create(
            ticket=self.ticket,
            sender=self.branch_user,
            message="Msg 2"
        )

        # Verify no system message was created since the ticket is assigned
        system_msgs = TicketMessage.objects.filter(ticket=self.ticket, is_system_message=True)
        self.assertEqual(system_msgs.count(), 0)

    def test_closed_ticket_creates_system_message(self):
        from tickets.models import TicketStatusHistory
        # Close the ticket
        self.ticket.status = Ticket.Status.CLOSED
        self.ticket.save()
        TicketStatusHistory.objects.create(
            ticket=self.ticket,
            status=Ticket.Status.CLOSED,
            event_type=TicketStatusHistory.EventType.STATUS_CHANGE,
            changed_by=self.support_user
        )

        # Verify that the closed system message was created
        system_msgs = TicketMessage.objects.filter(
            ticket=self.ticket, 
            is_system_message=True,
            message=f"Ticket closed by {self.support_user.username}"
        )
        self.assertEqual(system_msgs.count(), 1)

    def test_reopened_ticket_creates_system_message(self):
        from tickets.models import TicketStatusHistory
        # First close the ticket
        self.ticket.status = Ticket.Status.CLOSED
        self.ticket.save()
        TicketStatusHistory.objects.create(
            ticket=self.ticket,
            status=Ticket.Status.CLOSED,
            event_type=TicketStatusHistory.EventType.STATUS_CHANGE,
            changed_by=self.support_user
        )

        # Now reopen the ticket
        self.ticket.status = Ticket.Status.IN_PROGRESS
        self.ticket.save()
        TicketStatusHistory.objects.create(
            ticket=self.ticket,
            status=Ticket.Status.IN_PROGRESS,
            event_type=TicketStatusHistory.EventType.REOPENED,
            changed_by=self.branch_user
        )

        # Verify that the reopened system message was created
        system_msgs = TicketMessage.objects.filter(
            ticket=self.ticket, 
            is_system_message=True,
            message=f"Ticket reopened by {self.branch_user.username}"
        )
        self.assertEqual(system_msgs.count(), 1)


from django.urls import reverse
from core.models import Role

class TicketAuthorizationViewTests(TestCase):
    def setUp(self):
        # Create Branches
        self.branch_a = Branch.objects.create(code="BR-A", name="Branch A")
        self.branch_b = Branch.objects.create(code="BR-B", name="Branch B")

        # Create Departments
        self.dept_a = Department.objects.create(name="Dept A")
        self.dept_b = Department.objects.create(name="Dept B")

        # Create Categories
        self.category_a = Category.objects.create(
            department=self.dept_a,
            name="Category A",
            default_priority=Ticket.Priority.MEDIUM
        )
        self.category_b = Category.objects.create(
            department=self.dept_b,
            name="Category B",
            default_priority=Ticket.Priority.MEDIUM
        )

        # Create Role
        self.role_with_edit = Role.objects.create(
            name="Support Agent with Edit",
            can_update_ticket=True
        )

        # Create Users
        self.branch_user_a = User.objects.create_user(
            username="branch_user_a",
            email="branch_a@test.com",
            password="password123",
            user_type=User.UserType.BRANCH,
            branch=self.branch_a
        )
        self.branch_user_b = User.objects.create_user(
            username="branch_user_b",
            email="branch_b@test.com",
            password="password123",
            user_type=User.UserType.BRANCH,
            branch=self.branch_b
        )

        self.support_user_a = User.objects.create_user(
            username="support_user_a",
            email="support_a@test.com",
            password="password123",
            user_type=User.UserType.SUPPORT,
            department=self.dept_a,
            role=self.role_with_edit
        )
        self.support_user_b = User.objects.create_user(
            username="support_user_b",
            email="support_b@test.com",
            password="password123",
            user_type=User.UserType.SUPPORT,
            department=self.dept_b,
            role=self.role_with_edit
        )

        # Create Tickets
        self.ticket_a = Ticket.objects.create(
            ticket_number="TK-A",
            title="Ticket A",
            description="Branch A Ticket",
            branch=self.branch_a,
            department=self.dept_a,
            category=self.category_a,
            created_by=self.branch_user_a,
            client_name="Client A",
            client_phone="123456789"
        )
        self.ticket_b = Ticket.objects.create(
            ticket_number="TK-B",
            title="Ticket B",
            description="Branch B Ticket",
            branch=self.branch_b,
            department=self.dept_b,
            category=self.category_b,
            created_by=self.branch_user_b,
            client_name="Client B",
            client_phone="987654321"
        )

    def test_detail_view_branch_user_authorized(self):
        self.client.login(username="branch_user_a", password="password123")
        url = reverse("ticket_detail", kwargs={"ticket_id": self.ticket_a.id})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)

    def test_detail_view_branch_user_unauthorized(self):
        self.client.login(username="branch_user_a", password="password123")
        url = reverse("ticket_detail", kwargs={"ticket_id": self.ticket_b.id})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 403)

    def test_detail_view_support_user_authorized(self):
        self.client.login(username="support_user_a", password="password123")
        url = reverse("ticket_detail", kwargs={"ticket_id": self.ticket_a.id})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)

    def test_detail_view_support_user_unauthorized(self):
        self.client.login(username="support_user_a", password="password123")
        url = reverse("ticket_detail", kwargs={"ticket_id": self.ticket_b.id})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 403)

    def test_detail_view_nonexistent_ticket(self):
        self.client.login(username="branch_user_a", password="password123")
        url = reverse("ticket_detail", kwargs={"ticket_id": 99999})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 404)

    def test_edit_view_support_user_authorized(self):
        self.client.login(username="support_user_a", password="password123")
        url = reverse("ticket_update", kwargs={"ticket_id": self.ticket_a.id})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)

    def test_edit_view_support_user_unauthorized(self):
        self.client.login(username="support_user_a", password="password123")
        url = reverse("ticket_update", kwargs={"ticket_id": self.ticket_b.id})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 403)

    def test_edit_view_nonexistent_ticket(self):
        self.client.login(username="support_user_a", password="password123")
        url = reverse("ticket_update", kwargs={"ticket_id": 99999})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 404)


class TicketListSearchTests(TestCase):
    def setUp(self):
        self.branch = Branch.objects.create(code="SRCH", name="Search Branch")
        self.department = Department.objects.create(name="Search Dept")
        self.category = Category.objects.create(
            department=self.department,
            name="VPN Issues",
            default_priority=Ticket.Priority.MEDIUM,
        )
        self.role = Role.objects.create(name="Search Support", can_create_ticket=True)
        self.support = User.objects.create_user(
            username="search_agent",
            email="search@test.com",
            password="password123",
            user_type=User.UserType.SUPPORT,
            department=self.department,
            role=self.role,
        )
        self.creator = User.objects.create_user(
            username="branch_creator",
            email="creator@test.com",
            password="password123",
            user_type=User.UserType.BRANCH,
            branch=self.branch,
        )
        self.by_title = Ticket.objects.create(
            ticket_number="TK-2001",
            title="Cannot reset password",
            description="Login page error",
            branch=self.branch,
            department=self.department,
            category=self.category,
            created_by=self.creator,
            client_name="Alice Smith",
            client_phone="555-0100",
        )
        self.by_client = Ticket.objects.create(
            ticket_number="TK-2002",
            title="Printer offline",
            description="Hardware issue",
            branch=self.branch,
            department=self.department,
            category=self.category,
            created_by=self.creator,
            assigned_to=self.support,
            client_name="Bob Jones",
            client_phone="555-0199",
        )
        self.by_category = Ticket.objects.create(
            ticket_number="TK-2003",
            title="Network dropouts",
            description="Intermittent wifi",
            branch=self.branch,
            department=self.department,
            category=self.category,
            created_by=self.creator,
            client_name="Carol",
            client_phone="555-0111",
        )
        self.client.login(username="search_agent", password="password123")

    def test_search_by_client_name(self):
        response = self.client.get(reverse("tickets_list"), {"q": "Alice"})
        self.assertEqual(response.status_code, 200)
        tickets = list(response.context["tickets"])
        self.assertIn(self.by_title, tickets)
        self.assertNotIn(self.by_client, tickets)

    def test_search_by_phone_digits(self):
        response = self.client.get(reverse("tickets_list"), {"q": "5550199"})
        self.assertEqual(response.status_code, 200)
        tickets = list(response.context["tickets"])
        self.assertIn(self.by_client, tickets)

    def test_search_by_assignee_username(self):
        response = self.client.get(reverse("tickets_list"), {"q": "search_agent"})
        self.assertEqual(response.status_code, 200)
        tickets = list(response.context["tickets"])
        self.assertIn(self.by_client, tickets)
        self.assertNotIn(self.by_title, tickets)

    def test_multi_word_and_search(self):
        response = self.client.get(reverse("tickets_list"), {"q": "reset password"})
        self.assertEqual(response.status_code, 200)
        tickets = list(response.context["tickets"])
        self.assertIn(self.by_title, tickets)
        self.assertNotIn(self.by_client, tickets)

    def test_relevance_ranks_ticket_number_first(self):
        response = self.client.get(reverse("tickets_list"), {"q": "TK-2001"})
        self.assertEqual(response.status_code, 200)
        tickets = list(response.context["tickets"])
        self.assertEqual(tickets[0], self.by_title)

    def test_htmx_partial_updates_load_more(self):
        response = self.client.get(
            reverse("tickets_list"),
            {"q": "Alice"},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "tickets/list_live_partial.html")
        self.assertTrue(response.context["is_htmx"])
        self.assertContains(response, self.by_title.ticket_number)


from django.core.exceptions import ValidationError, PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.test import override_settings
from core.models import Role
from kb.models import Article
from tickets.access import user_can_view_ticket, user_can_pick_ticket, user_can_reopen_ticket


class SecurityXSSAndUploadTests(TestCase):
    def setUp(self):
        self.branch = Branch.objects.create(code="XSS", name="XSS Branch")
        self.department = Department.objects.create(name="XSS Dept")
        self.category = Category.objects.create(
            department=self.department, name="XSS Cat", default_priority=Ticket.Priority.MEDIUM
        )
        self.role = Role.objects.create(name="XSS Branch Role", can_send_message=True)
        self.user = User.objects.create_user(
            username="xss_branch",
            email="xss@test.com",
            password="password123",
            user_type=User.UserType.BRANCH,
            branch=self.branch,
            role=self.role,
        )
        self.ticket = Ticket.objects.create(
            ticket_number="TK-XSS-1",
            title="XSS ticket",
            description="desc",
            branch=self.branch,
            department=self.department,
            category=self.category,
            created_by=self.user,
            client_name="Client",
            client_phone="123",
        )

    def test_message_html_escaped_in_detail(self):
        TicketMessage.objects.create(
            ticket=self.ticket,
            sender=self.user,
            message='<img src=x onerror=alert(1)>',
        )
        self.client.login(username="xss_branch", password="password123")
        response = self.client.get(reverse("ticket_detail", kwargs={"ticket_id": self.ticket.id}))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "<img src=x onerror=alert(1)>")
        self.assertContains(response, "&lt;img src=x onerror=alert(1)&gt;")

    def test_reject_html_attachment(self):
        bad = SimpleUploadedFile("evil.html", b"<script>alert(1)</script>", content_type="text/html")
        with self.assertRaises(ValidationError):
            TicketMessage.objects.create(
                ticket=self.ticket,
                sender=self.user,
                message="",
                attachment=bad,
            )

    def test_reject_svg_attachment(self):
        bad = SimpleUploadedFile("evil.svg", b"<svg onload=alert(1)></svg>", content_type="image/svg+xml")
        with self.assertRaises(ValidationError):
            TicketMessage.objects.create(
                ticket=self.ticket,
                sender=self.user,
                message="",
                attachment=bad,
            )

    def test_accept_png_attachment(self):
        # Minimal 1x1 PNG
        png = (
            b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
            b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N"
            b"\x00\x00\x00\x00IEND\xaeB`\x82"
        )
        good = SimpleUploadedFile("ok.png", png, content_type="image/png")
        msg = TicketMessage.objects.create(
            ticket=self.ticket,
            sender=self.user,
            message="with image",
            attachment=good,
        )
        self.assertTrue(msg.attachment.name.endswith(".png"))


class SecurityTenancyTests(TestCase):
    def setUp(self):
        self.branch_a = Branch.objects.create(code="BA", name="Branch A")
        self.branch_b = Branch.objects.create(code="BB", name="Branch B")
        self.dept_a = Department.objects.create(name="Dept A")
        self.dept_b = Department.objects.create(name="Dept B")
        self.cat_a = Category.objects.create(
            department=self.dept_a, name="Cat A", default_priority=Ticket.Priority.MEDIUM
        )
        self.role_kb = Role.objects.create(
            name="KB Viewer", can_access_kb=True, can_pick_ticket=True, can_send_message=True
        )
        self.role_no_kb = Role.objects.create(
            name="No KB", can_access_kb=False, can_pick_ticket=True, can_send_message=True
        )
        self.branch_user_b_kb = User.objects.create_user(
            username="branch_b_kb",
            email="bbkb@test.com",
            password="password123",
            user_type=User.UserType.BRANCH,
            branch=self.branch_b,
            role=self.role_kb,
        )
        self.branch_user_b_nokk = User.objects.create_user(
            username="branch_b_nokb",
            email="bbnokb@test.com",
            password="password123",
            user_type=User.UserType.BRANCH,
            branch=self.branch_b,
            role=self.role_no_kb,
        )
        self.support_a = User.objects.create_user(
            username="support_a_pick",
            email="sa@test.com",
            password="password123",
            user_type=User.UserType.SUPPORT,
            department=self.dept_a,
            role=self.role_kb,
        )
        self.support_b = User.objects.create_user(
            username="support_b_pick",
            email="sb@test.com",
            password="password123",
            user_type=User.UserType.SUPPORT,
            department=self.dept_b,
            role=self.role_kb,
        )
        self.creator = User.objects.create_user(
            username="creator_a",
            email="ca@test.com",
            password="password123",
            user_type=User.UserType.BRANCH,
            branch=self.branch_a,
            role=self.role_no_kb,
        )
        self.ticket = Ticket.objects.create(
            ticket_number="TK-TEN-1",
            title="Tenancy ticket",
            description="desc",
            branch=self.branch_a,
            department=self.dept_a,
            category=self.cat_a,
            created_by=self.creator,
            status=Ticket.Status.OPEN,
            client_name="Client",
            client_phone="123",
        )
        self.article = Article.objects.create(
            title="Related article",
            content="<p>help</p>",
            is_published=True,
            related_ticket=self.ticket,
            created_by=self.support_a,
        )

    def test_kb_bypass_requires_can_access_kb(self):
        self.assertTrue(user_can_view_ticket(self.branch_user_b_kb, self.ticket))
        self.assertFalse(user_can_view_ticket(self.branch_user_b_nokk, self.ticket))

    def test_kb_bypass_http_allows_kb_user(self):
        self.client.login(username="branch_b_kb", password="password123")
        response = self.client.get(reverse("ticket_detail", kwargs={"ticket_id": self.ticket.id}))
        self.assertEqual(response.status_code, 200)

    def test_kb_bypass_http_denies_non_kb_user(self):
        self.client.login(username="branch_b_nokb", password="password123")
        response = self.client.get(reverse("ticket_detail", kwargs={"ticket_id": self.ticket.id}))
        self.assertEqual(response.status_code, 403)

    def test_cross_dept_cannot_pick(self):
        self.assertFalse(user_can_pick_ticket(self.support_b, self.ticket))
        self.client.login(username="support_b_pick", password="password123")
        response = self.client.post(reverse("pick_ticket", kwargs={"ticket_id": self.ticket.id}))
        self.assertEqual(response.status_code, 403)
        self.ticket.refresh_from_db()
        self.assertIsNone(self.ticket.assigned_to_id)

    def test_same_dept_can_pick(self):
        self.assertTrue(user_can_pick_ticket(self.support_a, self.ticket))
        self.client.login(username="support_a_pick", password="password123")
        response = self.client.post(reverse("pick_ticket", kwargs={"ticket_id": self.ticket.id}))
        self.assertIn(response.status_code, (200, 204, 302))
        self.ticket.refresh_from_db()
        self.assertEqual(self.ticket.assigned_to_id, self.support_a.id)

    def test_cross_dept_cannot_reopen(self):
        self.ticket.status = Ticket.Status.CLOSED
        self.ticket.assigned_to = self.support_a
        self.ticket.save()
        self.assertFalse(user_can_reopen_ticket(self.support_b, self.ticket))
        self.client.login(username="support_b_pick", password="password123")
        response = self.client.post(
            reverse("update_status", kwargs={"ticket_id": self.ticket.id}),
            {"status": Ticket.Status.IN_PROGRESS},
        )
        self.assertEqual(response.status_code, 403)

    def test_same_dept_can_reopen(self):
        self.ticket.status = Ticket.Status.CLOSED
        self.ticket.assigned_to = self.support_a
        self.ticket.save()
        self.assertTrue(user_can_reopen_ticket(self.support_a, self.ticket))
        self.client.login(username="support_a_pick", password="password123")
        response = self.client.post(
            reverse("update_status", kwargs={"ticket_id": self.ticket.id}),
            {"status": Ticket.Status.IN_PROGRESS},
        )
        self.assertIn(response.status_code, (200, 204, 302))
        self.ticket.refresh_from_db()
        self.assertEqual(self.ticket.status, Ticket.Status.IN_PROGRESS)


class TicketNumberCopyButtonTests(TestCase):
    def setUp(self):
        self.branch = Branch.objects.create(code="COPY", name="Copy Branch")
        self.department = Department.objects.create(name="Copy Dept")
        self.category = Category.objects.create(
            department=self.department, name="Copy Cat", default_priority=Ticket.Priority.MEDIUM
        )
        self.user = User.objects.create_user(
            username="copy_branch",
            email="copy@test.com",
            password="password123",
            user_type=User.UserType.BRANCH,
            branch=self.branch,
        )
        self.ticket = Ticket.objects.create(
            ticket_number="TK-COPY-1",
            title="Copy ticket",
            description="desc",
            branch=self.branch,
            department=self.department,
            category=self.category,
            created_by=self.user,
            client_name="Client",
            client_phone="0501234567",
        )

    def test_detail_page_has_copy_buttons_for_ticket_and_phone(self):
        self.client.login(username="copy_branch", password="password123")
        response = self.client.get(reverse("ticket_detail", kwargs={"ticket_id": self.ticket.id}))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="ticket-copy-number-btn"', count=2)
        self.assertContains(response, 'data-copy-text="TK-COPY-1"')
        self.assertContains(response, 'data-copy-text="0501234567"')
        self.assertContains(response, "copyTicketNumber")

    def test_drawer_partial_has_copy_buttons_for_ticket_and_phone(self):
        self.client.login(username="copy_branch", password="password123")
        response = self.client.get(reverse("ticket_drawer", kwargs={"ticket_id": self.ticket.id}))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="ticket-copy-number-btn"', count=2)
        self.assertContains(response, 'data-copy-text="TK-COPY-1"')
        self.assertContains(response, 'data-copy-text="0501234567"')

    def test_drawer_partial_starts_closed_and_opens_after_paint(self):
        self.client.login(username="copy_branch", password="password123")
        response = self.client.get(reverse("ticket_drawer", kwargs={"ticket_id": self.ticket.id}))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'class="ticket-offcanvas"')
        self.assertNotContains(response, 'class="ticket-offcanvas open"')
        self.assertContains(response, "requestAnimationFrame")


class TicketDrawerAnimationTests(TestCase):
    def test_offcanvas_slides_with_transform(self):
        css = (Path(settings.BASE_DIR) / "static" / "css" / "style.css").read_text(
            encoding="utf-8"
        )
        self.assertRegex(
            css,
            r"\.ticket-offcanvas \{[^}]*transform:\s*translateX\(-100%\)",
        )
        self.assertRegex(
            css,
            r"\.ticket-offcanvas \{[^}]*transition:\s*transform",
        )
        self.assertRegex(
            css,
            r"\.ticket-offcanvas\.open \{[^}]*transform:\s*translateX\(0\)",
        )
        self.assertNotRegex(
            css,
            r"\.ticket-offcanvas-overlay \{[^}]*display:\s*none",
        )

    def test_dark_mode_preserves_offcanvas_transform_transition(self):
        css = (Path(settings.BASE_DIR) / "static" / "css" / "dark-mode.css").read_text(
            encoding="utf-8"
        )
        self.assertRegex(
            css,
            r'\[data-theme="dark"\] \.ticket-offcanvas \{[^}]*transition:[^}]*transform',
        )
        self.assertRegex(
            css,
            r'\[data-theme="dark"\] \.ticket-offcanvas-overlay \{[^}]*transition:[^}]*opacity',
        )


class TicketDetailQueryOptimizationTests(TestCase):
    def setUp(self):
        self.branch = Branch.objects.create(code="QO", name="Query Opt Branch")
        self.department = Department.objects.create(name="Query Opt Dept")
        self.category = Category.objects.create(
            department=self.department,
            name="Query Opt Category",
            default_priority=Ticket.Priority.MEDIUM,
        )
        self.branch_user = User.objects.create_user(
            username="qo_branch",
            email="qo_branch@test.com",
            password="password123",
            user_type=User.UserType.BRANCH,
            branch=self.branch,
        )
        self.support_user = User.objects.create_user(
            username="qo_support",
            email="qo_support@test.com",
            password="password123",
            user_type=User.UserType.SUPPORT,
            department=self.department,
        )
        self.other_support = User.objects.create_user(
            username="qo_support2",
            email="qo_support2@test.com",
            password="password123",
            user_type=User.UserType.SUPPORT,
            department=self.department,
        )
        self.target = Ticket.objects.create(
            ticket_number="TK-QO-TARGET",
            title="Merge target",
            description="Target",
            branch=self.branch,
            department=self.department,
            category=self.category,
            created_by=self.branch_user,
            client_name="Client",
            client_phone="0500000000",
        )
        self.ticket = Ticket.objects.create(
            ticket_number="TK-QO-1",
            title="Query opt ticket",
            description="Needs replies",
            branch=self.branch,
            department=self.department,
            category=self.category,
            created_by=self.branch_user,
            client_name="Client",
            client_phone="0500000001",
            pending_transfer_to=self.support_user,
            pending_transfer_by=self.branch_user,
            merged_into=self.target,
        )
        first = TicketMessage.objects.create(
            ticket=self.ticket,
            sender=self.branch_user,
            message="First message",
        )
        TicketMessage.objects.create(
            ticket=self.ticket,
            sender=self.support_user,
            message="Reply to first",
            reply_to=first,
        )

    def test_detail_queryset_prefetches_reply_chain_and_transfer_fks(self):
        from tickets.template_views import TicketDetailView

        ticket = TicketDetailView().get_queryset().get(pk=self.ticket.pk)
        with self.assertNumQueries(0):
            self.assertEqual(ticket.merged_into.ticket_number, "TK-QO-TARGET")
            self.assertEqual(ticket.pending_transfer_to.username, "qo_support")
            self.assertEqual(ticket.pending_transfer_by.username, "qo_branch")
            messages = list(ticket.messages.all())
            reply = next(m for m in messages if m.reply_to_id)
            self.assertEqual(reply.reply_to.sender.username, "qo_branch")
            self.assertEqual(reply.sender.username, "qo_support")

    def test_detail_get_context_data_does_not_refetch_ticket(self):
        from django.db import connection
        from django.test import RequestFactory
        from django.test.utils import CaptureQueriesContext
        from tickets.template_views import TicketDetailView

        request = RequestFactory().get(f"/tickets/{self.ticket.id}/")
        request.user = self.branch_user
        view = TicketDetailView()
        view.setup(request, ticket_id=self.ticket.id)
        view.object = view.get_object()

        with CaptureQueriesContext(connection) as ctx:
            context = view.get_context_data()

        ticket_selects = [
            q["sql"]
            for q in ctx.captured_queries
            if q["sql"].lstrip().upper().startswith("SELECT") and "tickets_ticket" in q["sql"]
        ]
        self.assertEqual(ticket_selects, [])
        self.assertIs(context["ticket"], view.object)

    def test_drawer_get_context_data_does_not_refetch_ticket(self):
        from django.db import connection
        from django.test import RequestFactory
        from django.test.utils import CaptureQueriesContext
        from tickets.template_views import TicketDrawerPartialView

        request = RequestFactory().get(f"/tickets/{self.ticket.id}/drawer/")
        request.user = self.branch_user
        view = TicketDrawerPartialView()
        view.setup(request, ticket_id=self.ticket.id)
        view.object = view.get_object()

        with CaptureQueriesContext(connection) as ctx:
            context = view.get_context_data()

        ticket_selects = [
            q["sql"]
            for q in ctx.captured_queries
            if q["sql"].lstrip().upper().startswith("SELECT") and "tickets_ticket" in q["sql"]
        ]
        self.assertEqual(ticket_selects, [])
        self.assertIs(context["ticket"], view.object)

    def test_supporters_select_related_department(self):
        from django.test import RequestFactory
        from tickets.template_views import TicketDetailView

        request = RequestFactory().get(f"/tickets/{self.ticket.id}/")
        request.user = self.branch_user
        view = TicketDetailView()
        view.setup(request, ticket_id=self.ticket.id)
        view.object = view.get_object()
        supporters = list(view.get_context_data()["supporters"])
        self.assertTrue(any(s.department_id for s in supporters))
        with self.assertNumQueries(0):
            for supporter in supporters:
                if supporter.department_id:
                    _ = supporter.department.name

    def test_detail_view_renders_reply_quote(self):
        self.client.login(username="qo_branch", password="password123")
        response = self.client.get(reverse("ticket_detail", kwargs={"ticket_id": self.ticket.id}))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "qo_branch")
        self.assertContains(response, "Reply to first")


class DashboardFilterPartialTests(TestCase):
    def setUp(self):
        self.branch = Branch.objects.create(code="DASH", name="Dash Branch")
        self.department = Department.objects.create(name="Dash Department")
        self.admin = User.objects.create_user(
            username="dash_admin",
            email="dash_admin@test.com",
            password="password123",
            user_type=User.UserType.SUPPORT,
            is_superuser=True,
            is_staff=True,
        )
        self.role = Role.objects.create(name="Branch Dashboard", can_access_dashboard=True)
        self.branch_user = User.objects.create_user(
            username="dash_branch",
            email="dash_branch@test.com",
            password="password123",
            user_type=User.UserType.BRANCH,
            branch=self.branch,
            role=self.role,
        )
        self.client.login(username="dash_admin", password="password123")

    def test_full_page_uses_htmx_filters_without_full_reload(self):
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "tickets/dashboard.html")
        self.assertContains(response, 'id="dashboard-form"')
        self.assertContains(response, 'id="dashboard-live"')
        self.assertContains(response, 'hx-target="#dashboard-live"')
        self.assertContains(response, "vendor/chart.umd.min.js")
        self.assertNotContains(response, "cdn.jsdelivr.net/npm/chart.js")
        self.assertNotContains(response, "this.form.submit()")
        self.assertNotContains(response, "window.location.replace(window.location.pathname)")

    def test_htmx_request_returns_live_partial(self):
        response = self.client.get(
            reverse("dashboard"),
            {"department": "all"},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "tickets/partials/dashboard_live.html")
        self.assertContains(response, 'id="dashboard-live"')
        self.assertNotContains(response, 'id="dashboard-form"')
        self.assertNotContains(response, 'class="dashboard-stats"')
        self.assertNotContains(response, "<html")

    def test_htmx_filters_apply_to_partial(self):
        other_dept = Department.objects.create(name="Other Dash Dept")
        category = Category.objects.create(
            department=self.department,
            name="Dash Category",
            default_priority=Ticket.Priority.MEDIUM,
        )
        Ticket.objects.create(
            ticket_number="TK-DASH-1",
            title="Dash ticket",
            description="desc",
            branch=self.branch,
            department=self.department,
            category=category,
            created_by=self.admin,
            client_name="Client",
            client_phone="123",
        )
        response = self.client.get(
            reverse("dashboard"),
            {"department": other_dept.id},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total_tickets"], 1)
        self.assertEqual(sum(item["count"] for item in response.context["status_summary"]), 0)

    def test_headline_stats_ignore_dashboard_filters(self):
        other_dept = Department.objects.create(name="Other Dash Dept")
        other_branch = Branch.objects.create(code="DASH2", name="Dash Branch 2")
        category = Category.objects.create(
            department=self.department,
            name="Dash Category",
            default_priority=Ticket.Priority.MEDIUM,
        )
        Ticket.objects.create(
            ticket_number="TK-DASH-1",
            title="Dash ticket",
            description="desc",
            branch=self.branch,
            department=self.department,
            category=category,
            created_by=self.admin,
            client_name="Client",
            client_phone="123",
        )
        response = self.client.get(
            reverse("dashboard"),
            {
                "department": other_dept.id,
                "branch": other_branch.id,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total_tickets"], 1)
        self.assertEqual(response.context["branches_count"], 1)
        self.assertEqual(response.context["departments_count"], 1)
        self.assertContains(response, 'class="dashboard-stats"')
        self.assertEqual(sum(item["count"] for item in response.context["status_summary"]), 0)

    def test_branch_dashboard_htmx_returns_partial(self):
        self.client.logout()
        self.client.login(username="dash_branch", password="password123")
        response = self.client.get(reverse("dashboard"), HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "tickets/partials/branch_dashboard_live.html")
        self.assertContains(response, 'id="dashboard-live"')
        self.assertNotContains(response, "<html")

    def test_branch_dashboard_uses_local_chartjs(self):
        self.client.logout()
        self.client.login(username="dash_branch", password="password123")
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "tickets/branch_dashboard.html")
        self.assertContains(response, "vendor/chart.umd.min.js")
        self.assertNotContains(response, "cdn.jsdelivr.net/npm/chart.js")

    def test_export_excel_returns_workbook(self):
        category = Category.objects.create(
            department=self.department,
            name="Dash Category",
            default_priority=Ticket.Priority.MEDIUM,
        )
        Ticket.objects.create(
            ticket_number="TK-DASH-XLSX",
            title="Export ticket",
            description="desc",
            branch=self.branch,
            department=self.department,
            category=category,
            created_by=self.admin,
            client_name="Client",
            client_phone="123",
        )
        response = self.client.get(reverse("dashboard_export"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        from io import BytesIO
        import openpyxl

        workbook = openpyxl.load_workbook(BytesIO(response.content))
        self.assertIn("Tickets", workbook.sheetnames)
        self.assertIn("Status Summary", workbook.sheetnames)
        ticket_rows = list(workbook["Tickets"].iter_rows(min_row=2, values_only=True))
        self.assertTrue(any(row[0] == "TK-DASH-XLSX" for row in ticket_rows))
        self.assertTrue(
            any(row[5] == "Open" for row in ticket_rows),
            "status display values must be written as strings",
        )


class TicketCreateValidationTests(TestCase):
    def setUp(self):
        from tickets.forms import TicketCreateForm

        self.TicketCreateForm = TicketCreateForm
        self.branch = Branch.objects.create(code="CRT", name="Create Branch")
        self.department = Department.objects.create(name="Create Dept")
        self.category = Category.objects.create(
            department=self.department,
            name="Create Category",
            default_priority=Ticket.Priority.MEDIUM,
        )
        self.role = Role.objects.create(name="Create Role", can_create_ticket=True)
        self.user = User.objects.create_user(
            username="ticket_creator",
            email="creator@test.com",
            password="password123",
            user_type=User.UserType.BRANCH,
            branch=self.branch,
            role=self.role,
        )

    def test_missing_required_fields_name_the_field(self):
        form = self.TicketCreateForm(data={}, user=self.user)
        self.assertFalse(form.is_valid())
        for name in ("title", "description", "department", "category", "client_name", "client_phone"):
            self.assertIn(name, form.errors)
            label = str(form.fields[name].label)
            self.assertIn(label, str(form.errors[name]))

    def test_create_page_shows_named_field_errors(self):
        self.client.login(username="ticket_creator", password="password123")
        response = self.client.post(reverse("ticket_create"), {})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Title is required")
        self.assertContains(response, "Department is required")
        self.assertNotContains(response, "Please complete the following")

    def test_create_page_includes_client_validation(self):
        self.client.login(username="ticket_creator", password="password123")
        response = self.client.get(reverse("ticket_create"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "form-validation.js")


class DashboardAggregateTests(TestCase):
    def setUp(self):
        self.branch = Branch.objects.create(code="AGG", name="Agg Branch")
        self.department = Department.objects.create(name="Agg Department")
        self.category = Category.objects.create(
            department=self.department,
            name="Agg Category",
            default_priority=Ticket.Priority.MEDIUM,
        )
        self.admin = User.objects.create_user(
            username="agg_admin",
            email="agg_admin@test.com",
            password="password123",
            is_superuser=True,
            is_staff=True,
        )
        self.agent = User.objects.create_user(
            username="agg_agent",
            email="agg_agent@test.com",
            password="password123",
            user_type=User.UserType.SUPPORT,
            department=self.department,
        )
        self.slow_agent = User.objects.create_user(
            username="agg_slow",
            email="agg_slow@test.com",
            password="password123",
            user_type=User.UserType.SUPPORT,
            department=self.department,
        )

    def _ticket(self, number, **updates):
        ticket = Ticket.objects.create(
            ticket_number=number,
            title=number,
            description="desc",
            branch=self.branch,
            department=self.department,
            category=self.category,
            created_by=self.admin,
            client_name="Client",
            client_phone="123",
        )
        Ticket.objects.filter(pk=ticket.pk).update(**updates)
        return ticket

    @override_settings(TIME_ZONE="Asia/Riyadh")
    def test_chart_buckets_use_local_monday_weeks(self):
        from datetime import datetime

        from django.utils import timezone

        local_sunday = timezone.make_aware(datetime(2026, 3, 1, 13, 0))
        local_monday = timezone.make_aware(datetime(2026, 3, 2, 1, 30))
        self._ticket("TK-AGG-SUN", created_at=local_sunday)
        self._ticket("TK-AGG-MON", created_at=local_monday)
        self.client.force_login(self.admin)
        params = {"start_date": "2026-02-01", "end_date": "2026-03-31"}

        day = self.client.get(reverse("dashboard"), {**params, "date_view": "day"})
        day_counts = {item["label"]: item["count"] for item in day.context["tickets_by_date"]}
        self.assertEqual(day_counts["03/01/2026"], 1)
        self.assertEqual(day_counts["03/02/2026"], 1)

        week = self.client.get(reverse("dashboard"), {**params, "date_view": "week"})
        week_counts = {item["label"]: item["count"] for item in week.context["tickets_by_date"]}
        self.assertEqual(week_counts["Week of 02/23/2026"], 1)
        self.assertEqual(week_counts["Week of 03/02/2026"], 1)

        month = self.client.get(reverse("dashboard"), {**params, "date_view": "month"})
        month_counts = {item["label"]: item["count"] for item in month.context["tickets_by_date"]}
        self.assertEqual(month_counts["Mar 2026"], 2)
        self.assertEqual(month_counts["Feb 2026"], 0)

        year = self.client.get(reverse("dashboard"), {**params, "date_view": "year"})
        year_counts = {item["label"]: item["count"] for item in year.context["tickets_by_date"]}
        self.assertEqual(year_counts["2026"], 2)
        self.assertNotIn("recent_activity", day.context)

    @override_settings(TIME_ZONE="Asia/Riyadh")
    def test_leaderboard_working_time_matches_clamped_formula(self):
        from datetime import datetime, timedelta

        from django.utils import timezone

        created = timezone.make_aware(datetime(2026, 3, 10, 11, 0))
        self._ticket(
            "TK-AGG-WORK",
            created_at=created,
            picked_at=created + timedelta(hours=1),
            closed_at=created + timedelta(hours=3),
            status=Ticket.Status.CLOSED,
            assigned_to=self.agent,
            total_pending_duration_seconds=600,
        )
        self._ticket(
            "TK-AGG-CLAMP",
            created_at=created,
            picked_at=created,
            closed_at=created + timedelta(minutes=5),
            status=Ticket.Status.CLOSED,
            assigned_to=self.slow_agent,
            total_pending_duration_seconds=1000,
        )
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("dashboard"),
            {"start_date": "2026-03-01", "end_date": "2026-03-31", "date_view": "month"},
        )
        by_name = {row["name"]: row for row in response.context["agent_performance"]}
        agent = by_name["agg_agent"]
        self.assertEqual(agent["total_working_time"], "1h 50m")
        self.assertEqual(agent["score"], 75)
        self.assertEqual(agent["grade"], "B")
        self.assertEqual(by_name["agg_slow"]["total_working_time"], "--")
